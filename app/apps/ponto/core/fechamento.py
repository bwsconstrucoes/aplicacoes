# -*- coding: utf-8 -*-
"""
O FECHAMENTO DA FOLHA PARA A CONTABILIDADE (29/10/2026) — os relatórios e os arquivos .ps do Fortes.

O dono: *"o fechamento da folha é feito do dia 26 do mês anterior ao dia 25 do mês corrente (…) esse
período final é um fechamento, pode necessitar algum ajuste de ponto, então a gente deixa esses dias para
não ter perigo de dar errado, e a gente ter tempo de mandar com antecedência para a contabilidade."*

DOIS PEDAÇOS, porque o Fortes só aceita arquivo do mês corrente — e só do que é INTEIRO:
  · MÊS ANTERIOR (dia 26 ao fim do mês anterior): vai em planilha — as faltas (integrais e parciais) e o
    total de horas extras de cada um. *"O sistema não consegue mais lançar as faltas do mês anterior via
    arquivo, nem horas extras."*
  · MÊS CORRENTE (dia 1 ao 25): as FALTAS INTEIRAS e as HORAS EXTRAS INTEIRAS vão no arquivo .ps (dois
    arquivos: um de faltas, um de horas extras); as faltas parciais e as horas quebradas vão em planilha,
    que a contabilidade lança de outro jeito. *"Meia falta do dia, não consigo lançar. Meia hora também
    não."*
Tudo também num Excel só, com uma aba para cada parte — *"três arquivos: o relatório geral de faltas e
horas extras, e os dois arquivos"*.

AS CONTAS SÃO AS DO ESPELHO (a mesma régua do Quadro e dos relatórios): falta = dia previsto sem batida e
sem justificativa; falta parcial = o dia com batida faltando ou só meio período (vale o comprovado —
`apuracao.pares_comprovados`); atraso ou saída antecipada que passou da tolerância também é hora que
faltou, e vai junto das parciais, marcado. Hora extra = a do espelho.

ESCOLHAS DESTA PRIMEIRA VERSÃO (a confirmar com o dono — estão no HISTORICO):
  · o percentual da hora extra vem da Configuração (`hora_extra.py`): por obra, por estado ou o padrão da
    empresa, para semana, sábado, domingo e feriado — o dia vai pela obra em que a pessoa bateu nele;
  · a hora extra do mês corrente vai no .ps como o TOTAL do período por pessoa e percentual, com a data
    do último dia (25): o inteiro vai no arquivo, os minutos que sobram vão na planilha de horas quebradas
    (com `por_dia=True`, a conta é feita dia a dia);
  · quem tem BANCO DE HORAS não entra nas horas extras nem nas faltas parciais (vão para o banco); a
    falta do dia inteiro entra;
  · o ID Fortes vem do cadastro da Análise de SPs (`analisesps.colaborador.id_fortes`); sem ele, a pessoa
    fica fora do .ps e aparece na aba Conferência.

O FORMATO DO .ps (`prompt_geracao_ps.md`, do dono) é posicional: cabeçalho fixo, uma linha
`3 IIIIII EEE DDMMAAAA QQQQ` por evento e o rodapé com o total. ⚠️ Três pontos ficaram para conferir com um
arquivo real já importado: os espaços do cabeçalho, os zeros do rodapé e se a quantidade de horas é inteira
pura (`0012` = 12 h). Estão em constantes aqui embaixo, para trocar num lugar só.
"""
from __future__ import annotations

import datetime as dt
import io
import logging
import zipfile
from collections import defaultdict
from typing import Optional

from sqlalchemy.engine import Connection

from .. import db, horario
from ..erros import ErroDeValidacao
from . import apuracao, competencias, hora_extra, parametros, relatorios, vigencia
from .ocorrencias import Quem

logger = logging.getLogger("ponto.fechamento")

PARAMETRO_CORTE = "fechamento.dia_corte"
# A HORA EXTRA PEQUENA DESCONSIDERADA (30/10/2026, o dono: "às vezes tem muitos dias com 2 minutos, 1 minuto, oito
# minutos, até 12 minutos (…) eu queria poder decidir se isso vai entrar ou não"): a hora extra do dia de até N
# minutos não entra no fechamento. Escolhida na tela do fechamento (e guardada para a próxima vez); 0 = tudo entra.
PARAMETRO_HE_MINIMA = "fechamento.he_minima"
# AS HORAS A MENOS ABATIDAS DAS EXTRAS (30/10/2026, o dono: "de repente o cara tem meia hora negativa e tem uma hora
# de extra, aí ele vai ficar só com meia hora de extra"): ligado, o atraso e a saída antecipada (o dia completo com
# horas a menos) do período são abatidos das horas extras do mesmo período — primeiro das de menor percentual. A falta
# parcial e a falta inteira NÃO são abatidas: continuam falta. A mesma tolerância (`he_minima`) vale para o atraso.
PARAMETRO_ABATER = "fechamento.abater_negativas"
HE_MINIMA_MAXIMA = 120
CORTE_PADRAO = 25

EVENTO_FALTA = "001"
EVENTO_HE_100 = "003"
EVENTO_HE_50 = "006"
ROTULO_EVENTO = {"001": "Falta integral", "003": "Hora extra 100%", "004": "Hora extra 70%",
                 "005": "Hora extra 60%", "006": "Hora extra 50%"}

# ⚠️ Do exemplo do dono; conferir com um .ps real (ver o topo).
CABECALHO_PS = "07AC        PONTO     0013                                        "
RODAPE_PS_RESTO = "0000000000000000000000000000000.000000000000000.00"
FIM_DE_LINHA_PS = "\r\n"


# ---------------------------------------------------------------------------
# Os períodos
# ---------------------------------------------------------------------------
def he_minima(conn: Connection, pedida=None, por: str = "") -> int:
    """Os minutos de hora extra do dia que não entram. Com `pedida`, grava como a nova escolha."""
    if pedida is not None and str(pedida).strip() != "":
        try:
            n = int(pedida)
        except (TypeError, ValueError):
            raise ErroDeValidacao("minutos da hora extra desconsiderada ilegíveis", campo="he_minima") from None
        if not 0 <= n <= HE_MINIMA_MAXIMA:
            raise ErroDeValidacao(f"de 0 a {HE_MINIMA_MAXIMA} minutos", campo="he_minima")
        parametros.gravar(conn, PARAMETRO_HE_MINIMA, str(n), por)
        return n
    try:
        return max(0, int(parametros.ler(conn, PARAMETRO_HE_MINIMA, "0") or 0))
    except (TypeError, ValueError):
        return 0


def abater_negativas(conn: Connection, pedido=None, por: str = "") -> bool:
    """Abater os atrasos das horas extras? Com `pedido`, grava como a nova escolha."""
    if pedido is not None and str(pedido).strip() != "":
        valor = str(pedido).strip().lower() in ("1", "true", "sim", "on")
        parametros.gravar(conn, PARAMETRO_ABATER, "1" if valor else "0", por)
        return valor
    return (parametros.ler(conn, PARAMETRO_ABATER, "0") or "0") == "1"


_PCT_DO_EVENTO = {cod: pct for pct, cod in hora_extra.EVENTO_DO_PERCENTUAL.items()}


def abater(pool: dict, linhas: list[dict]) -> int:
    """PURA. Abate das horas extras (`pool`: {chave: minutos}, chave = evento ou (evento, data)) os minutos das
    `linhas` de atraso, em ordem; primeiro da hora extra de menor percentual. Muda os dois no lugar; devolve quanto
    abateu."""
    ordem = sorted(pool, key=lambda k: (_PCT_DO_EVENTO.get(k[0] if isinstance(k, tuple) else k, 999), str(k)))
    abatido = 0
    for ln in linhas:
        for k in ordem:
            if ln["minutos"] <= 0:
                break
            tira = min(pool[k], ln["minutos"])
            pool[k] -= tira
            ln["minutos"] -= tira
            abatido += tira
    for k in [k for k, v in pool.items() if v <= 0]:
        del pool[k]
    return abatido


def dia_de_corte(conn: Connection) -> int:
    try:
        n = int(parametros.ler(conn, PARAMETRO_CORTE, str(CORTE_PADRAO)) or CORTE_PADRAO)
    except (TypeError, ValueError):
        n = CORTE_PADRAO
    return n if 1 <= n <= 27 else CORTE_PADRAO


def periodos(competencia, corte: int = CORTE_PADRAO) -> dict:
    """PURA. A folha de `competencia` (o mês da folha): o pedaço do mês anterior (do dia seguinte ao corte
    ao fim do mês) e o do mês corrente (do dia 1 ao corte)."""
    comp = competencias.primeiro_dia(competencia)
    fim_anterior = comp - dt.timedelta(days=1)
    return {"competencia": comp,
            "anterior": (fim_anterior.replace(day=corte + 1), fim_anterior),
            "corrente": (comp, comp.replace(day=corte))}


# ---------------------------------------------------------------------------
# Regras puras
# ---------------------------------------------------------------------------
def obra_do_dia(d: dict, da_pessoa: Optional[str]) -> Optional[str]:
    """PURA. A obra em que a pessoa bateu no dia (a da primeira batida); sem batida, a que o espelho deu ao dia (a
    do dia com batida mais perto — 31/10/2026); sem nenhuma, a do cadastro."""
    return (next((b.get("obra") for b in d.get("batidas") or [] if b.get("obra")), None)
            or d.get("obra_do_dia") or da_pessoa)


def evento_da_hora_extra(d: dict, cfg: Optional[dict] = None, obra: Optional[str] = None,
                         uf: Optional[str] = None) -> str:
    """PURA. O evento do Fortes pelo percentual da regra (da obra, do estado ou o padrão — `hora_extra`)."""
    cfg = cfg or {"padrao": dict(hora_extra.PADRAO), "uf": {}, "obra": {}}
    # o feriado compensado da 12x36 (art. 59-A) não paga o percentual de feriado: a extra é a do dia comum
    feriado = None if d.get("feriado_compensado") else d.get("feriado")
    pct = hora_extra.percentual_do_dia(cfg, dt.date.fromisoformat(d["data"]), feriado, obra, uf)
    return hora_extra.EVENTO_DO_PERCENTUAL[pct]


def tem_banco(p: dict) -> bool:
    return p.get("regime_banco") not in (None, "", "SEM_BANCO")


def falta_do_dia(d: dict) -> Optional[dict]:
    """PURA. A falta do dia, se houver: {tipo, minutos}. Integral = o dia previsto sem batida; parcial = o
    dia com batida faltando ou meio período; atraso/saída = dia completo com débito além da tolerância."""
    if d.get("ocorrencia"):
        return None
    s = d.get("situacao")
    if s == "FALTA":
        return {"tipo": "INTEGRAL", "rotulo": "Falta integral", "minutos": d.get("previsto") or 0}
    if s in apuracao.SITUACOES_PARCIAIS and d.get("debito"):
        return {"tipo": "PARCIAL", "rotulo": "Falta parcial", "minutos": d["debito"]}
    if s == "OK" and d.get("debito"):
        return {"tipo": "ATRASO", "rotulo": "Atraso ou saída antecipada", "minutos": d["debito"]}
    return None


def hhmm(minutos: int) -> str:
    m = int(minutos or 0)
    return f"{m // 60}:{m % 60:02d}"


def id_fortes_ok(valor) -> tuple[Optional[str], Optional[str]]:
    """PURA. (id com 6 dígitos, None) ou (None, motivo)."""
    digitos = "".join(ch for ch in str(valor or "") if ch.isdigit())
    if not digitos or int(digitos) == 0:
        return None, "ID Fortes ausente"
    digitos = digitos.lstrip("0") or "0"
    if len(digitos) > 6:
        return None, f"ID Fortes com mais de 6 dígitos ({digitos})"
    return digitos.zfill(6), None


def linha_ps(id_fortes6: str, evento: str, data: dt.date, quantidade: int) -> str:
    """PURA. `3IIIIIIEEEDDMMAAAAQQQQ`."""
    if not 0 < int(quantidade) <= 9999:
        raise ValueError("quantidade fora de 1 a 9999")
    return f"3{id_fortes6}{evento}{data:%d%m%Y}{int(quantidade):04d}"


def arquivo_ps(linhas: list[str]) -> bytes:
    """PURA. Cabeçalho, as linhas e o rodapé com o total."""
    corpo = [CABECALHO_PS, *linhas, f"Z{len(linhas):05d}{RODAPE_PS_RESTO}"]
    return (FIM_DE_LINHA_PS.join(corpo) + FIM_DE_LINHA_PS).encode("ascii")


# ---------------------------------------------------------------------------
# A conta
# ---------------------------------------------------------------------------
def _ids_fortes(conn: Connection, cpfs: list[str]) -> dict[str, str]:
    if not cpfs or not db.um(conn, """SELECT 1 AS x FROM information_schema.columns WHERE table_schema = 'analisesps'
                                       AND table_name = 'colaborador' AND column_name = 'id_fortes'"""):
        return {}
    return {r["cpf"]: r["id_fortes"] for r in db.todos(
        conn, "SELECT cpf, id_fortes FROM analisesps.colaborador WHERE cpf = ANY(:c)", c=cpfs)}


def _so_digitos(cpf) -> str:
    return "".join(ch for ch in str(cpf or "") if ch.isdigit())


def montar(conn: Connection, quem: Quem, competencia, filtros: Optional[dict] = None, *, por_dia: bool = False,
           minima: Optional[int] = None, abate: Optional[bool] = None, calendario_todos: bool = False) -> dict:
    corte = dia_de_corte(conn)
    minima = he_minima(conn) if minima is None else int(minima)
    abate = abater_negativas(conn) if abate is None else bool(abate)
    per = periodos(competencia, corte)
    (ai, af), (ci, cf) = per["anterior"], per["corrente"]
    vale = vigencia.inicio(conn)
    if not relatorios._uma_por_vez.acquire(blocking=False):
        raise relatorios.Ocupado("outro relatório está sendo montado agora — tente de novo em instantes")
    try:
        # quem não bate ponto fica fora; e só CTPS entra nas contas de horas (29/10/2026, `relatorios.conta_horas`)
        pessoas = [p for p in relatorios._pessoas(conn, quem, filtros or {}, ai, cf)
                   if p.get("ativo_no_ponto", True) and relatorios.conta_horas(p)]
        inicio = max(ai, vale)
        espelhos = list(relatorios._espelhos(conn, pessoas, inicio, cf)) if cf >= inicio else []
    finally:
        relatorios._uma_por_vez.release()
    fortes = _ids_fortes(conn, [_so_digitos(p.get("cpf")) for p in pessoas])
    cfg = hora_extra.ler(conn)
    from . import cadastros
    obras = {str(o["codigo"]).upper(): o for o in cadastros.listar_obras(conn, so_ativas=False) if o.get("codigo")}

    r = {k: [] for k in ("ant_faltas", "ant_extras", "ant_extras_dias", "cor_faltas_integrais",
                         "cor_faltas_parciais", "cor_extras_integrais", "cor_extras_quebradas", "conferencia",
                         "he_dias", "neg_dias", "abatimentos")}
    trab_dias: list[dict] = []
    ps_faltas, ps_extras = defaultdict(int), defaultdict(int)
    pequenas: list[int] = []
    pequenas_neg: list[int] = []
    for p, esp in espelhos:
        base = relatorios._base(p)
        cpf = _so_digitos(p.get("cpf"))
        idf, sem_id = id_fortes_ok(fortes.get(cpf))
        base = {**base, "cpf": cpf, "id_fortes": idf or ""}
        banco_de_horas = tem_banco(p)
        if p.get("no_registro") and not (p.get("tipo_cadastro") or "").strip():
            r["conferencia"].append({**base, "aviso": "a coluna \"Tipo\" (D) da planilha está vazia — entrou como CTPS; "
                                                         "confira o cadastro"})
        if p.get("situacao") == "AFASTADO":
            r["conferencia"].append({**base, "aviso": "afastada no cadastro — confira se as faltas não são do "
                                                         "afastamento (lance a licença, se for)"})
        he_ant, he_cor = defaultdict(int), defaultdict(int)
        atrasos_ant, atrasos_cor = [], []
        teve_algo = False
        for d in esp["dias"]:
            dia = dt.date.fromisoformat(d["data"])
            anterior = ai <= dia <= af
            corrente = ci <= dia <= cf
            if not (anterior or corrente):
                continue
            if d.get("compensacao_a_cumprir"):
                # a folga já está abonada, mas o dia trabalhado ainda não chegou: se não cumprir, vira falta
                r["conferencia"].append({**base, "aviso": f"folga compensatória em {_br(d['data'])} ainda a cumprir — "
                                                             f"o dia trabalhado é {_br(d['compensacao_a_cumprir'])}; "
                                                             "se não for cumprido, a folga vira falta"})
            if "FORA_DO_CICLO" in (d.get("alertas") or []):
                # o 12x36 automático (30/10/2026): a noite fora do ciclo pede a decisão do DP
                r["conferencia"].append({**base, "aviso": f"trabalhou em {_br(d['data'])}, noite de folga do ciclo 12x36 — "
                                                             "troca de plantão: lance a compensação (a folga e o dia "
                                                             "trabalhado); extra: deixe como está; o ciclo mudou: dê a "
                                                             "escala de novo a partir desse dia"})
            f = falta_do_dia(d)
            # o atraso pequeno (até a mesma tolerância da hora extra) não entra (30/10/2026)
            atraso_pequeno = bool(f and f["tipo"] == "ATRASO" and minima and f["minutos"] <= minima)
            if f:
                # as horas a menos, no calendário da supervisão (30/10/2026: "os dias com hora negativa não estão
                # aparecendo"): todas, inclusive as de quem tem banco e as pequenas (marcadas)
                obn = obra_do_dia(d, p.get("obra_codigo"))
                obn = str(obn).upper() if obn else ""
                r["neg_dias"].append({**base, "data": d["data"], "obra_do_dia": obn,
                                      "obra_nome": (obras.get(obn) or {}).get("nome") or "", "minutos": f["minutos"],
                                      "tipo": f["tipo"], "rotulo": f["rotulo"], "banco": banco_de_horas,
                                      "desconsiderada": atraso_pequeno})
            if (d.get("trabalhado") or 0) > 0 and not f and not (d.get("extra") or d.get("extra_negada")):
                # o dia trabalhado sem hora a mais nem a menos (30/10/2026, o dono: "nos dias em que não há, mesmo
                # que tenha tido batida, fica vazio"): um ✓ no calendário da supervisão
                obt = obra_do_dia(d, p.get("obra_codigo"))
                obt = str(obt).upper() if obt else ""
                trab_dias.append({**base, "data": d["data"], "obra_do_dia": obt,
                                  "obra_nome": (obras.get(obt) or {}).get("nome") or "",
                                  "minutos": d.get("trabalhado") or 0, "banco": banco_de_horas})
            if atraso_pequeno and not banco_de_horas:
                pequenas_neg.append(f["minutos"])
                f = None
            if f and banco_de_horas and f["tipo"] != "INTEGRAL":
                f = None                     # vai para o banco de horas
            if f:
                teve_algo = True
                linha = {**base, "data": d["data"], "dia_semana": d.get("dia_semana"), "tipo": f["rotulo"],
                         "minutos": f["minutos"], "previsto": d.get("previsto") or 0,
                         "batidas": " ".join(b["hora"] for b in d.get("batidas") or [])}
                if f["tipo"] == "ATRASO":
                    (atrasos_ant if anterior else atrasos_cor).append(linha)
                if anterior:
                    r["ant_faltas"].append(linha)
                elif f["tipo"] == "INTEGRAL":
                    r["cor_faltas_integrais"].append(linha)
                    if idf:
                        ps_faltas[(idf, EVENTO_FALTA, dia)] += 1
                else:
                    r["cor_faltas_parciais"].append(linha)
            if d.get("extra") or d.get("extra_negada"):
                # para a supervisão conferir (29/10/2026): toda hora extra do dia, na obra em que foi feita —
                # inclusive a de quem tem banco de horas (marcada) e a negada na conferência (marcada, fora
                # das contas: o espelho já a zerou)
                ob = obra_do_dia(d, p.get("obra_codigo"))
                ob_up = str(ob).upper() if ob else None
                ev_sup = evento_da_hora_extra(d, cfg, ob_up, (obras.get(ob_up) or {}).get("uf"))
                negada = d.get("he_negada")
                r["he_dias"].append({**base, "data": d["data"], "obra_do_dia": ob_up or "",
                                     "obra_nome": (obras.get(ob_up) or {}).get("nome") or "",
                                     "minutos": d.get("extra") or d.get("extra_negada") or 0, "evento": ev_sup,
                                     "banco": banco_de_horas, "negada": bool(not d.get("extra") and negada),
                                     "desconsiderada": bool(d.get("extra") and d["extra"] <= minima),
                                     "negada_motivo": (negada or {}).get("motivo"),
                                     "negada_por": (negada or {}).get("por")})
            if d.get("extra") and d["extra"] <= minima and not banco_de_horas:
                pequenas.append(d["extra"])          # a hora extra pequena desconsiderada: fora das contas
            elif d.get("extra") and not banco_de_horas:
                teve_algo = True
                ev = ev_sup
                if anterior:
                    he_ant[ev] += d["extra"]
                    r["ant_extras_dias"].append({**base, "data": d["data"], "dia_semana": d.get("dia_semana"),
                                                 "evento": ev, "rotulo": ROTULO_EVENTO[ev], "minutos": d["extra"],
                                                 "feriado": d.get("feriado") or ""})
                else:
                    he_cor[(ev, dia if por_dia else cf)] += d["extra"]
        if abate and not banco_de_horas:
            for periodo, pool, linhas_atraso in (("anterior", he_ant, atrasos_ant), ("corrente", he_cor, atrasos_cor)):
                tirou = abater(pool, linhas_atraso)
                if tirou:
                    r["abatimentos"].append({**base, "periodo": periodo, "minutos": tirou})
        if he_ant:
            r["ant_extras"].append({**base, **{f"he{pct}": he_ant.get(cod, 0) for pct, cod in
                                               hora_extra.EVENTO_DO_PERCENTUAL.items()},
                                    "minutos": sum(he_ant.values())})
        for (ev, data), minutos in sorted(he_cor.items()):
            inteiras, resto = divmod(minutos, 60)
            comum = {**base, "evento": ev, "rotulo": ROTULO_EVENTO[ev], "data": data.isoformat(), "minutos_total": minutos}
            if inteiras:
                r["cor_extras_integrais"].append({**comum, "horas": inteiras})
                if idf:
                    ps_extras[(idf, ev, data)] += inteiras
            if resto:
                r["cor_extras_quebradas"].append({**comum, "minutos": resto})
        if teve_algo and banco_de_horas:
            r["conferencia"].append({**base, "aviso": "tem banco de horas: as horas extras e as faltas parciais "
                                                         "ficaram fora (vão para o banco); a falta integral entra"})
        if sem_id and (any(x["cpf"] == cpf for x in r["cor_faltas_integrais"])
                       or any(x["cpf"] == cpf for x in r["cor_extras_integrais"])):
            r["conferencia"].append({**base, "aviso": f"IGNORADA no arquivo .ps — {sem_id}; lance à mão ou "
                                                         "corrija o ID Fortes na Análise de SPs"})

    # a troca de plantão identificada sozinha (31/10/2026, trocas.py): para o DP confirmar antes de fechar
    from . import trocas
    por_id = {int(p["id"]): p for p, _ in espelhos}
    for t in trocas.detectar({int(p["id"]): esp for p, esp in espelhos}):
        for eu, outro in ((t["a"], t["b"]), (t["b"], t["a"])):
            r["conferencia"].append({**relatorios._base(por_id[eu["id"]]),
                                     "aviso": f"troca de plantão provável com {outro['nome']}: trabalhou a noite de "
                                              f"{_br(eu['cobriu'])} (do colega) e o colega a de {_br(eu['deixou'])} (dele) "
                                              "— confirme na janela da pessoa › Espelho do mês"})
    # o atraso todo abatido sai da planilha de atrasos (o que sobrou fica, para descontar)
    for chave in ("ant_faltas", "cor_faltas_parciais"):
        r[chave] = [x for x in r[chave] if x["minutos"] > 0]
    # o ✓ do dia trabalhado: só de quem já está no calendário (tem hora a mais ou a menos no período) — o calendário
    # é das horas extras; sem isto, entraria todo mundo que trabalhou
    # O CALENDÁRIO GERAL (31/10/2026, o dono: "só aparece o que tem hora extra. E se quisermos ver o calendário
    # geral, com ou sem hora, não tem como"): com `calendario_todos`, entra também quem só trabalhou no horário.
    no_calendario = {x["colaborador_id"] for x in r["he_dias"] + r["neg_dias"]}
    r["trab_dias"] = trab_dias if calendario_todos else [x for x in trab_dias if x["colaborador_id"] in no_calendario]
    r["calendario_todos"] = bool(calendario_todos)
    linhas_faltas = [linha_ps(i, e, d, q) for (i, e, d), q in sorted(ps_faltas.items(), key=lambda x: (x[0][0], x[0][2]))]
    linhas_extras = [linha_ps(i, e, d, min(q, 9999)) for (i, e, d), q in
                     sorted(ps_extras.items(), key=lambda x: (x[0][0], x[0][2], x[0][1]))]
    for chave in ("ant_faltas", "ant_extras_dias", "cor_faltas_integrais", "cor_faltas_parciais",
                  "cor_extras_integrais", "cor_extras_quebradas"):
        r[chave].sort(key=lambda x: (x["nome"], x.get("data") or ""))
    r["ant_extras"].sort(key=lambda x: x["nome"])
    return {
        "competencia": f"{per['competencia']:%Y-%m}", "corte": corte, "por_dia": por_dia,
        "anterior": {"de": ai.isoformat(), "ate": af.isoformat()},
        "corrente": {"de": ci.isoformat(), "ate": cf.isoformat()},
        "vale_desde": vale.isoformat(), "pessoas": len(pessoas), "gerado_em": horario.texto(horario.agora()),
        "regras_hora_extra": hora_extra.para_json(cfg)["regras"], "he_minima": minima, "abater": abate,
        **r, "ps_faltas": linhas_faltas, "ps_extras": linhas_extras,
        "totais": {
            "ant_faltas_integrais": sum(1 for x in r["ant_faltas"] if x["tipo"] == "Falta integral"),
            "ant_faltas_parciais": sum(1 for x in r["ant_faltas"] if x["tipo"] != "Falta integral"),
            "ant_extras_minutos": sum(x["minutos"] for x in r["ant_extras"]),
            "cor_faltas_integrais": len(r["cor_faltas_integrais"]),
            "cor_faltas_parciais": len(r["cor_faltas_parciais"]),
            "cor_extras_horas": sum(x["horas"] for x in r["cor_extras_integrais"]),
            "cor_extras_quebradas_minutos": sum(x["minutos"] for x in r["cor_extras_quebradas"]),
            "ps_faltas": len(linhas_faltas), "ps_extras": len(linhas_extras),
            "conferencia": len(r["conferencia"]),
            "he_desconsideradas_dias": len(pequenas), "he_desconsideradas_minutos": sum(pequenas),
            "atrasos_desconsiderados_dias": len(pequenas_neg), "atrasos_desconsiderados_minutos": sum(pequenas_neg),
            "abatido_minutos": sum(x["minutos"] for x in r["abatimentos"]),
        },
    }


# ---------------------------------------------------------------------------
# Os arquivos
# ---------------------------------------------------------------------------
PESSOA = [("Nome", "nome"), ("CPF", "cpf"), ("ID Fortes", "id_fortes"), ("Função", "funcao"), ("Obra", "obra")]


def _br(iso) -> str:
    return dt.date.fromisoformat(iso).strftime("%d/%m/%Y") if iso else ""


def _horas(m) -> float:
    return round((m or 0) / 60, 2)


# Cada parte: (título da aba, chave, colunas [(título, função da linha)]).
def _partes(f: dict) -> list[tuple[str, str, str, list]]:
    pessoa = [(t, (lambda k: lambda x: x.get(k) or "")(k)) for t, k in PESSOA]
    falta = [*pessoa, ("Data", lambda x: _br(x["data"])), ("Dia", lambda x: x.get("dia_semana") or ""),
             ("Tipo", lambda x: x["tipo"]), ("Horas que faltaram", lambda x: hhmm(x["minutos"])),
             ("Horas (decimal)", lambda x: _horas(x["minutos"])), ("Batidas do dia", lambda x: x.get("batidas") or "")]
    ant, cor = f["anterior"], f["corrente"]
    pa = f"{_br(ant['de'])} a {_br(ant['ate'])}"
    pc = f"{_br(cor['de'])} a {_br(cor['ate'])}"
    return [
        ("ant_faltas", "Ant - Faltas", f"Mês anterior ({pa}) — faltas integrais e parciais", falta),
        ("ant_extras", "Ant - Horas extras", f"Mês anterior ({pa}) — total de horas extras por pessoa",
         [*pessoa, *[(f"HE {pct}%", (lambda k: lambda x: hhmm(x.get(k)))(f"he{pct}")) for pct in hora_extra.PERCENTUAIS],
          ("Total", lambda x: hhmm(x["minutos"])), ("Total (decimal)", lambda x: _horas(x["minutos"]))]),
        ("ant_extras_dias", "Ant - HE por dia", f"Mês anterior ({pa}) — horas extras dia a dia (para conferir)",
         [*pessoa, ("Data", lambda x: _br(x["data"])), ("Dia", lambda x: x.get("dia_semana") or ""),
          ("Evento", lambda x: x["rotulo"]), ("Horas", lambda x: hhmm(x["minutos"])), ("Feriado", lambda x: x["feriado"])]),
        ("cor_faltas_integrais", "Faltas integrais", f"Mês corrente ({pc}) — faltas integrais (vão no arquivo .ps)", falta),
        ("cor_faltas_parciais", "Faltas parciais", f"Mês corrente ({pc}) — faltas parciais, atrasos e saídas (lançar à mão)", falta),
        ("cor_extras_integrais", "HE integrais", f"Mês corrente ({pc}) — horas extras inteiras (vão no arquivo .ps)",
         [*pessoa, ("Evento", lambda x: x["rotulo"]), ("Código", lambda x: x["evento"]),
          ("Data no arquivo", lambda x: _br(x["data"])), ("Horas inteiras", lambda x: x["horas"]),
          ("Total do período", lambda x: hhmm(x["minutos_total"]))]),
        ("cor_extras_quebradas", "HE quebradas", f"Mês corrente ({pc}) — os minutos que sobram das horas extras (lançar à mão)",
         [*pessoa, ("Evento", lambda x: x["rotulo"]), ("Código", lambda x: x["evento"]),
          ("Data", lambda x: _br(x["data"])), ("Minutos", lambda x: x["minutos"]),
          ("Total do período", lambda x: hhmm(x["minutos_total"]))]),
        ("conferencia", "Conferência", "Para conferir antes de mandar", [*pessoa, ("Aviso", lambda x: x["aviso"])]),
    ]


PARTES = ("ant_faltas", "ant_extras", "ant_extras_dias", "cor_faltas_integrais", "cor_faltas_parciais",
          "cor_extras_integrais", "cor_extras_quebradas", "conferencia")


def excel(f: dict, so: Optional[str] = None) -> bytes:
    """O relatório geral (todas as abas, com o resumo na frente) ou uma parte só (`so`)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    negrito, cab = Font(bold=True), PatternFill("solid", fgColor="E8EEF6")
    partes = [x for x in _partes(f) if not so or x[0] == so]
    if not partes:
        raise ErroDeValidacao("parte do fechamento desconhecida", campo="parte")
    primeira = True
    if not so:
        ws = wb.active
        ws.title = "Resumo"
        primeira = False
        t = f["totais"]
        ant, cor = f["anterior"], f["corrente"]
        linhas = [
            (f"Fechamento da folha de {f['competencia'][5:]}/{f['competencia'][:4]}", ""),
            ("Gerado em", f["gerado_em"]),
            ("", ""),
            (f"MÊS ANTERIOR ({_br(ant['de'])} a {_br(ant['ate'])}) — em planilha", ""),
            ("Faltas integrais", t["ant_faltas_integrais"]), ("Faltas parciais, atrasos e saídas", t["ant_faltas_parciais"]),
            ("Horas extras (total)", hhmm(t["ant_extras_minutos"])),
            ("", ""),
            (f"MÊS CORRENTE ({_br(cor['de'])} a {_br(cor['ate'])})", ""),
            ("Faltas integrais (arquivo .ps de faltas)", t["cor_faltas_integrais"]),
            ("Linhas no .ps de faltas", t["ps_faltas"]),
            ("Faltas parciais, atrasos e saídas (planilha)", t["cor_faltas_parciais"]),
            ("Horas extras inteiras (arquivo .ps de horas extras)", t["cor_extras_horas"]),
            ("Linhas no .ps de horas extras", t["ps_extras"]),
            ("Horas extras quebradas (planilha)", hhmm(t["cor_extras_quebradas_minutos"])),
            (f"Hora extra do dia de até {f.get('he_minima') or 0} min, desconsiderada",
             f"{t.get('he_desconsideradas_dias', 0)} dia(s), {hhmm(t.get('he_desconsideradas_minutos', 0))}"
             if f.get("he_minima") else "nenhuma (tudo entra)"),
            ("", ""),
            ("Para conferir (aba Conferência)", t["conferencia"]),
            ("", ""),
            *[("Percentual da hora extra" if i == 0 else "",
               f"{x['rotulo']}: semana {x['semana']}% · sábado {x['sabado']}% · domingo {x['domingo']}% · feriado {x['feriado']}%")
              for i, x in enumerate(f.get("regras_hora_extra") or [])],
            ("", "a hora extra do mês corrente vai no .ps " + ("dia a dia" if f.get("por_dia") else
                                                                 "como o total do período, com a data do dia " + _br(cor["ate"]))),
            ("", "quem tem banco de horas: horas extras e faltas parciais vão para o banco"),
            ("Atrasos abatidos das horas extras", hhmm(t.get("abatido_minutos", 0)) if f.get("abater")
             else "não (o abatimento está desligado: os atrasos vão para desconto)"),
            (f"Atrasos de até {f.get('he_minima') or 0} min, desconsiderados",
             f"{t.get('atrasos_desconsiderados_dias', 0)} dia(s), {hhmm(t.get('atrasos_desconsiderados_minutos', 0))}"
             if f.get("he_minima") else "nenhum"),
        ]
        for a, b in linhas:
            ws.append([a, b])
        ws["A1"].font = Font(bold=True, size=13)
        for c in ("A4", "A9"):
            ws[c].font = negrito
        ws.column_dimensions["A"].width = 52
        ws.column_dimensions["B"].width = 60
    for chave, aba, titulo, colunas in partes:
        ws = wb.active if primeira else wb.create_sheet()
        primeira = False
        ws.title = aba[:31]
        ws.append([titulo])
        ws["A1"].font = Font(bold=True, size=12)
        ws.append([c[0] for c in colunas])
        for cel in ws[2]:
            cel.font, cel.fill = negrito, cab
        for x in f[chave]:
            ws.append([fn(x) for _, fn in colunas])
        if not f[chave]:
            ws.append(["(nada neste período)"])
        for i, (nome, _) in enumerate(colunas):
            ws.column_dimensions[chr(65 + i) if i < 26 else "A"].width = max(12, min(45, len(nome) + 6))
        ws.freeze_panes = "A3"
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


# ---------------------------------------------------------------------------
# AS HORAS EXTRAS PARA A SUPERVISÃO CONFERIR (29/10/2026). O dono: *"a gente sempre encaminha o relatório para
# os encarregados e supervisores de obra (…) para saber se tem coerência aquela quantidade de horas extras de
# cada funcionário (…) visualmente é uma espécie de calendário, cada coluna do dia (…) agrupado por obra (…) em
# cada dia a quantidade de horas extras e um totalizador ao final (…) em paisagem, para caber tudo."*
# O período é o da folha inteira (26 a 25); a obra é a do dia (onde a pessoa bateu), então quem trabalhou em
# duas obras aparece nas duas, cada uma com os seus dias.
# ---------------------------------------------------------------------------
def supervisao(f: dict) -> dict:
    """PURA. {dias: [iso], obras: [{obra, nome, pessoas: [{…, dias: {iso: min}, total}], por_dia, total}]}."""
    de, ate = dt.date.fromisoformat(f["anterior"]["de"]), dt.date.fromisoformat(f["corrente"]["ate"])
    dias = [(de + dt.timedelta(days=i)).isoformat() for i in range((ate - de).days + 1)]
    grupos: dict = {}

    def pessoa_do_grupo(x):
        g = grupos.setdefault(x["obra_do_dia"] or "", {"obra": x["obra_do_dia"] or "", "nome": x.get("obra_nome") or "",
                                                       "pessoas": {}, "por_dia": defaultdict(int), "total": 0,
                                                       "atrasos": 0})
        return g, g["pessoas"].setdefault(x["colaborador_id"], {
            "nome": x["nome"], "funcao": x.get("funcao") or "", "cpf": x.get("cpf"), "banco": x.get("banco"),
            "dias": defaultdict(int), "negadas": {}, "pequenas": {}, "negativos": {}, "trabalhou": {}, "total": 0,
            "atrasos": 0})

    # as horas a menos (30/10/2026): em vermelho no calendário; o atraso (dia completo com horas a menos) soma em
    # "atrasos" e, com o abatimento ligado, sai do saldo; falta parcial e falta inteira só aparecem
    for x in f.get("neg_dias") or []:
        g, p = pessoa_do_grupo(x)
        p["negativos"][x["data"]] = {"minutos": x["minutos"], "tipo": x["tipo"], "desconsiderada": x.get("desconsiderada")}
        if x["tipo"] == "ATRASO" and not x.get("desconsiderada"):
            p["atrasos"] += x["minutos"]
            g["atrasos"] += x["minutos"]
    for x in f.get("he_dias") or []:
        g, p = pessoa_do_grupo(x)
        if x.get("negada"):                 # negada na conferência: aparece, mas fora da soma
            p["negadas"][x["data"]] = x["minutos"]
            continue
        if x.get("desconsiderada"):         # pequena demais (até a tolerância escolhida): idem
            p["pequenas"][x["data"]] = x["minutos"]
            continue
        p["dias"][x["data"]] += x["minutos"]
        p["total"] += x["minutos"]
        g["por_dia"][x["data"]] += x["minutos"]
        g["total"] += x["minutos"]
    # o dia trabalhado sem hora a mais nem a menos (30/10/2026): ✓, só de quem já está no calendário
    no_calendario = {x["colaborador_id"] for x in (f.get("he_dias") or []) + (f.get("neg_dias") or [])}
    for x in f.get("trab_dias") or []:
        if f.get("calendario_todos") or x["colaborador_id"] in no_calendario:
            pessoa_do_grupo(x)[1]["trabalhou"][x["data"]] = x["minutos"]
    obras = []
    for g in sorted(grupos.values(), key=lambda g: (not g["obra"], g["obra"])):
        obras.append({**g, "pessoas": sorted(({**p, "dias": dict(p["dias"])} for p in g["pessoas"].values()),
                                              key=lambda p: p["nome"]),
                      "por_dia": dict(g["por_dia"])})
    for g in obras:
        for p in g["pessoas"]:
            p["saldo"] = p["total"] - (p["atrasos"] if f.get("abater") and not p.get("banco") else 0)
    return {"dias": dias, "obras": obras, "total": sum(g["total"] for g in obras)}


def excel_supervisao(f: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    s = supervisao(f)
    wb = Workbook()
    ws = wb.active
    ws.title = "Horas extras por obra"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    negrito, cab, obra_fill = Font(bold=True), PatternFill("solid", fgColor="E8EEF6"), PatternFill("solid", fgColor="D9E2F0")
    azul, vermelho, cinza = Font(color="1D4ED8"), Font(color="B91C1C"), Font(color="9CA3AF")
    trabalhou = Font(color="6B7280")
    comp = f["competencia"]
    ws.append([f"Horas extras para conferência da supervisão — folha de {comp[5:]}/{comp[:4]} "
               f"({_br(s['dias'][0])} a {_br(s['dias'][-1])})"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Gerado em {f['gerado_em']}. Em cada dia, as horas extras daquele dia (h:mm), na obra em que a pessoa "
               "bateu o ponto. (banco) = a pessoa tem banco de horas: a hora extra vai para o banco, não para a folha. "
               + (f"Entre parênteses = até {f['he_minima']} min, desconsiderada. " if f.get("he_minima") else "")
               + "\"negada\" = negada na conferência. As duas ficam fora do total. Em vermelho, as horas a menos "
               "(atraso ou saída antecipada; \"falta\" = o dia inteiro). ✓ = trabalhou no horário, sem hora a mais "
               "nem a menos. "
               + ("O saldo é a hora extra menos os atrasos (o abatimento está ligado)." if f.get("abater")
                  else "Os atrasos não são abatidos das horas extras (o abatimento está desligado).")])
    dias_cab = [dt.date.fromisoformat(d).strftime("%d/%m") for d in s["dias"]]
    for g in s["obras"]:
        ws.append([])
        ws.append([f"Obra {g['obra'] or 'sem obra'}" + (f" — {g['nome']}" if g["nome"] else "")])
        linha = ws.max_row
        for cel in ws[linha]:
            cel.font, cel.fill = Font(bold=True, size=12), obra_fill
        ws.append(["Funcionário", "Função", *dias_cab, "Extras", "Atrasos", "Saldo"])
        for cel in ws[ws.max_row]:
            cel.font, cel.fill = negrito, cab
            cel.alignment = Alignment(horizontal="center")
        for p in g["pessoas"]:
            celulas, fontes = [], []
            for d in s["dias"]:
                ng = p["negativos"].get(d)
                if p["dias"].get(d):
                    celulas.append(hhmm(p["dias"][d])); fontes.append(azul)
                elif p["negadas"].get(d):
                    celulas.append(f"negada {hhmm(p['negadas'][d])}"); fontes.append(cinza)
                elif p["pequenas"].get(d):
                    celulas.append(f"({hhmm(p['pequenas'][d])})"); fontes.append(cinza)
                elif ng:
                    celulas.append("falta" if ng["tipo"] == "INTEGRAL" else
                                   (f"(-{hhmm(ng['minutos'])})" if ng.get("desconsiderada") else f"-{hhmm(ng['minutos'])}"))
                    fontes.append(cinza if ng.get("desconsiderada") else vermelho)
                elif p["trabalhou"].get(d):
                    celulas.append("✓"); fontes.append(trabalhou)
                else:
                    celulas.append(""); fontes.append(None)
            ws.append([p["nome"] + (" (banco)" if p.get("banco") else ""), p["funcao"], *celulas,
                       hhmm(p["total"]), f"-{hhmm(p['atrasos'])}" if p["atrasos"] else "",
                       (("-" if p["saldo"] < 0 else "") + hhmm(abs(p["saldo"]))) if f.get("abater") else ""])
            linha = ws.max_row
            for i, fonte in enumerate(fontes):
                if fonte is not None:
                    ws.cell(row=linha, column=3 + i).font = fonte
        ws.append(["Total da obra", "", *[hhmm(g["por_dia"][d]) if g["por_dia"].get(d) else "" for d in s["dias"]],
                   hhmm(g["total"]), f"-{hhmm(g['atrasos'])}" if g["atrasos"] else "",
                   hhmm(sum(p["saldo"] for p in g["pessoas"])) if f.get("abater") else ""])
        for cel in ws[ws.max_row]:
            cel.font = negrito
    if not s["obras"]:
        ws.append(["Nenhuma hora extra neste período."])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18
    for i in range(len(s["dias"]) + 1):
        ws.column_dimensions[get_column_letter(3 + i)].width = 6.5
    ws.column_dimensions[get_column_letter(3 + len(s["dias"]))].width = 8
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def nome_do_arquivo(f: dict, qual: str) -> str:
    comp = f["competencia"].replace("-", "")
    return {"geral": f"fechamento_folha_{comp}.xlsx", "ps_faltas": f"faltas_{comp}.ps",
            "ps_extras": f"horas_extras_{comp}.ps", "zip": f"fechamento_folha_{comp}.zip",
            "supervisao": f"horas_extras_para_supervisao_{comp}.xlsx"}.get(
        qual, f"fechamento_{comp}_{qual}.xlsx")


def arquivo(f: dict, qual: str) -> tuple[bytes, str, str]:
    """(conteúdo, nome, tipo) — o geral, um dos .ps, uma parte em Excel, ou os três num .zip."""
    xlsx = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if qual == "geral":
        return excel(f), nome_do_arquivo(f, qual), xlsx
    if qual == "supervisao":
        return excel_supervisao(f), nome_do_arquivo(f, qual), xlsx
    if qual in ("ps_faltas", "ps_extras"):
        return arquivo_ps(f[qual]), nome_do_arquivo(f, qual), "text/plain"
    if qual == "zip":
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for q in ("geral", "ps_faltas", "ps_extras"):
                conteudo, nome, _ = arquivo(f, q)
                z.writestr(nome, conteudo)
        return buf.getvalue(), nome_do_arquivo(f, qual), "application/zip"
    if qual in PARTES:
        return excel(f, so=qual), nome_do_arquivo(f, qual), xlsx
    raise ErroDeValidacao("arquivo desconhecido", campo="arquivo")
