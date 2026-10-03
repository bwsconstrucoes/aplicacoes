# -*- coding: utf-8 -*-
"""
Leitura das planilhas de importação (xlsx/csv) e as regras de cada importador.

Os cabeçalhos são reconhecidos sem acento, sem caixa e sem espaço: "Centro de
Custo", "centro_custo" e "CENTRO CUSTO" são a mesma coluna. Linha com erro é
relatada com o número dela e NÃO impede as boas de entrar. Reexecutar não
duplica: a chave da pessoa é o CPF, a da obra é o código.

As duas funções `importar_*` recebem as linhas já lidas e um `gravar`: com
False (o padrão dos scripts) só dizem o que fariam.
"""
from __future__ import annotations

import csv
import io
import logging
import os
import unicodedata
from typing import Iterable

from sqlalchemy.engine import Connection

from app.apps.erp.core.cadastros.validadores import cpf_valido, somente_digitos

from ..erros import ErroDeValidacao
from . import cadastros, geo

logger = logging.getLogger("ponto.importacao")

# Erros que uma LINHA pode causar sem derrubar a importação inteira.
ErroDoPontoOuValor = (ErroDeValidacao, ValueError)

SINONIMOS = {
    "cpf": {"cpf", "documento", "cpf_do_colaborador"},
    "nome": {"nome", "colaborador", "nome_do_colaborador", "funcionario", "nome_da_obra"},
    "obra": {"obra", "codigo_da_obra", "obra_principal", "cod_obra", "local"},
    "centro_custo": {"centro_custo", "centro_de_custo", "cc", "centrocusto"},
    "tipo_jornada": {"tipo_jornada", "jornada", "tipo_de_jornada", "turno"},
    "codigo": {"codigo", "cod", "codigo_obra", "id_obra"},
    "latitude": {"latitude", "lat"},
    "longitude": {"longitude", "lon", "lng", "long"},
    "raio_metros": {"raio_metros", "raio", "raio_m", "cerca_metros"},
}

JORNADAS_POR_APELIDO = {
    "PADRAO": "PADRAO_4", "PADRAO_4": "PADRAO_4", "4": "PADRAO_4", "NORMAL": "PADRAO_4",
    "VIGIA_DIURNO": "VIGIA_DIURNO_2", "VIGIA_DIURNO_2": "VIGIA_DIURNO_2", "DIURNO": "VIGIA_DIURNO_2",
    "VIGIA_NOTURNO": "VIGIA_NOTURNO_2", "VIGIA_NOTURNO_2": "VIGIA_NOTURNO_2", "NOTURNO": "VIGIA_NOTURNO_2",
}


def normalizar_cabecalho(texto) -> str:
    base = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    base = base.strip().lower().replace("-", " ").replace("/", " ")
    return "_".join(base.split())


def mapear_colunas(cabecalhos: Iterable) -> dict[str, int]:
    """Nome canônico → índice da coluna. Coluna desconhecida é ignorada."""
    mapa: dict[str, int] = {}
    for indice, bruto in enumerate(cabecalhos):
        chave = normalizar_cabecalho(bruto)
        for canonico, apelidos in SINONIMOS.items():
            if chave in apelidos and canonico not in mapa:
                mapa[canonico] = indice
                break
    return mapa


def _valor(celula) -> str:
    if celula is None:
        return ""
    if isinstance(celula, float) and celula.is_integer():
        return str(int(celula))
    return str(celula).strip()


def ler_tabela(caminho: str) -> list[dict]:
    """Lê .xlsx (primeira aba) ou .csv (separador ; ou , detectado). Devolve
    uma lista de dicionários com os nomes canônicos e o número da linha."""
    extensao = os.path.splitext(caminho)[1].lower()
    if extensao in (".xlsx", ".xlsm"):
        linhas_brutas = _ler_xlsx(caminho)
    elif extensao in (".csv", ".txt"):
        with open(caminho, "r", encoding="utf-8-sig", newline="") as f:
            linhas_brutas = ler_csv(f.read())
    else:
        raise ErroDeValidacao(f"formato não suportado: {extensao or '(sem extensão)'}; "
                              "use .xlsx ou .csv")
    return montar_registros(linhas_brutas)


def _ler_xlsx(caminho: str) -> list[list]:
    from openpyxl import load_workbook
    livro = load_workbook(caminho, read_only=True, data_only=True)
    try:
        aba = livro.worksheets[0]
        return [list(linha) for linha in aba.iter_rows(values_only=True)]
    finally:
        livro.close()


def ler_csv(conteudo: str) -> list[list]:
    amostra = conteudo[:4096]
    try:
        dialeto = csv.Sniffer().sniff(amostra, delimiters=";,\t")
    except csv.Error:
        dialeto = csv.excel
        dialeto.delimiter = ";" if amostra.count(";") >= amostra.count(",") else ","
    return [linha for linha in csv.reader(io.StringIO(conteudo), dialeto)]


def montar_registros(linhas_brutas: list[list]) -> list[dict]:
    if not linhas_brutas:
        return []
    mapa = mapear_colunas(linhas_brutas[0])
    if not mapa:
        raise ErroDeValidacao("nenhuma coluna reconhecida no cabeçalho: "
                              + ", ".join(str(c) for c in linhas_brutas[0] if c))
    registros = []
    for numero, linha in enumerate(linhas_brutas[1:], start=2):
        if not linha or all(_valor(c) == "" for c in linha):
            continue
        registro = {"linha": numero}
        for canonico, indice in mapa.items():
            registro[canonico] = _valor(linha[indice]) if indice < len(linha) else ""
        registros.append(registro)
    return registros


def jornada_de(texto: str) -> str:
    chave = normalizar_cabecalho(texto).upper()
    if not chave:
        return "PADRAO_4"
    if chave in JORNADAS_POR_APELIDO:
        return JORNADAS_POR_APELIDO[chave]
    raise ErroDeValidacao(f"tipo de jornada desconhecido: {texto!r}", campo="tipo_jornada")


# ---------------------------------------------------------------------------
# Colaboradores
# ---------------------------------------------------------------------------
def importar_colaboradores(conn: Connection, registros: list[dict], *, gravar: bool) -> dict:
    """Pessoa que não existe no ERP é CRIADA lá (nome, CPF, obra); a jornada e o
    centro de custo vão para `ponto.colaborador_config`. Pessoa que já existe
    não tem nada do ERP alterado — se a obra da planilha divergir, vira aviso."""
    relatorio = {"criados": [], "atualizados": [], "erros": [], "avisos": [], "gravou": gravar}
    for r in registros:
        try:
            cpf = somente_digitos(r.get("cpf", ""))
            if not cpf_valido(cpf):
                raise ErroDeValidacao(f"CPF inválido: {r.get('cpf')!r}")
            nome = (r.get("nome") or "").strip()
            jornada = jornada_de(r.get("tipo_jornada", ""))
            centro_custo = (r.get("centro_custo") or "").strip() or None
            obra = cadastros.resolver_obra(conn, r.get("obra")) if r.get("obra") else None
            if r.get("obra") and not obra:
                raise ErroDeValidacao(f"obra não cadastrada no ERP: {r.get('obra')!r}")
            pessoa = cadastros.colaborador_por_cpf(conn, cpf)
            if pessoa:
                if obra and pessoa["obra_id"] != obra["id"]:
                    relatorio["avisos"].append(
                        f"linha {r['linha']}: {pessoa['nome']} está na obra "
                        f"{pessoa.get('obra_codigo') or '(sem obra)'} no ERP e a planilha diz "
                        f"{obra['codigo']} — o ERP manda; corrija lá se for o caso")
                if gravar:
                    cadastros.gravar_config_colaborador(
                        conn, pessoa["id"], tipo_jornada=jornada, centro_custo=centro_custo)
                relatorio["atualizados"].append(f"linha {r['linha']}: {pessoa['nome']} ({jornada})")
            else:
                if not nome:
                    raise ErroDeValidacao("pessoa nova sem nome")
                if gravar:
                    novo_id = cadastros.criar_colaborador_no_erp(
                        conn, nome=nome, cpf=cpf, obra_id=(obra["id"] if obra else None))
                    cadastros.gravar_config_colaborador(
                        conn, novo_id, tipo_jornada=jornada, centro_custo=centro_custo)
                relatorio["criados"].append(
                    f"linha {r['linha']}: {nome} ({jornada}, obra {obra['codigo'] if obra else '-'})")
        except ErroDoPontoOuValor as e:
            relatorio["erros"].append(f"linha {r['linha']}: {e}")
    logger.info("Ponto: importação de colaboradores — %d criados, %d atualizados, %d erros (%s)",
                len(relatorio["criados"]), len(relatorio["atualizados"]), len(relatorio["erros"]),
                "gravado" if gravar else "simulação")
    return relatorio


# ---------------------------------------------------------------------------
# Obras
# ---------------------------------------------------------------------------
def importar_obras(conn: Connection, registros: list[dict], *, gravar: bool,
                   sobrescrever_coordenadas: bool = False) -> dict:
    """Obra que não existe no ERP é criada (código e nome). Coordenadas vazias
    ficam vazias. Coordenada que já existe no ERP só é trocada com
    `sobrescrever_coordenadas` — senão vira aviso."""
    relatorio = {"criadas": [], "atualizadas": [], "erros": [], "avisos": [], "gravou": gravar}
    for r in registros:
        try:
            codigo = (r.get("codigo") or r.get("obra") or "").strip()
            nome = (r.get("nome") or "").strip()
            if not codigo:
                raise ErroDeValidacao("obra sem código")
            lat, lon = (r.get("latitude") or "").strip(), (r.get("longitude") or "").strip()
            tem_coordenada = bool(lat or lon)
            if tem_coordenada and not geo.coordenada_valida(lat.replace(",", "."), lon.replace(",", ".")):
                raise ErroDeValidacao(f"latitude/longitude inválidas: {lat!r}, {lon!r}")
            raio_txt = (r.get("raio_metros") or "").strip()
            raio = int(float(raio_txt.replace(",", "."))) if raio_txt else None
            centro_custo = (r.get("centro_custo") or "").strip() or None
            obra = cadastros.obra_por_codigo(conn, codigo)
            if obra:
                if tem_coordenada:
                    if obra["latitude"] is None or sobrescrever_coordenadas:
                        if gravar:
                            cadastros.gravar_coordenadas_da_obra(
                                conn, obra["id"], geo.decimal_ou_none(lat), geo.decimal_ou_none(lon))
                    elif (float(obra["latitude"]) != float(lat.replace(",", "."))
                          or float(obra["longitude"]) != float(lon.replace(",", "."))):
                        relatorio["avisos"].append(
                            f"linha {r['linha']}: obra {codigo} já tem coordenada no ERP "
                            "diferente da planilha — mantida a do ERP (use "
                            "--sobrescrever-coordenadas para trocar)")
                if gravar:
                    cadastros.gravar_config_obra(conn, obra["id"], raio_metros=raio,
                                                 centro_custo=centro_custo)
                relatorio["atualizadas"].append(f"linha {r['linha']}: {codigo} — {obra['nome']}")
            else:
                if not nome:
                    raise ErroDeValidacao(f"obra nova {codigo} sem nome")
                if gravar:
                    novo_id = cadastros.criar_obra_no_erp(conn, codigo=codigo, nome=nome)
                    if tem_coordenada:
                        cadastros.gravar_coordenadas_da_obra(
                            conn, novo_id, geo.decimal_ou_none(lat), geo.decimal_ou_none(lon))
                    cadastros.gravar_config_obra(conn, novo_id, raio_metros=raio,
                                                 centro_custo=centro_custo)
                relatorio["criadas"].append(
                    f"linha {r['linha']}: {codigo} — {nome}"
                    + ("" if tem_coordenada else " (sem coordenada)"))
        except ErroDoPontoOuValor as e:
            relatorio["erros"].append(f"linha {r['linha']}: {e}")
    logger.info("Ponto: importação de obras — %d criadas, %d atualizadas, %d erros (%s)",
                len(relatorio["criadas"]), len(relatorio["atualizadas"]), len(relatorio["erros"]),
                "gravado" if gravar else "simulação")
    return relatorio



def resumo(relatorio: dict) -> str:
    """Texto para o terminal."""
    partes = []
    for chave in ("criados", "criadas", "atualizados", "atualizadas", "avisos", "erros"):
        itens = relatorio.get(chave)
        if itens is None:
            continue
        partes.append(f"\n== {chave.upper()} ({len(itens)}) ==")
        partes.extend(f"  {i}" for i in itens)
    partes.append("\n" + ("GRAVADO." if relatorio.get("gravou")
                          else "SIMULAÇÃO — nada foi gravado. Use --gravar para gravar."))
    return "\n".join(partes)
