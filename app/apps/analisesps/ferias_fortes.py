# -*- coding: utf-8 -*-
"""
A "LISTAGEM DE FÉRIAS" DO FORTES — importar as férias do mês (05/10/2026).

O dono: *"aí nós temos a lista de férias das pessoas (…) tem um determinado
local que tem o gozo das férias (…) para a gente entender se aquela pessoa terá
direito a auxílio alimentação e auxílio transporte, porque se ela tiver de férias
é para ser descontado proporcionalmente (…) eu vou pedir do mês anterior e do mês
atual (…) jogar dois arquivos, ele faz a leitura, compreende também se aquela
informação já foi cadastrada ou não (…) o que já está cadastrado elimina (…) e
até verificar se tem alguma mudança, alguma diferença nas férias de alguém."*

O ARQUIVO (Fortes Pessoal, `.xls` antigo, uma aba, várias páginas):

    Listagem de Férias                              Pag.: 1
    Empresa: BWS CONSTRUCOES LTDA - CNPJ …
    Iniciadas entre 01/09/2026 a 30/09/2026
    Código | Empregado | Evento | | Referência | | Provento | Desconto
    000705 | NOME
    Cargo: AUX. ADMINISTRATIVO
           |           | 110 Remuneração de Férias | | 30 dia(s) | | 1794.0 |
           …
                                      FGTS: 191,36 | | Líquido a receber: | 1948.85
    Período Aquisitivo: 01/10/2024 a 30/09/2025 | | | Gozo: 01/09/2026 a
    30/09/2026 | Retorno: 01/10/2026 | | Abono: 0 dia
    …
    Total Geral (9 empregados)

O que vale para o desconto é o GOZO (os dias de descanso; o abono pecuniário —
dias vendidos — já fica fora dele). Os períodos vão para a mesma tabela de
férias que a tela "Feriados e férias" e o cálculo do auxílio já usam.

A pessoa é achada pelo CÓDIGO DO FORTES (coluna BU da ficha); sem ele, pelo
nome, só se houver UM igual no cadastro. Quem não for achado não entra — e a
tela diz quem é.
"""
from __future__ import annotations

import datetime as dt
import json
import logging
import re
import unicodedata
from decimal import Decimal

logger = logging.getLogger("analisesps.ferias_fortes")

MAXIMO_DO_ARQUIVO = 10 * 1024 * 1024
MAXIMO_DE_ARQUIVOS = 6

_DATA = r"(\d{2}/\d{2}/\d{4})"
_RE_INICIADAS = re.compile(r"iniciadas\s+entre\s+" + _DATA + r"\s+a\s+" + _DATA, re.I)
_RE_AQUISITIVO = re.compile(r"per[ií]odo\s+aquisitivo:\s*" + _DATA + r"\s+a\s+" + _DATA, re.I)
_RE_GOZO = re.compile(r"gozo:\s*" + _DATA + r"\s+a\s+" + _DATA, re.I)
_RE_RETORNO = re.compile(r"retorno:\s*" + _DATA, re.I)
_RE_ABONO = re.compile(r"abono:\s*(\d+)", re.I)
_RE_EVENTO = re.compile(r"^(\d{1,4})\s+(.+)$")


class ErroDasFerias(RuntimeError):
    """A frase vai inteira para a tela."""


def _texto(v) -> str:
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v if v is not None else "").split())


def _sem_acento(texto) -> str:
    cru = unicodedata.normalize("NFKD", _texto(texto))
    return "".join(c for c in cru if not unicodedata.combining(c)).upper()


def _data(texto) -> dt.date | None:
    try:
        d, m, a = str(texto).split("/")
        return dt.date(int(a), int(m), int(d))
    except (ValueError, TypeError):
        return None


def _valor(v) -> Decimal | None:
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return Decimal(str(v)).quantize(Decimal("0.01"))
    from . import formatos
    return formatos.para_numero(v)


# ---------------------------------------------------------------------------
# LER
# ---------------------------------------------------------------------------
def _linhas_do_arquivo(conteudo: bytes, nome: str = "") -> list:
    if not conteudo:
        raise ErroDasFerias(f"{nome or 'arquivo'}: arquivo vazio.")
    if len(conteudo) > MAXIMO_DO_ARQUIVO:
        raise ErroDasFerias(f"{nome or 'arquivo'}: acima de 10 MB — não parece a "
                            "Listagem de Férias.")
    if conteudo[:2] == b"PK":            # .xlsx
        from openpyxl import load_workbook
        import io
        livro = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        aba = livro.worksheets[0]
        return [list(l) for l in aba.iter_rows(values_only=True)]
    try:
        import xlrd
        livro = xlrd.open_workbook(file_contents=conteudo)
        aba = livro.sheet_by_index(0)
        return [[aba.cell_value(r, c) for c in range(aba.ncols)] for r in range(aba.nrows)]
    except Exception as e:  # noqa: BLE001 — a tela precisa da frase
        raise ErroDasFerias(f"{nome or 'arquivo'}: não foi possível abrir como Excel "
                            f"(.xls): {e}") from e


def interpretar(linhas, nome: str = "") -> dict:
    """`{periodo: (início, fim) | None, registros: [...]}` — um registro por
    colaborador: codigo, nome, cargo, aquisitivo, inicio, fim (do gozo),
    retorno, abono_dias, liquido, eventos."""
    periodo = None
    registros = []
    atual = None
    titulo_ok = False
    for bruta in linhas:
        cel = [_texto(c) for c in (list(bruta) + [""] * 8)[:8]]
        juntos = " ".join(c for c in cel if c)
        if not juntos:
            continue
        if "listagem de f" in juntos.lower():
            titulo_ok = True
        if not periodo:
            m = _RE_INICIADAS.search(juntos)
            if m:
                periodo = (_data(m.group(1)), _data(m.group(2)))
                continue
        if cel[0].lower().startswith("total geral"):
            break
        # O começo de um colaborador: o código (só dígitos) e o nome.
        if re.fullmatch(r"\d{3,6}", cel[0]) and cel[1] and not cel[2]:
            atual = {"codigo": cel[0].zfill(6), "nome": cel[1], "cargo": "",
                     "eventos": [], "liquido": None, "aquisitivo": "",
                     "inicio": None, "fim": None, "retorno": None, "abono_dias": 0}
            registros.append(atual)
            continue
        if atual is None:
            continue
        if cel[0].lower().startswith("cargo:"):
            atual["cargo"] = cel[0].split(":", 1)[1].strip()
            continue
        m = _RE_EVENTO.match(cel[2])
        if m:
            atual["eventos"].append({
                "codigo": m.group(1), "descricao": m.group(2), "referencia": cel[4],
                "provento": str(_valor(bruta[6]) or "") if len(bruta) > 6 else "",
                "desconto": str(_valor(bruta[7]) or "") if len(bruta) > 7 else ""})
            continue
        if "quido a receber" in juntos.lower():
            atual["liquido"] = _valor(bruta[7]) if len(bruta) > 7 else None
            continue
        m = _RE_GOZO.search(juntos)
        if m:
            atual["inicio"], atual["fim"] = _data(m.group(1)), _data(m.group(2))
            a = _RE_AQUISITIVO.search(juntos)
            if a:
                atual["aquisitivo"] = f"{a.group(1)} a {a.group(2)}"
            r = _RE_RETORNO.search(juntos)
            atual["retorno"] = _data(r.group(1)) if r else None
            b = _RE_ABONO.search(juntos)
            atual["abono_dias"] = int(b.group(1)) if b else 0
            atual = None
    if not titulo_ok:
        raise ErroDasFerias(f"{nome or 'arquivo'}: não é a \"Listagem de Férias\" do "
                            "Fortes (o título não foi encontrado).")
    validos = [r for r in registros if r["inicio"] and r["fim"]]
    sem_gozo = [r["nome"] for r in registros if not (r["inicio"] and r["fim"])]
    return {"periodo": periodo, "registros": validos, "sem_gozo": sem_gozo,
            "arquivo": nome}


def ler(conteudo: bytes, nome: str = "") -> dict:
    return interpretar(_linhas_do_arquivo(conteudo, nome), nome)


# ---------------------------------------------------------------------------
# COMPARAR COM O QUE JÁ ESTÁ CADASTRADO
# ---------------------------------------------------------------------------
def _cadastro() -> tuple:
    """(`{código do Fortes: (cpf, nome)}`, `{nome sem acento: [(cpf, nome)]}`)."""
    from . import colaboradores
    from .db import consultar
    por_codigo, por_nome = {}, {}
    tem_codigo = colaboradores.tem_id_fortes()
    for linha in consultar(
            "SELECT cpf, nome" + (", id_fortes" if tem_codigo else "")
            + " FROM analisesps.colaborador"):
        cpf, nome = linha[0], linha[1] or ""
        if tem_codigo and linha[2]:
            por_codigo[colaboradores.normalizar_id_fortes(linha[2])] = (cpf, nome)
        por_nome.setdefault(_sem_acento(nome), []).append((cpf, nome))
    return por_codigo, por_nome


def _periodos_cadastrados(cpfs) -> dict:
    """`{cpf: [{id, inicio, fim, origem}]}`."""
    from .db import consultar, tem_coluna
    cpfs = sorted(set(cpfs))
    if not cpfs:
        return {}
    com_origem = tem_coluna("ferias", "origem")
    marcas = ", ".join("?" for _ in cpfs)
    saida: dict = {}
    for l in consultar(
            "SELECT id, cpf, inicio, fim" + (", origem" if com_origem else "")
            + f" FROM analisesps.ferias WHERE cpf IN ({marcas}) ORDER BY inicio",
            tuple(cpfs)):
        saida.setdefault(l[1], []).append(
            {"id": l[0], "inicio": l[2], "fim": l[3],
             "origem": (l[4] if com_origem else "") or ""})
    return saida


def analisar(arquivos) -> dict:
    """Lê os arquivos (`[(nome, bytes)]`) e diz o que cada período é — NÃO
    GRAVA NADA:

      novas       — o período não existe: entra;
      iguais      — já cadastrado com as mesmas datas: fica como está;
      alteradas   — a pessoa tem outro período que se CRUZA com este (datas
                    mudaram): o antigo é substituído pelo do arquivo;
      sem_cadastro— código e nome não acham ninguém no cadastro: não entra;
      sumiram     — importadas antes, com início dentro do mês de um destes
                    arquivos, e que o arquivo não traz mais: só aviso (podem ter
                    sido canceladas), nada é apagado.
    """
    from .folha_rateio import cpf_bonito
    if not arquivos:
        raise ErroDasFerias("nenhum arquivo enviado.")
    if len(arquivos) > MAXIMO_DE_ARQUIVOS:
        raise ErroDasFerias(f"no máximo {MAXIMO_DE_ARQUIVOS} arquivos por vez.")
    lidos = [ler(conteudo, nome) for nome, conteudo in arquivos]

    # O mesmo período em dois arquivos (o mês anterior e o atual se repetem)
    # conta uma vez só.
    vistos: dict = {}
    for lido in lidos:
        rotulo = (f"Fortes — férias iniciadas de {lido['periodo'][0]:%d/%m/%Y} a "
                  f"{lido['periodo'][1]:%d/%m/%Y}" if lido["periodo"] and all(lido["periodo"])
                  else f"Fortes — {lido['arquivo']}")
        for r in lido["registros"]:
            chave = (r["codigo"], r["inicio"], r["fim"])
            vistos.setdefault(chave, dict(
                r, origem=rotulo, arquivo=lido["arquivo"],
                mes_do_arquivo=(lido["periodo"][0] if lido["periodo"] and lido["periodo"][0]
                                else dt.date.min)))

    por_codigo, por_nome = _cadastro()
    registros = list(vistos.values())
    for r in registros:
        achado = por_codigo.get(r["codigo"])
        if not achado:
            mesmos = por_nome.get(_sem_acento(r["nome"])) or []
            achado = mesmos[0] if len(mesmos) == 1 else None
            r["casou_por"] = "nome" if achado else ""
        else:
            r["casou_por"] = "código"
        r["cpf"] = achado[0] if achado else ""
        r["cpf_bonito"] = cpf_bonito(r["cpf"]) if r["cpf"] else ""
        r["dias"] = (r["fim"] - r["inicio"]).days + 1

    cadastrados = _periodos_cadastrados(r["cpf"] for r in registros if r["cpf"])
    saida = {"novas": [], "iguais": [], "alteradas": [], "sem_cadastro": [],
             "sumiram": [], "conflitos": [],
             "arquivos": [{"nome": l["arquivo"], "periodo": l["periodo"],
                           "quantos": len(l["registros"]), "sem_gozo": l["sem_gozo"]}
                          for l in lidos]}
    # Dois períodos do arquivo para a mesma pessoa que se cruzam (datas
    # mudaram entre o arquivo do mês anterior e o do atual): vale o do arquivo
    # de mês mais recente, e a tela diz.
    por_pessoa: dict = {}
    for r in sorted(registros, key=lambda x: (x["codigo"], x["inicio"])):
        por_pessoa.setdefault(r["codigo"], []).append(r)
    descartados = set()
    for lista in por_pessoa.values():
        for i, a in enumerate(lista):
            for b in lista[i + 1:]:
                if a["inicio"] <= b["fim"] and b["inicio"] <= a["fim"]:
                    velho = a if a["mes_do_arquivo"] < b["mes_do_arquivo"] else b
                    novo = b if velho is a else a
                    descartados.add(id(velho))
                    saida["conflitos"].append(
                        f"{a['nome']}: {velho['inicio']:%d/%m/%Y} a {velho['fim']:%d/%m/%Y} "
                        f"({velho['arquivo']}) × {novo['inicio']:%d/%m/%Y} a "
                        f"{novo['fim']:%d/%m/%Y} ({novo['arquivo']}) — vale o segundo.")
    for r in registros:
        if id(r) in descartados:
            continue
        if not r["cpf"]:
            saida["sem_cadastro"].append(r)
            continue
        dele = cadastrados.get(r["cpf"]) or []
        if any(p["inicio"] == r["inicio"] and p["fim"] == r["fim"] for p in dele):
            saida["iguais"].append(r)
            continue
        cruzam = [p for p in dele if p["inicio"] <= r["fim"] and r["inicio"] <= p["fim"]]
        if cruzam:
            r["antes"] = cruzam
            saida["alteradas"].append(r)
        else:
            saida["novas"].append(r)

    # O que foi importado antes e o arquivo não traz mais.
    janelas = [l["periodo"] for l in lidos if l["periodo"] and all(l["periodo"])]
    if janelas:
        no_arquivo = {(r["cpf"], r["inicio"], r["fim"]) for r in registros if r["cpf"]}
        for cpf, periodos, nome in _importadas_nas_janelas(janelas):
            for p in periodos:
                if (cpf, p["inicio"], p["fim"]) not in no_arquivo and not any(
                        r["cpf"] == cpf and r["inicio"] <= p["fim"] and p["inicio"] <= r["fim"]
                        for r in registros):
                    saida["sumiram"].append({"cpf": cpf, "cpf_bonito": cpf_bonito(cpf),
                                             "nome": nome, **p})
    for chave in ("novas", "iguais", "alteradas", "sem_cadastro"):
        saida[chave].sort(key=lambda r: (r["nome"], r["inicio"]))
    saida["a_gravar"] = len(saida["novas"]) + len(saida["alteradas"])
    return saida


def _importadas_nas_janelas(janelas) -> list:
    """`[(cpf, [períodos], nome)]` importados do Fortes com início dentro das
    janelas ("iniciadas entre") dos arquivos."""
    from .db import consultar, tem_coluna
    if not tem_coluna("ferias", "origem"):
        return []
    saida: dict = {}
    for ini, fim in janelas:
        for l in consultar(
                "SELECT id, cpf, nome, inicio, fim, origem FROM analisesps.ferias "
                " WHERE origem LIKE 'Fortes%' AND inicio BETWEEN ? AND ?", (ini, fim)):
            item = saida.setdefault(l[1], ([], l[2]))
            if not any(p["id"] == l[0] for p in item[0]):
                item[0].append({"id": l[0], "inicio": l[3], "fim": l[4], "origem": l[5]})
    return [(cpf, periodos, nome) for cpf, (periodos, nome) in saida.items()]


# ---------------------------------------------------------------------------
# GRAVAR
# ---------------------------------------------------------------------------
def importar(arquivos, quem: str = "") -> dict:
    """Refaz a análise e grava as novas e as alteradas (o período antigo que se
    cruza é substituído), numa transação só. Devolve a análise e o que gravou."""
    from . import folha_calendario
    from .db import conexao, tem_coluna
    if not folha_calendario._pronto():
        raise ErroDasFerias('tabela de férias não encontrada. Clique em "Aplicar '
                            'atualizações do banco" em Configurações.')
    analise = analisar(arquivos)
    completas = tem_coluna("ferias", "origem")
    gravadas = substituidas = 0
    with conexao() as conn:
        for r in analise["novas"] + analise["alteradas"]:
            for antigo in r.get("antes") or []:
                conn.execute("DELETE FROM analisesps.ferias WHERE id = ?", (antigo["id"],))
                substituidas += 1
            observacao = (f"Fortes: aquisitivo {r['aquisitivo']}"
                          + (f", abono {r['abono_dias']} dia(s)" if r["abono_dias"] else ""))
            if completas:
                conn.execute(
                    "INSERT INTO analisesps.ferias (cpf, nome, inicio, fim, observacao, "
                    "  criado_por, origem, codigo_fortes, aquisitivo, retorno, abono_dias, "
                    "  liquido, eventos) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (r["cpf"], r["nome"][:160], r["inicio"], r["fim"], observacao[:300],
                     str(quem or "")[:120], r["origem"][:200], r["codigo"],
                     r["aquisitivo"], r["retorno"], int(r["abono_dias"] or 0),
                     r["liquido"], json.dumps(r["eventos"], ensure_ascii=False)))
            else:
                conn.execute(
                    "INSERT INTO analisesps.ferias (cpf, nome, inicio, fim, observacao, "
                    "  criado_por) VALUES (?,?,?,?,?,?)",
                    (r["cpf"], r["nome"][:160], r["inicio"], r["fim"],
                     (observacao + " — " + r["origem"])[:300], str(quem or "")[:120]))
            gravadas += 1
        conn.commit()
    logger.info("Férias: %d período(s) importado(s) do Fortes (%d substituído(s)) por %s.",
                gravadas, substituidas, quem or "(sem nome)")
    return {**analise, "gravadas": gravadas, "substituidas": substituidas}


def para_tela(analise: dict) -> dict:
    """A análise com datas em texto, para o JSON da tela."""
    def r_(r):
        return {"nome": r["nome"], "codigo": r["codigo"], "cpf": r.get("cpf_bonito") or "",
                "cargo": r.get("cargo") or "", "casou_por": r.get("casou_por") or "",
                "inicio": f"{r['inicio']:%d/%m/%Y}", "fim": f"{r['fim']:%d/%m/%Y}",
                "dias": r.get("dias") or 0, "abono": r.get("abono_dias") or 0,
                "aquisitivo": r.get("aquisitivo") or "",
                "liquido": str(r["liquido"]) if r.get("liquido") is not None else "",
                "arquivo": r.get("arquivo") or "",
                "antes": [f"{p['inicio']:%d/%m/%Y} a {p['fim']:%d/%m/%Y}"
                          for p in r.get("antes") or []]}
    saida = {k: [r_(r) for r in analise[k]]
             for k in ("novas", "iguais", "alteradas", "sem_cadastro")}
    saida["sumiram"] = [{"nome": p["nome"], "cpf": p["cpf_bonito"],
                         "inicio": f"{p['inicio']:%d/%m/%Y}", "fim": f"{p['fim']:%d/%m/%Y}"}
                        for p in analise["sumiram"]]
    saida["conflitos"] = analise["conflitos"]
    saida["arquivos"] = [{"nome": a["nome"], "quantos": a["quantos"],
                          "sem_gozo": a["sem_gozo"],
                          "periodo": (f"{a['periodo'][0]:%d/%m/%Y} a {a['periodo'][1]:%d/%m/%Y}"
                                      if a["periodo"] and all(a["periodo"]) else "")}
                         for a in analise["arquivos"]]
    saida["a_gravar"] = analise["a_gravar"]
    for k in ("gravadas", "substituidas"):
        if k in analise:
            saida[k] = analise[k]
    return saida
