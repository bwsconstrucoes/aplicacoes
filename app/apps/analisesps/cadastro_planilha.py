# -*- coding: utf-8 -*-
"""
AS PLANILHAS DE CADASTRO DO BEEVALE E DA SOMAPAY — 03/10/2026.

O dono: *"às vezes tem pessoas novas que não têm cadastro no BeeVale ainda, e a
gente gera para poder fazer o cadastro dessas pessoas que estão sendo pagas (…)
planilha de cadastro BeeVale, planilha do SomaPay (…) a gente selecionar algumas
pessoas. Às vezes pode acontecer de eu querer pagar um arquivo, aí, opa, dois não
estão cadastrados ainda. Aí vou lá, cadastro, gera a planilha dessas duas pessoas,
cadastra, e processa novamente."*

- **BeeVale** (`modelo_import2.xlsx` que ele mandou): Nome completo, Nome
  Impresso no Cartão, CPF, Email, Data de nascimento, DDI, Celular. O e-mail é o
  mesmo do arquivo de pagamento (`<cpf>@bwsconstrucoes.com.br`) — é a chave da
  pessoa no portal.
- **SomaPay** ("modelo de registro de funcionário" do portal): o modelo deles,
  preenchido a partir da linha 12, como o arquivo de pagamento (o portal recusa
  arquivo montado do zero — ver `folha_geracao.somapay_xlsx`). Pede RG, órgão,
  nome da mãe, endereço…, que vêm da ficha (`colaborador.documentos`, migração
  047). O que faltar fica em branco e é DITO na resposta.
"""
from __future__ import annotations

import datetime as dt
import io
import re
import zipfile
from pathlib import Path

from . import colaboradores
from .folha_rateio import so_digitos

BEEVALE = "beevale"
SOMAPAY = "somapay"
DESTINOS = (BEEVALE, SOMAPAY)

COLUNAS_BEEVALE = ["Nome completo", "Nome Impresso no Cartão", "CPF", "Email",
                   "Data de nascimento", "DDI", "Celular"]
DOMINIO = "@bwsconstrucoes.com.br"
DDI = "55"
# O nome no cartão: o portal limita o tamanho; primeiro e último nome cabem.
TETO_DO_NOME_NO_CARTAO = 26

MODELO_SOMAPAY = Path(__file__).with_name("modelos") / "somapay_cadastro_modelo.xlsx"
PRIMEIRA_LINHA_SOMAPAY = 12
# A ordem das colunas do modelo (A…W), com o que vai em cada uma.
COLUNAS_SOMAPAY = ["cpf", "nome", "nascimento", "matricula", "rg", "rg_emissao",
                   "rg_orgao", "rg_uf", "nome_mae", "sexo", "admissao", "cep",
                   "logradouro", "bairro", "numero", "complemento", "cidade",
                   "estado", "ddd", "celular", "centro_de_custo", "setor",
                   "tipo_contrato"]
OBRIGATORIOS_SOMAPAY = {"cpf": "CPF", "nome": "nome", "nascimento": "nascimento",
                        "matricula": "matrícula", "rg": "RG", "rg_emissao":
                        "emissão do RG", "rg_orgao": "órgão emissor", "rg_uf":
                        "UF do RG", "nome_mae": "nome da mãe", "admissao":
                        "admissão", "ddd": "DDD", "celular": "celular",
                        "tipo_contrato": "tipo de contrato"}
# A lista do modelo (validação da coluna W).
TIPOS_DE_CONTRATO = ("CLT (tempo determinado)", "CLT (tempo Indeterminado)", "PJ",
                     "Autônomo (RPA)", "Jovem aprendiz", "Estágio", "Temporário",
                     "Pro labore")


class ErroDoCadastro(RuntimeError):
    """A frase vai inteira para a tela."""


# ---------------------------------------------------------------------------
# Feitio de dado
# ---------------------------------------------------------------------------
def _sem_acento(texto) -> str:
    import unicodedata
    cru = unicodedata.normalize("NFKD", str(texto or ""))
    return "".join(c for c in cru if not unicodedata.combining(c)).lower()


def nome_no_cartao(nome: str) -> str:
    """Primeiro e último nome, em maiúsculas, no tamanho do cartão."""
    partes = [p for p in str(nome or "").upper().split() if p]
    if not partes:
        return ""
    curto = " ".join([partes[0], partes[-1]]) if len(partes) > 1 else partes[0]
    return curto[:TETO_DO_NOME_NO_CARTAO]


def _data(valor) -> str:
    if isinstance(valor, (dt.date, dt.datetime)):
        return valor.strftime("%d/%m/%Y")
    texto = str(valor or "").strip()
    achado = re.match(r"^(\d{4})-(\d{2})-(\d{2})", texto)
    if achado:
        return f"{achado.group(3)}/{achado.group(2)}/{achado.group(1)}"
    return texto


def telefone(celular) -> tuple:
    """(DDD, número) do celular da ficha — o DDD de 2 dígitos e o resto."""
    d = so_digitos(celular)
    if d.startswith("55") and len(d) in (12, 13):
        d = d[2:]
    if len(d) in (10, 11):
        return d[:2], d[2:]
    return "", d


def celular_formatado(celular) -> str:
    ddd, numero = telefone(celular)
    if not ddd:
        return numero
    meio = len(numero) - 4
    return f"({ddd}) {numero[:meio]}-{numero[meio:]}"


def tipo_de_contrato(texto) -> str:
    """O tipo da ficha na lista do modelo da SomaPay."""
    t = _sem_acento(texto)
    if "rpa" in t or "autonom" in t or "prestador" in t:
        return "Autônomo (RPA)"
    if "pro labore" in t or "pro-labore" in t or "prolabore" in t:
        return "Pro labore"
    if "estag" in t:
        return "Estágio"
    if "aprendiz" in t:
        return "Jovem aprendiz"
    if "tempor" in t:
        return "Temporário"
    if t.strip() == "pj" or "pessoa juridica" in t:
        return "PJ"
    if "determinado" in t and "indeterminado" not in t:
        return "CLT (tempo determinado)"
    if "clt" in t or "ctps" in t:
        return "CLT (tempo Indeterminado)"
    return ""


def pessoas(cpfs, nomes: dict | None = None) -> list:
    """Os dados de cada CPF pedido, do cadastro. Quem não está no cadastro vem
    só com o CPF e o nome da tela (`nomes`)."""
    nomes = nomes or {}
    fichas = colaboradores.documentos_de(cpfs)
    saida = []
    for cpf in dict.fromkeys(so_digitos(c) for c in cpfs or []):
        if len(cpf) != 11:
            continue
        f = fichas.get(cpf) or {"cpf": cpf, "nome": nomes.get(cpf, ""),
                                "documentos": {}, "sem_cadastro": True}
        saida.append(f)
    return saida


# ---------------------------------------------------------------------------
# BeeVale
# ---------------------------------------------------------------------------
def beevale_xlsx(lista) -> tuple:
    """(bytes, avisos) — a planilha de cadastro do BeeVale."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    planilha = Workbook()
    aba = planilha.active
    aba.title = "Colaboradores"
    aba.append(COLUNAS_BEEVALE)
    avisos = []
    for p in lista:
        nome = " ".join(str(p.get("nome") or "").split()).upper()
        faltam = [r for r, ok in (("cadastro", not p.get("sem_cadastro")),
                                  ("nascimento", p.get("nascimento")),
                                  ("celular", p.get("celular"))) if not ok]
        if faltam:
            avisos.append(f"{nome or p['cpf']}: sem {', '.join(faltam)}")
        aba.append([nome, nome_no_cartao(nome), p["cpf"], f"{p['cpf']}{DOMINIO}",
                    _data(p.get("nascimento")), DDI, celular_formatado(p.get("celular"))])
    for coluna in range(1, len(COLUNAS_BEEVALE) + 1):
        for celula in aba[get_column_letter(coluna)]:
            celula.number_format = "@"
    aba.freeze_panes = "A2"
    memoria = io.BytesIO()
    planilha.save(memoria)
    return memoria.getvalue(), avisos


# ---------------------------------------------------------------------------
# SomaPay — o modelo deles, preenchido
# ---------------------------------------------------------------------------
def _linha_somapay(p: dict) -> dict:
    docs = p.get("documentos") or {}
    ddd, numero = telefone(p.get("celular"))
    sexo = _sem_acento(docs.get("sexo"))[:1].upper()
    return {
        "cpf": p["cpf"], "nome": " ".join(str(p.get("nome") or "").split()).upper(),
        "nascimento": _data(p.get("nascimento")),
        "matricula": so_digitos(p.get("matricula")) or p["cpf"],
        "rg": so_digitos(docs.get("rg")), "rg_emissao": _data(docs.get("rg_emissao")),
        "rg_orgao": _orgao(docs.get("rg_orgao"), docs.get("rg_uf")),
        "rg_uf": str(docs.get("rg_uf") or "").strip().upper()[:2],
        "nome_mae": str(docs.get("nome_mae") or "").strip().upper(),
        "sexo": sexo if sexo in ("M", "F") else "",
        "admissao": _data(p.get("data_admissao")),
        "cep": so_digitos(docs.get("cep")),
        "logradouro": str(docs.get("logradouro") or "").strip(),
        "bairro": str(docs.get("bairro") or "").strip(),
        "numero": str(docs.get("numero") or "").strip(),
        "complemento": str(docs.get("complemento") or "").strip(),
        "cidade": str(docs.get("cidade") or "").strip(),
        "estado": str(docs.get("estado") or "").strip().upper()[:2],
        "ddd": ddd, "celular": numero,
        "centro_de_custo": "", "setor": "",
        "tipo_contrato": tipo_de_contrato(p.get("tipo_contrato")),
    }


def _orgao(orgao, uf) -> str:
    """"SSP-CE" → "SSP": o modelo pede o órgão sem traço, e a UF vai à parte."""
    limpo = re.sub(r"[^A-Za-z]", "", _sem_acento(orgao)).upper()
    uf = str(uf or "").strip().upper()[:2]
    if uf and len(limpo) > len(uf) + 1 and limpo.endswith(uf):
        limpo = limpo[:-len(uf)]
    return limpo


def _letra(n: int) -> str:
    letras = ""
    while n > 0:
        n, resto = divmod(n - 1, 26)
        letras = chr(65 + resto) + letras
    return letras


def somapay_xlsx(lista) -> tuple:
    """(bytes, avisos) — o modelo de cadastro da SomaPay, preenchido."""
    from xml.sax.saxutils import escape

    with zipfile.ZipFile(MODELO_SOMAPAY) as modelo:
        partes = [(info, modelo.read(info.filename)) for info in modelo.infolist()]
    caminho = "xl/worksheets/sheet1.xml"
    folha = next(d for i, d in partes if i.filename == caminho).decode("utf-8")

    avisos = []
    for n, p in enumerate(lista):
        linha = _linha_somapay(p)
        faltam = [rotulo for campo, rotulo in OBRIGATORIOS_SOMAPAY.items()
                  if not linha.get(campo)]
        if p.get("sem_cadastro"):
            avisos.append(f"{linha['nome'] or p['cpf']}: fora do cadastro de colaboradores")
        elif faltam:
            avisos.append(f"{linha['nome'] or p['cpf']}: sem {', '.join(faltam)}")
        r = PRIMEIRA_LINHA_SOMAPAY + n
        for i, campo in enumerate(COLUNAS_SOMAPAY, start=1):
            valor = str(linha.get(campo) or "")
            if not valor:
                continue
            ref = f"{_letra(i)}{r}"
            celula = re.compile(rf'<c r="{ref}" s="(\d+)"/>')
            texto = "".join(c for c in valor if c >= " ")
            folha, trocas = celula.subn(
                lambda m, ref=ref, texto=texto: (
                    f'<c r="{ref}" s="{m.group(1)}" t="inlineStr"><is>'
                    f'<t xml:space="preserve">{escape(texto)}</t></is></c>'),
                folha, count=1)
            if not trocas:
                raise ErroDoCadastro(
                    f"o modelo de cadastro da SomaPay não tem a célula {ref} — "
                    "o arquivo do portal mudou; mande o modelo novo.")
    memoria = io.BytesIO()
    with zipfile.ZipFile(memoria, "w", zipfile.ZIP_DEFLATED) as saida:
        for info, dados in partes:
            saida.writestr(info, folha.encode("utf-8") if info.filename == caminho
                           else dados)
    return memoria.getvalue(), avisos


def gerar(destino: str, cpfs, nomes: dict | None = None) -> tuple:
    """(bytes, nome do arquivo, avisos) da planilha de cadastro."""
    from .horario import agora
    destino = str(destino or "").strip().lower()
    if destino not in DESTINOS:
        raise ErroDoCadastro('escolha a planilha: BeeVale ou SomaPay.')
    lista = pessoas(cpfs, nomes)
    if not lista:
        raise ErroDoCadastro("selecione ao menos um colaborador.")
    if len(lista) > 2000:
        raise ErroDoCadastro("mais de 2.000 colaboradores numa planilha de cadastro.")
    quando = agora().strftime("%Y-%m-%d_%H%M")
    if destino == BEEVALE:
        conteudo, avisos = beevale_xlsx(lista)
        return conteudo, f"Cadastro BeeVale - {len(lista)} colaborador(es) - {quando}.xlsx", avisos
    conteudo, avisos = somapay_xlsx(lista)
    # A mesma extensão do modelo do portal (o conteúdo é do Excel novo).
    return conteudo, f"Cadastro SomaPay - {len(lista)} colaborador(es) - {quando}.xls", avisos
