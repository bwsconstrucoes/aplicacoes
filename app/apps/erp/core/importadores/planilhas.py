# ============================================================================
# BWS ERP — core/importadores/planilhas.py
# Importação de OBRAS e CATEGORIAS (plano financeiro) a partir de CSV
# exportado das planilhas Google. O de-para coluna→campo é explícito e
# apresentado ao usuário ANTES da importação (governança combinada).
#
# Formato esperado (cabeçalhos, maiúsc./minúsc. indiferente):
#   obras.csv:      codigo,nome,cno,municipio,uf,codigo_omie_depto
#   categorias.csv: codigo,descricao,codigo_omie,dedutivel_padrao,tipos_permitidos
#     - tipos_permitidos: lista separada por ; (ex.: T1_MATERIAL_NFE;T2_SERVICO_NFSE)
#     - vazio em tipos_permitidos = aceita todos os tipos
# ============================================================================
from __future__ import annotations

import csv
import io
from typing import Any, Optional

from sqlalchemy.orm import Session

from app.apps.erp.core.cadastros import categorias as svc_cat
from app.apps.erp.core.cadastros import obras as svc_obra
from app.apps.erp.core.comum.auditoria import ErroValidacao
from app.apps.erp.db.models.cadastros import Usuario


def ler_tabela(conteudo: bytes, *,
               colunas_esperadas: tuple[str, ...] = ()) -> list[dict[str, str]]:
    """Lê a planilha venha ela em CSV ou em Excel (.xlsx).

    Existe porque exportar para CSV é um passo a mais que ninguém lembra de
    fazer — e quando faz, o acento e o ponto e vírgula viram problema. O
    formato é reconhecido pelo CONTEÚDO, não pela extensão: arquivo .xlsx
    começa com a assinatura "PK" de arquivo compactado. Assim o nome do
    arquivo pode estar errado que a leitura funciona igual.

    Devolve o mesmo formato do CSV — uma lista de dicionários com o cabeçalho
    em minúsculas — para o resto do importador não precisar saber a diferença.
    `openpyxl` já é dependência do serviço (relatório do painel em Excel):
    nenhuma biblioteca nova entra por causa disto.
    """
    if conteudo[:2] == b"PK":
        return _ler_xlsx(conteudo, colunas_esperadas)
    return _ler_csv(conteudo)


def _ler_xlsx(conteudo: bytes,
              colunas_esperadas: tuple[str, ...] = ()) -> list[dict[str, str]]:
    """Primeira aba da pasta, em modo de leitura (não carrega tudo em memória
    de uma vez — a planilha de insumos da BWS tem mais de 3 mil linhas e a
    instância divide 2 GB com os outros módulos).

    QUANDO A ABA CERTA NÃO É A PRIMEIRA (18/09/2026): a planilha do banco de
    preços tem várias abas, e a que interessa não abre primeiro. Quem chama diz
    quais colunas espera, e a aba é escolhida por elas — não pela posição.
    Pedir para o dono reordenar as abas da planilha dele seria transformar um
    detalhe nosso em trabalho dele.
    """
    from openpyxl import load_workbook

    try:
        pasta = load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    except Exception as e:
        raise ErroValidacao(f"Não consegui abrir a planilha Excel: {e}") from e
    if not pasta.worksheets:
        raise ErroValidacao("A planilha Excel não tem nenhuma aba.")
    aba = _aba_certa(pasta, colunas_esperadas)
    linhas: list[dict[str, str]] = []
    cabecalho: list[str] = []
    alvo = {_simples(c) for c in colunas_esperadas}
    for bruta in aba.iter_rows(values_only=True):
        valores = ["" if v is None else str(v).strip() for v in bruta]
        if not cabecalho:
            if not any(valores):
                continue                 # linha em branco antes do cabeçalho
            # COM COLUNAS ESPERADAS, o cabeçalho é a linha que as traz — não a
            # primeira linha escrita. Planilha de gente tem título, logotipo e
            # painel de filtro em cima da tabela, e tomar a primeira linha como
            # cabeçalho faria a importação inteira ler colunas chamadas
            # "Preços de Insumos Cotados" e "".
            if alvo and not alvo.issubset({_simples(v) for v in valores if v}):
                continue
            cabecalho = [v.strip().lower() for v in valores]
            continue
        if not any(valores):
            continue                     # linha em branco no meio ou no fim
        linhas.append({c: (valores[i] if i < len(valores) else "")
                       for i, c in enumerate(cabecalho) if c})
    pasta.close()
    return linhas


def _aba_certa(pasta, colunas_esperadas: tuple[str, ...]):
    """A primeira aba cujo cabeçalho traz as colunas pedidas; se nenhuma
    trouxer, a primeira da pasta — que é o comportamento de sempre."""
    if not colunas_esperadas:
        return pasta.worksheets[0]
    alvo = {_simples(c) for c in colunas_esperadas}
    for aba in pasta.worksheets:
        # O cabeçalho às vezes não é a primeira linha (título, logo, filtro em
        # cima). Vinte linhas é folga suficiente e não custa leitura.
        for i, bruta in enumerate(aba.iter_rows(values_only=True)):
            if i >= 20:
                break
            achadas = {_simples(str(v)) for v in bruta if v is not None}
            if alvo.issubset(achadas):
                return aba
    return pasta.worksheets[0]


def _simples(texto: str) -> str:
    """Comparável: sem acento, sem caixa, sem espaço sobrando. "Nº Mapa" e
    "nº mapa" têm de cair no mesmo lugar."""
    import unicodedata
    bruto = unicodedata.normalize("NFKD", (texto or "").strip().lower())
    sem_acento = "".join(c for c in bruto if not unicodedata.combining(c))
    return " ".join(sem_acento.split())


def _ler_csv(conteudo: bytes) -> list[dict[str, str]]:
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            texto = conteudo.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ErroValidacao("CSV com codificação não reconhecida.")
    amostra = texto[:2048]
    delim = ";" if amostra.count(";") > amostra.count(",") else ","
    leitor = csv.DictReader(io.StringIO(texto), delimiter=delim)
    return [{(k or "").strip().lower(): (v or "").strip() for k, v in linha.items()}
            for linha in leitor]


def importar_obras_csv(s: Session, conteudo: bytes,
                       usuario: Optional[Usuario]) -> dict[str, Any]:
    linhas = _ler_csv(conteudo)
    criadas, rejeitadas = 0, []
    for i, ln in enumerate(linhas, start=2):     # linha 1 = cabeçalho
        try:
            svc_obra.criar(s, {
                "codigo": ln.get("codigo"), "nome": ln.get("nome"),
                "cno": ln.get("cno"), "municipio": ln.get("municipio"),
                "uf": ln.get("uf"), "codigo_omie_depto": ln.get("codigo_omie_depto"),
                "objeto": ln.get("objeto"), "cliente": ln.get("cliente"),
                "cnpj_cliente": ln.get("cnpj_cliente"), "contrato": ln.get("contrato"),
                "valor_contrato": ln.get("valor_contrato"),
                "aliquota_iss": ln.get("aliquota_iss"), "tributacao": ln.get("tributacao"),
                "data_inicio": ln.get("data_inicio"), "data_termino": ln.get("data_termino"),
                "orgao_resumido": ln.get("orgao_resumido"), "ref_pipefy": ln.get("ref_pipefy"),
                "origem": "IMPORTACAO_CSV",
            }, usuario)
            criadas += 1
        except ErroValidacao as e:
            rejeitadas.append({"linha": i, "codigo": ln.get("codigo"), "motivo": str(e)})
    return {"no_arquivo": len(linhas), "criadas": criadas, "rejeitadas": rejeitadas}


def importar_categorias_csv(s: Session, conteudo: bytes,
                            usuario: Optional[Usuario]) -> dict[str, Any]:
    linhas = _ler_csv(conteudo)
    criadas, rejeitadas = 0, []
    for i, ln in enumerate(linhas, start=2):
        try:
            tipos = [t.strip() for t in (ln.get("tipos_permitidos") or "").split(";") if t.strip()]
            dedutivel = (ln.get("dedutivel_padrao") or "sim").strip().lower() not in ("nao", "não", "n", "0", "false")
            svc_cat.criar(s, {
                "codigo": ln.get("codigo"), "descricao": ln.get("descricao"),
                "codigo_omie": ln.get("codigo_omie"),
                "natureza": ln.get("natureza") or "RESULTADO",
                "tipos_permitidos": tipos, "dedutivel_padrao": dedutivel,
                "origem": "IMPORTACAO_CSV",
            }, usuario)
            criadas += 1
        except ErroValidacao as e:
            rejeitadas.append({"linha": i, "codigo": ln.get("codigo"), "motivo": str(e)})
    return {"no_arquivo": len(linhas), "criadas": criadas, "rejeitadas": rejeitadas}
