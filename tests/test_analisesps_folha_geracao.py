# -*- coding: utf-8 -*-
"""
Os arquivos de pagamento da folha — BeeVale e SomaPay (27/09/2026).

⚠️ ESTES TESTES VALEM DINHEIRO DE GENTE. O que eles travam:

  - **o mesmo CPF duas vezes** é permitido no BeeVale (naturezas diferentes) e é
    TRAVA no SomaPay. Palavras do dono: *"o Soma não aceita"*. Um arquivo recusado
    depois de subir custa a rodada inteira — gerar, subir, criar card, avisar;
  - **um arquivo por conta, sempre.** Juntar contas pagaria da conta errada, e
    ninguém veria até o extrato;
  - **o CPF vai como TEXTO e formatado.** Como número, perde o zero da frente — e
    aí o portal paga outra pessoa, ou ninguém;
  - **valor zero não entra.** O portal recusa a linha, e a pessoa sairia do
    arquivo sem explicação nenhuma.
"""
import io
from decimal import Decimal as D

import pytest

from app.apps.analisesps import folha_geracao as g

GERLANIO = "997.133.493-34"
ANA = "035.134.413-63"


def _linha(cpf=GERLANIO, nome="GERLANIO GOMES LIMA", conta="50024",
           verba="alimentacao", valor="330.00", obra="CREPEOLINDA"):
    return {"cpf": cpf, "nome": nome, "conta": conta, "verba": verba,
            "valor": valor, "obra": obra}


def _abrir(conteudo: bytes):
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(conteudo)).active


# ---------------------------------------------------------------------------
# A DIVISÃO EM ARQUIVOS
# ---------------------------------------------------------------------------
def test_um_arquivo_por_conta_SEMPRE():
    """Não é escolha: a conta define de onde o dinheiro sai."""
    lotes = g.montar_lotes([_linha(conta="50024"), _linha(cpf=ANA, nome="ANA",
                                                          conta="50025")],
                           g.BEEVALE, juntar_verbas=True)
    assert [l["conta"] for l in lotes] == ["50024", "50025"]
    assert all(l["quantos"] == 1 for l in lotes)


def test_o_beevale_JUNTA_verbas_quando_ele_pede():
    """*"De repente eu quero gerar alimentação e transporte no mesmo arquivo do
    BeeVale. Ao invés de fazer três pagamentos, a gente faz só um."*"""
    lotes = g.montar_lotes([_linha(verba="alimentacao", valor="330.00"),
                            _linha(verba="transporte", valor="200.00")],
                           g.BEEVALE, juntar_verbas=True)
    assert len(lotes) == 1, "um arquivo só"
    assert lotes[0]["verbas"] == ["alimentacao", "transporte"]
    assert lotes[0]["total"] == D("530.00")
    # ⚠️ A MESMA PESSOA APARECE DUAS VEZES, e pode: a natureza da verba é outra.
    assert lotes[0]["quantos"] == 2
    assert {i["natureza"] for i in lotes[0]["linhas"]} == {
        "Despesas com Alimentação", "Despesas com Transporte"}


def test_o_beevale_SEPARA_quando_ele_nao_pede():
    lotes = g.montar_lotes([_linha(verba="alimentacao"),
                            _linha(verba="transporte", valor="200.00")],
                           g.BEEVALE, juntar_verbas=False)
    assert len(lotes) == 2
    assert sorted(v for l in lotes for v in l["verbas"]) == [
        "alimentacao", "transporte"]


def test_o_somapay_SEPARA_SOZINHO_mesmo_se_pedirem_para_juntar():
    """⚠️ A TRAVA DO SOMAPAY. Ele não aceita o mesmo CPF duas vezes, então não
    existe "juntar" — e o sistema separa em vez de deixar marcar e dar erro."""
    lotes = g.montar_lotes([_linha(verba="alimentacao"),
                            _linha(verba="transporte", valor="200.00")],
                           g.SOMAPAY, juntar_verbas=True)
    assert len(lotes) == 2, "separou sozinho"
    assert all(l["quantos"] == 1 for l in lotes)
    assert all(not l["criticas"] for l in lotes)


def test_duas_linhas_da_MESMA_natureza_viram_uma():
    """É o que o `BeeVale.gs` faz: consolida por conta + CPF + natureza. E é isto
    que impede o CPF repetido no SomaPay."""
    lotes = g.montar_lotes([_linha(valor="330.00"), _linha(valor="70.00")],
                           g.SOMAPAY)
    assert len(lotes) == 1
    assert lotes[0]["quantos"] == 1
    assert lotes[0]["linhas"][0]["valor"] == D("400.00")
    assert lotes[0]["linhas"][0]["juntou"] == 2


def test_valor_zero_ou_negativo_nao_entra():
    """O portal recusa a linha, e a pessoa sairia do arquivo sem explicação."""
    lotes = g.montar_lotes([_linha(valor="0"), _linha(cpf=ANA, nome="ANA",
                                                      valor="-10.00"),
                            _linha(valor="100.00")], g.BEEVALE)
    assert len(lotes) == 1
    assert lotes[0]["quantos"] == 1
    assert lotes[0]["total"] == D("100.00")


def test_pessoa_sem_CPF_vira_CRITICA_e_TRAVA_o_lote():
    """Sem CPF o portal não tem como pagar — e gerar assim gasta a rodada."""
    lotes = g.montar_lotes([_linha(cpf="123")], g.BEEVALE)
    assert lotes[0]["pode_gerar"] is False
    assert any("sem CPF" in c for c in lotes[0]["criticas"])


def test_pessoa_sem_NOME_vira_critica():
    """O banco confere o nome contra o CPF e recusa a linha."""
    lotes = g.montar_lotes([_linha(nome="")], g.BEEVALE)
    assert lotes[0]["pode_gerar"] is False
    assert any("sem nome" in c for c in lotes[0]["criticas"])


def test_linha_sem_conta_diz_o_que_consertar():
    """Obra sem conta é estado legítimo do cadastro — o recado tem de dizer onde
    se conserta, não só que faltou."""
    lotes = g.montar_lotes([_linha(conta="")], g.BEEVALE)
    assert lotes[0]["pode_gerar"] is False
    assert any("C. Diários" in c for c in lotes[0]["criticas"])


def test_destino_desconhecido_e_recusado():
    with pytest.raises(g.ErroDaGeracao) as erro:
        g.montar_lotes([_linha()], "pix")
    assert "BeeVale e SomaPay" in str(erro.value)


def test_o_teto_de_linhas_protege_de_laco_torto(monkeypatch):
    monkeypatch.setattr(g, "MAXIMO_POR_ARQUIVO", 1)
    lotes = g.montar_lotes([_linha(), _linha(cpf=ANA, nome="ANA")], g.BEEVALE)
    assert lotes[0]["pode_gerar"] is False
    assert any("teto" in c for c in lotes[0]["criticas"])


def test_o_resumo_diz_quantos_arquivos_e_quanto_ANTES_de_gerar():
    """Pedido dele: *"mostrar, antes de gerar, quantos arquivos vão sair e com que
    total cada um"*."""
    lotes = g.montar_lotes([_linha(conta="50024"),
                            _linha(cpf=ANA, nome="ANA", conta="50025",
                                   valor="100.00")], g.BEEVALE)
    resumo = g.resumo_dos_lotes(lotes)
    assert resumo["arquivos"] == 2
    assert resumo["total"] == D("430.00")
    assert resumo["pessoas"] == 2
    assert resumo["pode_gerar"] is True
    assert resumo["com_critica"] == []


def test_o_resumo_de_lista_vazia_nao_libera_a_geracao():
    """"Nada a pagar" não é "pode gerar": geraria um arquivo com só o cabeçalho, e
    o portal aceitaria calado."""
    assert g.resumo_dos_lotes([])["pode_gerar"] is False
    assert g.resumo_dos_lotes(None)["arquivos"] == 0


# ---------------------------------------------------------------------------
# O ARQUIVO DO SOMAPAY
# ---------------------------------------------------------------------------
def test_o_arquivo_do_somapay_e_o_MODELO_DO_SOMA_preenchido():
    """⚠️ Em 01/10/2026 o arquivo montado aqui não passou no portal. O modelo do
    site do Soma tem a aba "Planilha de Folha de Pagamento", título e instruções
    em cima, cabeçalho na linha 11 e dados a partir da 12."""
    aba = _abrir(g.somapay_xlsx([{"cpf": GERLANIO, "nome": "GERLANIO GOMES LIMA",
                                  "valor": "1126.60"}]))
    assert aba.title == "Planilha de Folha de Pagamento"
    assert "PLANILHA DE FOLHA DE PAGAMENTO" in aba["A1"].value
    assert aba["A4"].value == "Instruções:"
    assert [c.value for c in aba[11]][:3] == [
        "Nome do funcionário", "CPF* (obrigatório)", "Valor* (obrigatório)"]
    assert aba["A12"].value == "GERLANIO GOMES LIMA"
    assert aba["B12"].value == "99713349334", "só os 11 dígitos, como no modelo"
    assert aba["C12"].value == pytest.approx(1126.60)
    assert aba["C12"].number_format == "[$R$ -416]#,##0.00"
    assert aba["A13"].value is None


def test_fora_a_aba_de_dados_o_arquivo_e_IGUAL_ao_modelo():
    """Logotipo, estilos, propriedades: tudo é o arquivo deles."""
    import zipfile
    gerado = zipfile.ZipFile(io.BytesIO(g.somapay_xlsx(
        [{"cpf": ANA, "nome": "ANA", "valor": "10.00"}])))
    modelo = zipfile.ZipFile(g.MODELO_SOMAPAY)
    assert gerado.namelist() == modelo.namelist()
    diferentes = [n for n in modelo.namelist()
                  if gerado.read(n) != modelo.read(n)]
    assert diferentes == ["xl/worksheets/sheet1.xml"]


def test_o_CPF_vai_como_TEXTO_e_com_o_ZERO_da_frente_no_somapay():
    """⚠️ Como número, um CPF que começa com zero perde o zero — e o portal paga
    outra pessoa, ou ninguém."""
    aba = _abrir(g.somapay_xlsx([{"cpf": ANA, "nome": "ANA", "valor": "10.00"}]))
    assert aba["B12"].number_format == "@"
    assert aba["B12"].data_type == "s"
    assert aba["B12"].value == "03513441363"


def test_nome_com_caractere_especial_nao_quebra_o_arquivo():
    aba = _abrir(g.somapay_xlsx([{"cpf": ANA, "nome": "JOSÉ & <FILHO>",
                                  "valor": "10.00"}]))
    assert aba["A12"].value == "JOSÉ & <FILHO>"


def test_mais_gente_que_as_linhas_do_modelo_ganha_linhas_novas():
    """O modelo tem linhas até a 1000; o que passar disso não pode sumir."""
    linhas = [{"cpf": f"{n:011d}", "nome": f"P{n}", "valor": "1.00"}
              for n in range(1, 1201)]
    aba = _abrir(g.somapay_xlsx(linhas))
    assert aba["A12"].value == "P1"
    assert aba["A1000"].value == "P989"
    assert aba["A1211"].value == "P1200"
    assert aba["B1211"].value == "00000001200"


def test_o_arquivo_do_somapay_tem_o_nome_terminando_em_xls_como_o_modelo():
    lote = g.montar_lotes([_linha(verba="folha")], g.SOMAPAY)[0]
    assert g.nome_do_arquivo(lote, 2026, 9, "quinzena").endswith(".xls")


def test_o_somapay_RECUSA_o_arquivo_com_CPF_repetido():
    """A trava, com o nome de quem está repetido: é o que permite consertar."""
    with pytest.raises(g.ErroDaGeracao) as erro:
        g.somapay_xlsx([{"cpf": GERLANIO, "nome": "A", "valor": "10.00"},
                        {"cpf": GERLANIO, "nome": "A", "valor": "20.00"}])
    assert "997.133.493-34" in str(erro.value)
    assert "SomaPay" in str(erro.value)


# ---------------------------------------------------------------------------
# O ARQUIVO DO BEEVALE
# ---------------------------------------------------------------------------
def test_o_arquivo_do_beevale_usa_as_MESMAS_colunas_do_fluxo_das_SPs():
    """Duas listas de colunas divergiriam no dia em que o portal mudasse uma — e só
    uma seria corrigida."""
    from app.apps.analisesps.beevale import COLUNAS_PAGAMENTO

    aba = _abrir(g.beevale_xlsx([{"cpf": GERLANIO, "nome": "GERLANIO",
                                  "valor": "330.00", "obra": "CREPEOLINDA"}]))
    assert [c.value for c in aba[1]] == COLUNAS_PAGAMENTO
    assert aba["A2"].value == "GERLANIO"
    assert aba["B2"].value == "99713349334@bwsconstrucoes.com.br"
    assert aba["C2"].value == "Produção"
    assert aba["D2"].value == "Livre"
    assert aba["E2"].value == pytest.approx(330.00)
    assert aba["F2"].value == "Mensal"
    assert aba["G2"].value == 0
    assert aba["H2"].value == "997.133.493-34"
    assert aba["J2"].value == "CREPEOLINDA"
    assert aba["K2"].value == "BWS"


def test_o_email_do_beevale_e_SEMPRE_o_mesmo_para_a_mesma_pessoa():
    """É chave no portal: dois e-mails diferentes criariam duas carteiras para a
    mesma pessoa, e metade do dinheiro ficaria numa delas."""
    assert g.email_do_cpf("997.133.493-34") == g.email_do_cpf("99713349334")


def test_o_beevale_aceita_o_mesmo_CPF_duas_vezes():
    """Pode, se a natureza da verba for diferente — e é justamente o que permite
    juntar alimentação e transporte num pagamento só."""
    conteudo = g.beevale_xlsx([
        {"cpf": GERLANIO, "nome": "GERLANIO", "valor": "330.00"},
        {"cpf": GERLANIO, "nome": "GERLANIO", "valor": "200.00"}])
    aba = _abrir(conteudo)
    assert aba.max_row == 3


# ---------------------------------------------------------------------------
# O NOME DO ARQUIVO
# ---------------------------------------------------------------------------
def test_o_nome_do_arquivo_diz_TUDO_o_que_identifica_ele():
    """⚠️ Já subiram o arquivo da conta errada — está escrito no comentário do
    script do dono. O nome é a primeira defesa contra isso."""
    lote = g.montar_lotes([_linha(verba="alimentacao"),
                           _linha(verba="transporte", valor="200.00")],
                          g.BEEVALE, juntar_verbas=True)[0]
    nome = g.nome_do_arquivo(lote, 2026, 9, "quinzena")
    assert nome.startswith("BeeVale - 09-2026 - Quinzena")
    assert "Alimentação+Transporte" in nome
    assert "conta 50024" in nome
    assert nome.endswith(".xlsx")


def test_o_nome_do_arquivo_nao_cria_pasta_no_drive():
    """Barra no nome vira pasta, e o arquivo desaparece de onde se procura."""
    lote = {"destino": g.SOMAPAY, "conta": "50024/2", "verbas": ["folha"]}
    assert "/" not in g.nome_do_arquivo(lote, 2026, 9, "fim_de_mes")


def test_arquivo_de_lote_vazio_e_recusado():
    with pytest.raises(g.ErroDaGeracao):
        g.arquivo_do_lote({"destino": g.BEEVALE, "linhas": []})


def test_arquivo_do_lote_escolhe_o_layout_do_destino():
    lote_bee = g.montar_lotes([_linha()], g.BEEVALE)[0]
    lote_soma = g.montar_lotes([_linha()], g.SOMAPAY)[0]
    assert _abrir(g.arquivo_do_lote(lote_bee)).title == "Pagamento BeeVale"
    assert _abrir(g.arquivo_do_lote(lote_soma)).title == "Planilha de Folha de Pagamento"


# ---------------------------------------------------------------------------
# O ARQUIVO DE ANÁLISE — o segundo arquivo
#
# *"Tem que ter no mínimo o do arquivo de pagamento e um de análise da folha. Com
#  as informações separadas, agrupadas, qual obra, qual funcionário, rateio de
#  folhas, se for o caso."*
# ---------------------------------------------------------------------------
def _analise(lotes, **k):
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(g.analise_xlsx(lotes, **k)))


def test_a_analise_tem_as_QUATRO_visoes_que_ele_pediu():
    lotes = g.montar_lotes([_linha(), _linha(cpf=ANA, nome="ANA",
                                             obra="CREPEAREIAS",
                                             valor="170.00")], g.BEEVALE)
    livro = _analise(lotes, ano=2026, mes=9, tipo="quinzena")
    assert livro.sheetnames == ["Resumo", "Por obra", "Por funcionário",
                               "Rateio"]


def test_a_analise_mostra_CADA_OBRA_da_pessoa_mesmo_paga_pela_mesma_conta():
    """⚠️ Conserto de 01/10/2026. O arquivo junta a pessoa por conta — então quem
    tinha duas obras pagas pela mesma conta aparecia só na primeira, com o
    dinheiro todo nela, e os dias saíam zerados. O detalhe (uma linha por pessoa
    e obra) é o que a análise usa agora."""
    detalhe = [_linha(valor="200.00", obra="CREPEOLINDA"),
               _linha(valor="130.00", obra="CREPEAREIAS")]
    for linha, dias in zip(detalhe, (6, 5)):
        linha.update({"dias": dias, "origem": "ponto"})
    lotes = g.montar_lotes(detalhe, g.BEEVALE)
    assert lotes[0]["quantos"] == 1, "no arquivo, a pessoa é uma linha só"

    livro = _analise(lotes, ano=2026, mes=9, tipo="quinzena", detalhe=detalhe)
    por_obra = {r[0].value: r[2].value for r in livro["Por obra"].iter_rows(min_row=2)}
    assert por_obra == {"CREPEOLINDA": 200, "CREPEAREIAS": 130}
    gente = [[c.value for c in r] for r in livro["Por funcionário"].iter_rows(min_row=2)]
    assert sorted((l[4], l[5], l[6]) for l in gente) == [
        ("CREPEAREIAS", 5, 130), ("CREPEOLINDA", 6, 200)]
    assert {l[7] for l in gente} == {"ponto"}


def test_a_analise_agrupa_por_obra_e_por_funcionario():
    lotes = g.montar_lotes([_linha(valor="330.00"),
                            _linha(cpf=ANA, nome="ANA", obra="CREPEAREIAS",
                                   valor="170.00")], g.BEEVALE)
    livro = _analise(lotes, ano=2026, mes=9, tipo="quinzena")

    por_obra = [[c.value for c in r] for r in livro["Por obra"].iter_rows()]
    assert por_obra[0] == ["Obra", "Pessoas", "Total", "% do pagamento"]
    assert por_obra[1][:3] == ["CREPEOLINDA", 1, 330]
    assert por_obra[2][:3] == ["CREPEAREIAS", 1, 170]

    gente = [[c.value for c in r] for r in livro["Por funcionário"].iter_rows()]
    assert gente[0][:5] == ["Nome", "CPF", "Conta", "Verba", "Obra"]
    assert "997.133.493-34" in {l[1] for l in gente[1:]}


def test_o_rateio_da_analise_FECHA_100_por_cento():
    """⚠️ É O NÚMERO QUE VAI NO CARD. Um rateio que soma 99,9999% é recusado, e o
    card volta para a mão de alguém."""
    lotes = g.montar_lotes([
        _linha(valor="100.00", obra="A"),
        _linha(cpf=ANA, nome="ANA", valor="100.00", obra="B"),
        _linha(cpf="123.456.789-09", nome="JOÃO", valor="100.00", obra="C"),
    ], g.BEEVALE)
    livro = _analise(lotes, ano=2026, mes=9, tipo="quinzena")
    linhas = [[c.value for c in r] for r in livro["Rateio"].iter_rows()][1:]
    assert len(linhas) == 3
    assert sum(l[1] for l in linhas) == pytest.approx(100.0)


def test_o_percentual_tem_SETE_casas_como_o_card_usa():
    saida = g.percentuais_por_obra([{"obra": "A", "total": "100.00"},
                                    {"obra": "B", "total": "200.00"}])
    assert [o["obra"] for o in saida] == ["B", "A"], "da maior para a menor"
    assert sum(o["percentual"] for o in saida) == D("100")
    assert saida[0]["percentual"].as_tuple().exponent == -7


def test_a_sobra_do_percentual_vai_para_a_MAIOR_fatia():
    """Mesma regra da sobra do centavo: a maior fatia é a que menos sente."""
    saida = g.percentuais_por_obra([{"obra": "A", "total": "100.00"},
                                    {"obra": "B", "total": "100.00"},
                                    {"obra": "C", "total": "100.00"}])
    assert sum(o["percentual"] for o in saida) == D("100")


def test_percentual_de_pagamento_vazio_nao_estoura():
    assert g.percentuais_por_obra([]) == []
    assert g.percentuais_por_obra([{"obra": "A", "total": "0"}]) == []


def test_os_AVISOS_entram_no_arquivo_de_analise():
    """⚠️ Aviso que só existe na tela não sobrevive à rodada. Quem abrir a análise
    para explicar uma diferença precisa ver o que estava avisado na hora."""
    lotes = g.montar_lotes([_linha(conta="")], g.BEEVALE)
    livro = _analise(lotes, ano=2026, mes=9, tipo="quinzena")
    texto = " ".join(str(c.value) for r in livro["Resumo"].iter_rows()
                     for c in r if c.value)
    assert "Avisos da geração" in texto
    assert "C. Diários" in texto


def test_a_analise_usa_a_apropriacao_quando_ela_existe():
    """Com a apropriação, a aba por funcionário mostra os dias e de onde veio cada
    valor — que é o que se audita."""
    lotes = g.montar_lotes([_linha()], g.BEEVALE)
    livro = _analise(lotes, ano=2026, mes=9, tipo="quinzena", pessoas=[{
        "cpf": "99713349334", "nome": "GERLANIO", "nome_cadastro": "GERLANIO G",
        "por_obra": [{"obra": "CREPEOLINDA", "dias": 8, "valor": D("330.00"),
                      "origem": "ponto"}]}])
    linhas = [[c.value for c in r]
              for r in livro["Por funcionário"].iter_rows()][1:]
    assert linhas[0][0] == "GERLANIO G", "o nome do cadastro é o que fica"
    assert linhas[0][5] == 8
    assert linhas[0][7] == "ponto"


def test_a_analise_diz_a_competencia_e_o_pagamento():
    """Arquivo de análise sem competência é arquivo que ninguém sabe de quando é."""
    lotes = g.montar_lotes([_linha()], g.SOMAPAY)
    livro = _analise(lotes, ano=2026, mes=9, tipo="fim_de_mes")
    texto = " ".join(str(c.value) for r in livro["Resumo"].iter_rows()
                     for c in r if c.value)
    assert "09/2026" in texto
    assert "Fim de mês" in texto
    assert "SomaPay" in texto
