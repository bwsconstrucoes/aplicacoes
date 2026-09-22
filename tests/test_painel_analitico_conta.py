# -*- coding: utf-8 -*-
"""
Despesas Analítico: o filtro pela conta de pagamento, e o relatório com abas.

22/09/2026, o dono: "coloque no despesas analítico filtro pra conta de
pagamento. Gostaria ainda que você melhorasse o relatório, tá muito pobre. Acho
que pode ter outras abas."

De passagem, um defeito: a barra do Analítico não levava adiante a conta
escolhida na barra lateral — quem filtrava uma conta e clicava em Aplicar via
o filtro sumir, calado. É provavelmente por isso que o filtro "não existia".
"""
from __future__ import annotations

import io
import os

import pytest

pytestmark = pytest.mark.banco

SENHA_MESTRE = "senha-do-dono"


def _despesa(conn, cod, obra, conta, credor, valor, data="2025-03-10"):
    conn.execute(
        "INSERT INTO fato (codigo_lancamento, tipo, analise, situacao,"
        " situacao_vencimento, categoria, grupo, departamento, projeto,"
        " razao_social, conta_corrente, data, ano, pago_recebido,"
        " a_pagar_receber, juros, multa)"
        " VALUES (?,'2. Contas a Pagar','DRE','PAGO','Quitado','Serviços',"
        "         'Custo',?,'ALFA',?,?,?,2025,?,0,0,0)",
        (cod, obra, credor, conta, data, valor))


@pytest.fixture()
def base():
    from tests.conftest import VARIAVEL_BANCO_TESTE, url_de_teste_segura
    bruto = os.environ.get(VARIAVEL_BANCO_TESTE, "").strip()
    if not bruto:
        pytest.skip(f"{VARIAVEL_BANCO_TESTE} não definida — testes com banco pulados")
    os.environ["DATABASE_URL"] = url_de_teste_segura(bruto)

    from app.apps.painel import consultas
    from app.apps.painel import db as painel_db
    from app.apps.painel import migracoes_runner
    painel_db._engine = None
    assert not migracoes_runner.aplicar_pendentes().get("erro")
    with painel_db.conexao() as conn:
        conn.execute("TRUNCATE TABLE fato")
        _despesa(conn, 901, "OBRA A", "Itaú", "FORNECEDOR DO ITAU", -1000)
        _despesa(conn, 902, "OBRA A", "Itaú", "FORNECEDOR DO ITAU", -500, "2025-04-10")
        _despesa(conn, 903, "OBRA A", "Bradesco", "FORNECEDOR DO BRADESCO", -200)
        _despesa(conn, 904, "OBRA B", "Bradesco", "FORNECEDOR DO BRADESCO", -300)
        conn.commit()
    consultas.esquecer_listas()
    yield
    with painel_db.conexao() as conn:
        conn.execute("TRUNCATE TABLE fato")
        conn.commit()


def _cliente(monkeypatch):
    monkeypatch.setenv("PAINEL_SENHA", SENHA_MESTRE)
    from app.main import create_app
    app = create_app()
    app.config.update(TESTING=True)
    cliente = app.test_client()
    cliente.post("/painel/entrar", data={"usuario": "", "senha": SENHA_MESTRE})
    return cliente


def test_o_filtro_de_conta_recorta_a_lista(base, monkeypatch):
    html = _cliente(monkeypatch).get("/painel/analitico?conta=Ita%C3%BA").get_data(as_text=True)
    assert "FORNECEDOR DO ITAU" in html
    assert "FORNECEDOR DO BRADESCO" not in html


def test_a_caixa_da_tela_mostra_a_conta_escolhida(base, monkeypatch):
    html = _cliente(monkeypatch).get("/painel/analitico?conta=Ita%C3%BA").get_data(as_text=True)
    assert 'name="conta"' in html
    assert 'value="Itaú" selected' in html


def test_varias_contas_da_barra_lateral_nao_se_perdem_ao_aplicar(base, monkeypatch):
    """O defeito: a caixa desta tela só cabe uma conta, e as duas marcadas na
    barra lateral sumiam ao clicar em Aplicar. Agora viajam escondidas."""
    html = _cliente(monkeypatch).get(
        "/painel/analitico?conta=Ita%C3%BA&conta=Bradesco").get_data(as_text=True)
    assert html.count('type="hidden" name="conta"') == 2
    assert "2 contas, pela barra lateral" in html


def test_o_dre_tambem_leva_a_conta_adiante(base, monkeypatch):
    html = _cliente(monkeypatch).get("/painel/dre?conta=Ita%C3%BA").get_data(as_text=True)
    assert 'type="hidden" name="conta" value="Itaú"' in html


def test_o_relatorio_tem_uma_aba_por_recorte_e_elas_fecham_com_a_lista(base, monkeypatch):
    import openpyxl
    r = _cliente(monkeypatch).get("/painel/baixar/analitico")
    assert r.status_code == 200
    livro = openpyxl.load_workbook(io.BytesIO(r.get_data()), read_only=True)
    nomes = livro.sheetnames
    for esperada in ("Resumo", "Por grupo", "Por categoria", "Por credor",
                     "Por conta de pagamento", "Por obra", "Por mes", "Despesas Analitico"):
        assert esperada in nomes, nomes

    por_conta = {linha[0]: linha for linha in livro["Por conta de pagamento"].iter_rows(
        min_row=2, values_only=True)}
    # nome, lancamentos, pago, a pagar, juros e multa, total, %
    assert por_conta["Itaú"][1] == 2 and por_conta["Itaú"][5] == -1500
    assert por_conta["Bradesco"][1] == 2 and por_conta["Bradesco"][5] == -500

    por_mes = [linha[0] for linha in livro["Por mes"].iter_rows(min_row=2, values_only=True)]
    assert por_mes == ["03/2025", "04/2025"], "o mês tem de sair em ordem cronológica"

    lancamentos = list(livro["Despesas Analitico"].iter_rows(min_row=2, values_only=True))
    assert len(lancamentos) == 4


def test_o_relatorio_respeita_o_filtro_de_conta(base, monkeypatch):
    import openpyxl
    r = _cliente(monkeypatch).get("/painel/baixar/analitico?conta=Ita%C3%BA")
    livro = openpyxl.load_workbook(io.BytesIO(r.get_data()), read_only=True)
    contas = [l[0] for l in livro["Por conta de pagamento"].iter_rows(min_row=2, values_only=True)]
    assert contas == ["Itaú"]
    # varias linhas "Filtro" — por isso lista, e nao dicionario
    resumo = list(livro["Resumo"].iter_rows(min_row=2, values_only=True))
    assert ("Total", -1500) in resumo
    assert any("Itaú" in str(v) for _k, v in resumo), "o filtro tem de estar escrito no arquivo"


def test_o_pdf_do_analitico_sai(base, monkeypatch):
    r = _cliente(monkeypatch).get("/painel/baixar/analitico?formato=pdf")
    assert r.status_code == 200
    assert r.get_data()[:4] == b"%PDF"
