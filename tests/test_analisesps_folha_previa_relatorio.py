# -*- coding: utf-8 -*-
"""
01/10/2026 — a prévia do arquivo de pagamento, o relatório da folha (Excel e PDF)
e a lateral da folha no padrão das Solicitações.

Pedidos dele, no mesmo dia:
  *"se eu quiser gerar um arquivo de pagamento sem fechar, como fazer? até pra
   saber como tá saindo"*
  *"relatórios excel e pdf (…) o relatório do que eu visualizo em tela, além de
   poder ver o agrupamento do pagamento. Por obra, por conta e etc."*
  *"Ajuste o filtro do sidebar no mesmo padrão que o de solicitações"*
  *"o botão apagar esta folha tá integralmente vermelho, texto e fundo"*
"""
import io
import zipfile
from decimal import Decimal as D
from pathlib import Path

from tests.test_analisesps_telas import app  # noqa: F401 — a fixture
from tests.test_analisesps_telas import (SENHA_CONSULTA, _como_mestre,
                                         _dias_do_mes, _dias_em_duas_obras,
                                         _preparar_folha_aberta, como)


# ---------------------------------------------------------------------------
# O BOTÃO APAGAR
# ---------------------------------------------------------------------------
def test_o_botao_de_apagar_tem_FUNDO_BRANCO_e_texto_vermelho():
    """`.btn.perigo` pinta o fundo de vermelho; o secundário perigoso tem de
    voltar o fundo para branco, senão vira texto vermelho em fundo vermelho."""
    css = Path("app/apps/analisesps/static/analisesps.css").read_text(encoding="utf-8")
    regra = [l for l in css.splitlines() if l.startswith(".btn.secundario.perigo")]
    assert regra and "background: var(--branco)" in regra[0]
    assert "color: var(--vermelho)" in regra[0]


# ---------------------------------------------------------------------------
# A PRÉVIA DO PAGAMENTO
# ---------------------------------------------------------------------------
def _sem_drive_nem_log(monkeypatch):
    """A prévia não pode subir arquivo nem registrar no log — se tentar, estoura."""
    from app.apps.analisesps import drive, folha_pagamento

    def proibido(*a, **k):
        raise AssertionError("a prévia tentou subir ou registrar")
    monkeypatch.setattr(drive, "subir_arquivo", proibido)
    monkeypatch.setattr(folha_pagamento, "_registrar", proibido)


def test_a_previa_sai_SEM_FECHAR_e_nao_sobe_nem_registra(app, monkeypatch):
    _preparar_folha_aberta(monkeypatch, dias=_dias_do_mes(), fechamento=None)
    _sem_drive_nem_log(monkeypatch)

    r = _como_mestre(app).get("/analisesps/folha/1/previa-pagamento?destino=beevale")
    assert r.status_code == 200, r.get_data(as_text=True)[:500]
    assert r.mimetype == "application/zip"
    assert "PREVIA" in r.headers["Content-Disposition"]

    pacote = zipfile.ZipFile(io.BytesIO(r.data))
    nomes = pacote.namelist()
    assert "LEIA-ME.txt" in nomes
    arquivos = [n for n in nomes if n.endswith(".xlsx")]
    # Um de pagamento (a conta 7011-4) e um de análise, todos marcados.
    assert len(arquivos) == 2
    assert all(n.startswith("PREVIA - NAO SUBIR - ") for n in arquivos)
    assert any("7011-4" in n for n in arquivos)
    leia = pacote.read("LEIA-ME.txt").decode("utf-8-sig")
    assert "NÃO ENVIAR AO PORTAL" in leia


def test_a_previa_respeita_QUEM_VOCE_TIROU(app, monkeypatch):
    """A prévia usa as mesmas marcações da tela: quem está fora não entra."""
    from app.apps.analisesps import folha_pagamento as fpg
    _preparar_folha_aberta(monkeypatch, dias=_dias_do_mes(), ajustes={
        "99713349334": {"fora": True, "motivo": "teste"}})
    _, apropriado, linhas = fpg.linhas_da_previa(1)
    assert "99713349334" not in {l["cpf"] for l in linhas}


def test_a_previa_SAI_MESMO_COM_AVISO_e_o_aviso_vai_no_leia_me(app, monkeypatch):
    """Sem conta para a obra o arquivo de verdade não sai; a prévia sai, dizendo."""
    from app.apps.analisesps import folha_pagamento
    _preparar_folha_aberta(monkeypatch, dias=_dias_do_mes())
    monkeypatch.setattr(folha_pagamento, "conta_por_obra", lambda: {})
    _sem_drive_nem_log(monkeypatch)

    r = _como_mestre(app).get("/analisesps/folha/1/previa-pagamento?destino=somapay")
    assert r.status_code == 200
    leia = zipfile.ZipFile(io.BytesIO(r.data)).read("LEIA-ME.txt").decode("utf-8-sig")
    assert "AVISO" in leia and "sem conta de pagamento" in leia
    assert "NÃO é gerado" in leia


def test_a_previa_usa_A_MESMA_REGRA_do_fechamento():
    """As linhas da prévia e as do fechamento saem da mesma função."""
    from app.apps.analisesps import folha_apropriacao_guardada as guardada
    apropriado = {"pessoas": [
        {"cpf": "1", "nome": "A", "por_obra": [{"obra": "x", "dias": 2, "valor": "10"}]},
        {"cpf": "2", "nome": "B", "fora": True,
         "por_obra": [{"obra": "y", "dias": 1, "valor": "5"}]}]}
    linhas = guardada.linhas_do_apropriado(apropriado)
    assert [l[0] for l in linhas] == ["1"]
    assert linhas[0][2] == "X"


def test_a_previa_e_SO_DO_MESTRE(app):
    from app.apps.analisesps import auth
    assert auth.e_so_do_mestre("analisesps.folha_previa_pagamento") is True
    r = como(app, SENHA_CONSULTA).get("/analisesps/folha/1/previa-pagamento")
    assert r.status_code in (302, 403, 404)


def test_o_botao_da_previa_esta_na_LATERAL_fora_do_formulario_dos_filtros(app, monkeypatch):
    """Trocar o destino da prévia não pode recarregar a tela (o formulário dos
    filtros se aplica sozinho a cada mudança)."""
    _preparar_folha_aberta(monkeypatch, dias=_dias_do_mes())
    html = _como_mestre(app).get("/analisesps/folha/1").get_data(as_text=True)
    principal = html.index('<main class="principal">')
    assert html.index('id="baixar-previa"') < principal
    assert html.index('id="previa-destino"') < html.index('id="form-filtros"')


# ---------------------------------------------------------------------------
# A LATERAL: caixinhas, como nas Solicitações
# ---------------------------------------------------------------------------
def test_os_filtros_da_folha_sao_CAIXINHAS_como_nas_solicitacoes(app, monkeypatch):
    _preparar_folha_aberta(monkeypatch, dias=_dias_do_mes())
    html = _como_mestre(app).get("/analisesps/folha/1").get_data(as_text=True)
    lateral = html[:html.index('<main class="principal">')]
    for nome in ("situacao", "conta", "obra", "obra_cadastro", "fase", "origem"):
        assert f'type="checkbox" name="{nome}"' in lateral, nome
    # Nem lista de escolha, nem bolinha, no formulário dos filtros.
    form = lateral[lateral.index('id="form-filtros"'):]
    form = form[:form.index("</form>")]
    assert "<select" not in form
    assert 'type="radio"' not in form
    assert "<h2>Filtros</h2>" in form and ">Limpar</a>" in form


def test_duas_obras_marcadas_mostram_quem_esta_em_UMA_OU_OUTRA(app, monkeypatch):
    _preparar_folha_aberta(monkeypatch, dias=_dias_em_duas_obras())
    cliente = _como_mestre(app)
    so_uma = cliente.get("/analisesps/folha/1?obra=XYZ9").get_data(as_text=True)
    assert "GERLANIO" in so_uma
    nenhuma = cliente.get("/analisesps/folha/1?obra=NADA").get_data(as_text=True)
    assert "Nenhum colaborador encontrado com esses filtros" in nenhuma
    duas = cliente.get("/analisesps/folha/1?obra=NADA&obra=XYZ9").get_data(as_text=True)
    assert "GERLANIO" in duas


def test_filtrar_aceita_LISTA_e_TEXTO():
    from app.apps.analisesps import folha_gestao as fg
    pessoas = [
        {"nome_na_tela": "A", "nome_contabilidade": "A", "fase": "ativos",
         "situacao": "paga", "por_obra": [{"obra": "X"}]},
        {"nome_na_tela": "B", "nome_contabilidade": "B", "fase": "afastados",
         "situacao": "fora", "por_obra": [{"obra": "Y"}]},
    ]
    assert len(fg._filtrar(pessoas, {"fase": "ativos"})) == 1
    assert len(fg._filtrar(pessoas, {"fase": ["ativos", "afastados"]})) == 2
    assert len(fg._filtrar(pessoas, {"situacao": ["fora"], "obra": ["x"]})) == 0
    assert len(fg._filtrar(pessoas, {"situacao": [], "obra": []})) == 2


def test_acao_na_lateral_e_BOTAO_nao_link_sublinhado(app, monkeypatch):
    """*"Botões com sublinhado outros sem"* e *"todas as folhas importadas é um
    texto"*."""
    _preparar_folha_aberta(monkeypatch, dias=_dias_do_mes())
    html = _como_mestre(app).get("/analisesps/folha/1").get_data(as_text=True)
    lateral = html[:html.index('<main class="principal">')]
    assert 'class="btn secundario"\n       href="/analisesps/folha/importar?lista=1"' in lateral \
        or 'class="btn secundario" href="/analisesps/folha/importar?lista=1"' in lateral
    assert "Divisão por obra e conta" in lateral
    assert "ver só elas" not in lateral and "ver quem" not in lateral


def test_o_termo_PRECISA_DA_SUA_MAO_saiu_da_tela(app, monkeypatch):
    """*"Precisa da sua mão antes de pagar, isso não é termo pra usar em sistema"*."""
    _preparar_folha_aberta(monkeypatch, dias=_dias_do_mes())
    html = _como_mestre(app).get("/analisesps/folha/1").get_data(as_text=True)
    assert "sua mão" not in html


# ---------------------------------------------------------------------------
# O RELATÓRIO
# ---------------------------------------------------------------------------
def test_o_relatorio_em_EXCEL_traz_a_lista_e_os_agrupamentos(app, monkeypatch):
    from openpyxl import load_workbook
    from app.apps.analisesps import folha_pagamento
    _preparar_folha_aberta(monkeypatch, dias=_dias_em_duas_obras())
    monkeypatch.setattr(folha_pagamento, "conta_por_obra",
                        lambda: {"CRE1": "7011-4", "XYZ9": "22069-1"})

    r = _como_mestre(app).get("/analisesps/folha/1/relatorio.xlsx")
    assert r.status_code == 200, r.get_data(as_text=True)[:500]
    livro = load_workbook(io.BytesIO(r.data))
    assert livro.sheetnames == ["Resumo", "Pessoas", "Por obra", "Por conta",
                                "Por obra da contabilidade", "Por setor"]
    nomes = [c.value for c in livro["Pessoas"]["A"][1:]]
    assert any("GERLANIO" in (n or "") for n in nomes)
    contas = {l[0].value for l in livro["Por conta"].iter_rows(min_row=2)}
    assert {"7011-4", "22069-1"} <= contas


def test_o_relatorio_sai_COM_OS_FILTROS_DA_TELA(app, monkeypatch):
    from openpyxl import load_workbook
    _preparar_folha_aberta(monkeypatch, dias=_dias_do_mes())
    r = _como_mestre(app).get("/analisesps/folha/1/relatorio.xlsx?q=LUELIA")
    livro = load_workbook(io.BytesIO(r.data))
    nomes = [c.value for c in livro["Pessoas"]["A"][1:]]
    assert nomes and all("LUELIA" in n for n in nomes)
    resumo = " ".join(str(c.value) for linha in livro["Resumo"].iter_rows()
                      for c in linha if c.value)
    assert 'Procura: "LUELIA"' in resumo
    assert "filtrado" in r.headers["Content-Disposition"]


def test_o_relatorio_POR_CONTA_traz_so_a_parte_daquela_conta(app, monkeypatch):
    """O dono, 02/10/2026: *"e se eu quiser só o PDF de uma determinada conta
    (…) divididos, ou juntos"* — e sem quem não entra no pagamento."""
    import zipfile
    from openpyxl import load_workbook
    from app.apps.analisesps import folha_pagamento
    _preparar_folha_aberta(monkeypatch, dias=_dias_em_duas_obras())
    monkeypatch.setattr(folha_pagamento, "conta_por_obra",
                        lambda: {"CRE1": "7011-4", "XYZ9": "22069-1"})
    cliente = _como_mestre(app)

    so = cliente.get("/analisesps/folha/1/relatorio.xlsx?relatorio_conta=7011-4")
    assert "conta 7011-4" in so.headers["Content-Disposition"]
    livro = load_workbook(io.BytesIO(so.data))
    contas = {l[0].value for l in livro["Por conta"].iter_rows(min_row=2)}
    assert contas == {"7011-4", "Total"}, "só a conta pedida"

    tudo = load_workbook(io.BytesIO(cliente.get(
        "/analisesps/folha/1/relatorio.xlsx").data))
    total = [l[3].value for l in tudo["Por conta"].iter_rows(min_row=2)
             if l[0].value == "Total"][0]
    soma = 0
    todas = [l[0].value for l in tudo["Por conta"].iter_rows(min_row=2)
             if l[0].value != "Total"]
    for conta in todas:
        um = load_workbook(io.BytesIO(cliente.get(
            "/analisesps/folha/1/relatorio.xlsx",
            query_string={"relatorio_conta": conta}).data))
        soma += [l[3].value for l in um["Por conta"].iter_rows(min_row=2)
                 if l[0].value == "Total"][0]
    assert round(soma, 2) == round(total, 2), "as contas somam o total"

    pacote = cliente.get("/analisesps/folha/1/relatorio.pdf?relatorio_conta=__cada")
    nomes = zipfile.ZipFile(io.BytesIO(pacote.data)).namelist()
    assert any("7011-4" in n for n in nomes) and any("22069-1" in n for n in nomes)
    assert all(n.endswith(".pdf") for n in nomes)


def test_o_relatorio_em_PDF_sai(app, monkeypatch):
    _preparar_folha_aberta(monkeypatch, dias=_dias_do_mes())
    r = _como_mestre(app).get("/analisesps/folha/1/relatorio.pdf")
    assert r.status_code == 200
    assert r.data.startswith(b"%PDF")


def test_formato_desconhecido_responde_404(app, monkeypatch):
    _preparar_folha_aberta(monkeypatch, dias=_dias_do_mes())
    r = _como_mestre(app).get("/analisesps/folha/1/relatorio.doc")
    assert r.status_code == 404


def test_os_botoes_do_relatorio_LEVAM_OS_FILTROS(app, monkeypatch):
    _preparar_folha_aberta(monkeypatch, dias=_dias_em_duas_obras())
    html = _como_mestre(app).get(
        "/analisesps/folha/1?obra=CRE1&obra=XYZ9").get_data(as_text=True)
    i = html.index('id="relatorio-xlsx"')
    trecho = html[i:i + 300]
    assert "obra=CRE1" in trecho and "obra=XYZ9" in trecho


def test_agrupamentos_SO_SOMAM_QUEM_VAI_RECEBER_e_dividem_por_dia():
    from app.apps.analisesps import folha_relatorio as fr
    pessoas = [
        {"cpf": "1", "entra": True, "valor": D("100.00"), "filial": "F1",
         "setor_curto": "S", "situacao_rotulo": "vai receber",
         "por_obra": [{"obra": "A", "dias": 3, "valor": D("60.00")},
                      {"obra": "B", "dias": 2, "valor": D("40.00")}]},
        {"cpf": "2", "entra": False, "valor": D("500.00"), "filial": "F1",
         "situacao_rotulo": "fora do pagamento",
         "por_obra": [{"obra": "A", "dias": 5, "valor": D("500.00")}]},
    ]
    g = fr.agrupamentos(pessoas, {"A": "111", "B": "222"})
    assert g["total_pago"] == D("100.00")
    por_obra = {o["obra"]: o for o in g["por_obra"]}
    assert por_obra["A"]["valor"] == D("60.00") and por_obra["A"]["dias"] == 3
    assert por_obra["A"]["percentual"] == D("60.00")
    por_conta = {c["conta"]: c for c in g["por_conta"]}
    assert por_conta["222"]["valor"] == D("40.00")
    # Por situação conta TODO MUNDO da lista.
    situacoes = {s["nome"]: s for s in g["por_situacao"]}
    assert situacoes["fora do pagamento"]["valor"] == D("500.00")
    assert g["por_filial"][0]["valor"] == D("100.00")


def test_obra_sem_conta_aparece_PRIMEIRO_no_por_conta():
    from app.apps.analisesps import folha_relatorio as fr
    pessoas = [{"cpf": "1", "entra": True, "valor": D("10"),
                "por_obra": [{"obra": "A", "dias": 1, "valor": D("5")},
                             {"obra": "Z", "dias": 1, "valor": D("5")}]}]
    g = fr.agrupamentos(pessoas, {"A": "111"})
    assert g["por_conta"][0]["conta"] == fr.SEM_CONTA
