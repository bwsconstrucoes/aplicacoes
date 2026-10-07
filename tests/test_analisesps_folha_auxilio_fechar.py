# -*- coding: utf-8 -*-
"""
Fechar o auxílio — 01/10/2026.

O botão "Conferir e gerar" da tela de alimentação e transporte levava a uma tela
que dizia "nada fechado": nenhuma tela fechava essas verbas, e o gerador só paga
verba fechada. Estes testes provam o caminho inteiro, com banco.
"""
import datetime as dt
from decimal import Decimal as D

import pytest

pytestmark = pytest.mark.banco

ATIVO, SAIU = "99713349334", "03513441363"


@pytest.fixture
def banco_auxilio(banco_analisesps):
    from app.apps.analisesps.db import conexao
    with conexao() as conn:
        conn.execute("INSERT INTO analisesps.referencias_rateio (tipo, nome, codigo) "
                     " VALUES ('obra', 'CREPEOLINDA', '111')")
        conn.execute("INSERT INTO analisesps.contas_diarios (codigo, conta_pagamento) "
                     " VALUES ('CREPEOLINDA', '50024')")
        for cpf, nome, fase, saida in [
                (ATIVO, "ATIVO UM", "Colaboradores Ativos", None),
                (SAIU, "SAIU UM", "Colaboradores Desligados", dt.date(2026, 8, 20))]:
            conn.execute(
                "INSERT INTO analisesps.colaborador "
                "  (cpf, nome, fase, data_saida, obra_codigo, valor_alimentacao, "
                "   modo_alimentacao) VALUES (?,?,?,?,?,?,?)",
                (cpf, nome, fase, saida, "CREPEOLINDA", D("300.00"), "Mês"))
        conn.commit()
    # Sem ponto, a obra é escolhida por ele (03/10/2026) — a do cadastro não
    # entra sozinha.
    from app.apps.analisesps import folha_auxilio as fx
    fx.gravar_extras(fx.ALIMENTACAO, 2026, 9, ATIVO, obra="CREPEOLINDA")
    return banco_analisesps


def test_fechar_a_alimentacao_libera_o_arquivo(banco_auxilio):
    from app.apps.analisesps import folha_auxilio as fx, folha_pagamento as fp
    feito = fx.fechar(fx.ALIMENTACAO, 2026, 9, "fim_de_mes", quem="MARCELO")
    assert feito["pessoas"] == 1 and feito["total"] == D("300.00")
    linhas = fp.linhas_para_pagar(2026, 9, "fim_de_mes", ["alimentacao"])
    assert [(l["cpf"], l["conta"], l["valor"]) for l in linhas] == [
        (ATIVO, "50024", D("300.00"))]
    assert fx.calcular(fx.ALIMENTACAO, 2026, 9)["fechamento"]["tipo"] == "fim_de_mes"


def test_quem_ja_SAIU_nao_entra_no_fechamento(banco_auxilio):
    from app.apps.analisesps import folha_auxilio as fx
    calculado = fx.calcular(fx.ALIMENTACAO, 2026, 9)
    saiu = next(p for p in calculado["pessoas"] if p["cpf"] == SAIU)
    assert saiu["desligado"] and not saiu["pagar"]


def test_fechar_num_pagamento_TIRA_o_fechamento_do_outro(banco_auxilio):
    """Senão a mesma verba do mês sairia duas vezes."""
    from app.apps.analisesps import (folha_apropriacao_guardada as guardada,
                                     folha_auxilio as fx)
    fx.fechar(fx.ALIMENTACAO, 2026, 9, "quinzena")
    fx.fechar(fx.ALIMENTACAO, 2026, 9, "fim_de_mes")
    assert guardada.fechamento(2026, 9, "quinzena", "alimentacao") is None
    assert guardada.fechamento(2026, 9, "fim_de_mes", "alimentacao") is not None


def test_gerar_o_auxilio_sobe_o_RELATORIO_em_PDF_da_conta(banco_auxilio, monkeypatch):
    """Dono, 03/10/2026: o relatório em PDF gerado junto, baixável em Arquivos
    gerados e com link no card."""
    from app.apps.analisesps import beevale, drive, folha_pagamento as fp
    subidos = []
    monkeypatch.setattr(beevale, "pasta_do_drive", lambda: ("PASTA", "teste"))
    monkeypatch.setattr(drive, "subir_arquivo", lambda conteudo, nome, pasta, **k:
                        subidos.append((nome, k.get("mime"), conteudo[:5]))
                        or {"id": f"d{len(subidos)}", "link": f"https://drive/{len(subidos)}"})
    fp.gerar_direto("alimentacao", {"ano": 2026, "mes": 9, "pagamento": "fim_de_mes"},
                    "beevale", quem="MARCELO")
    pdfs = [s for s in subidos if s[1] == fp.MIME_PDF]
    assert len(pdfs) == 1 and pdfs[0][2] == b"%PDF-" and "50024" in pdfs[0][0]
    rodada = fp.rodadas()[0]
    assert rodada["relatorios"]["50024"]["link"]
    assert rodada["total"] == D("300.00"), "o cadastro não soma no total"
    # A planilha de cadastro da conta, gerada junto (07/10/2026).
    cadastros = [s for s in subidos if s[0].startswith("Cadastro conta 50024")]
    assert len(cadastros) == 1 and "BeeVale" in cadastros[0][0]
    assert rodada["cadastros"]["50024"]["link"]
    assert rodada["cadastros"]["50024"]["id"] in rodada["ids"], "excluir leva junto"


def test_a_REGRA_DE_RATEIO_divide_o_auxilio_entre_as_obras(banco_auxilio):
    """Dono, 03/10/2026: *"o rateio das obras serve sim para alimentação e
    transporte e diaristas"*. A regra manda sobre o ponto, como na folha."""
    from app.apps.analisesps import folha_auxilio as fx, folha_rateio as fr
    # A obra escolhida à mão manda sobre a regra; tirada, vale a regra.
    fx.gravar_extras(fx.ALIMENTACAO, 2026, 9, ATIVO, obra="")
    fr.gravar({"nome": "Supervisores", "obras": [
        {"obra": "CREPEOLINDA", "percentual": "60"},
        {"obra": "CREPEAREIAS", "percentual": "40"}],
        "pessoas": [{"cpf": ATIVO}]}, "MARCELO")
    calculado = fx.calcular(fx.ALIMENTACAO, 2026, 9)
    p = next(x for x in calculado["pessoas"] if x["cpf"] == ATIVO)
    assert p["obra_de_onde"] == "regra" and p["regra"] == "Supervisores"
    assert [(x["obra"], x["valor"]) for x in p["rateio"]] == [
        ("CREPEOLINDA", D("180.00")), ("CREPEAREIAS", D("120.00"))]
    por_obra = {o["obra"]: o["total"] for o in calculado["por_obra"]}
    assert por_obra == {"CREPEOLINDA": D("180.00"), "CREPEAREIAS": D("120.00")}

    fx.fechar(fx.ALIMENTACAO, 2026, 9, "fim_de_mes", quem="MARCELO")
    from app.apps.analisesps import folha_pagamento as fp
    pagas = fp.linhas_para_pagar(2026, 9, "fim_de_mes", ["alimentacao"])
    assert sorted((l["obra"], l["valor"]) for l in pagas) == [
        ("CREPEAREIAS", D("120.00")), ("CREPEOLINDA", D("180.00"))]


def test_A_TRAVA_E_O_ARQUIVO_gerar_de_novo_pede_confirmacao(banco_auxilio, monkeypatch):
    """O dono, 06/10/2026, ao tirar o botão "Fechar": *"a trava é o arquivo que a
    gente gerou"*. Nada impedia gerar o mesmo auxílio duas vezes (pagaria duas)."""
    from app.apps.analisesps import beevale, drive, folha_pagamento as fp
    subidos = []
    monkeypatch.setattr(beevale, "pasta_do_drive", lambda: ("PASTA", "teste"))
    monkeypatch.setattr(drive, "subir_arquivo", lambda conteudo, nome, pasta, **k:
                        subidos.append(nome)
                        or {"id": f"d{len(subidos)}", "link": f"https://drive/{len(subidos)}"})
    monkeypatch.setattr(drive, "mover_para_lixeira", lambda arquivo: None)
    pedido = {"ano": 2026, "mes": 9, "pagamento": "fim_de_mes"}
    assert fp.resumo_direto("alimentacao", pedido, "beevale")["resumo"]["ja_gerado"] == []
    primeira = fp.gerar_direto("alimentacao", pedido, "beevale", quem="MARCELO")

    resumo = fp.resumo_direto("alimentacao", pedido, "beevale")["resumo"]
    assert len(resumo["ja_gerado"]) == 1 and not resumo["pode_gerar"]
    # O auxílio é um por mês: na quinzena também conta.
    quinzena = dict(pedido, pagamento="quinzena")
    assert fp.resumo_direto("alimentacao", quinzena, "beevale")["resumo"]["ja_gerado"]
    with pytest.raises(fp.ErroDoPagamento) as e:
        fp.gerar_direto("alimentacao", pedido, "beevale", quem="MARCELO")
    assert "já foi gerado" in str(e.value) and "MARCELO" in str(e.value)
    antes = len(subidos)
    fp.gerar_direto("alimentacao", pedido, "beevale", quem="MARCELO", forcar=True)
    assert len(subidos) > antes, "com a confirmação, gera"

    # Excluir as gerações destrava.
    fp.excluir_arquivos([a["id"] for a in fp.log()], quem="MARCELO")
    assert fp.resumo_direto("alimentacao", pedido, "beevale")["resumo"]["ja_gerado"] == []
    assert primeira


@pytest.mark.parametrize("tipo", ["alimentacao", "transporte"])
def test_o_DESMARCADO_e_salvo_NAO_sai_no_arquivo(banco_auxilio, monkeypatch, tipo):
    """O dono, 06/10/2026: *"o que eu tô selecionando pra pagar não tá afetando o
    que eu gero pra pagar. Revisa tanto alimentação quanto transporte."*"""
    from app.apps.analisesps import beevale, drive, folha_auxilio as fx, folha_pagamento as fp
    from app.apps.analisesps.db import conexao
    OUTRO = "11144477735"
    with conexao() as conn:
        conn.execute("UPDATE analisesps.colaborador SET valor_transporte = 200, "
                     " modo_transporte = 'Mensal' WHERE cpf = ?", (ATIVO,))
        conn.execute(
            "INSERT INTO analisesps.colaborador (cpf, nome, fase, obra_codigo, "
            "  valor_alimentacao, modo_alimentacao, valor_transporte, modo_transporte) "
            " VALUES (?, 'OUTRO DOIS', 'Colaboradores Ativos', 'CREPEOLINDA', 300, "
            "  'Mês', 200, 'Mensal')", (OUTRO,))
        conn.commit()
    for cpf in (ATIVO, OUTRO):
        fx.gravar_extras(tipo, 2026, 9, cpf, obra="CREPEOLINDA")
    monkeypatch.setattr(beevale, "pasta_do_drive", lambda: ("PASTA", "teste"))
    subidos = []
    monkeypatch.setattr(drive, "subir_arquivo", lambda conteudo, nome, pasta, **k:
                        subidos.append(nome) or {"id": f"d{len(subidos)}", "link": "x"})
    assert {p["cpf"] for p in fx.calcular(tipo, 2026, 9)["pessoas"] if p["pagar"]} == {ATIVO, OUTRO}

    fx.salvar_selecao(tipo, 2026, 9, [{"cpf": ATIVO, "pagar": True},
                                      {"cpf": OUTRO, "pagar": False}], quem="MARCELO")
    pedido = {"ano": 2026, "mes": 9, "pagamento": "fim_de_mes"}
    assert fp.resumo_direto(tipo, pedido, "beevale")["resumo"]["pessoas"] == 1
    fp.gerar_direto(tipo, pedido, "beevale", quem="MARCELO")
    assert [l["cpf"] for l in fp.linhas_para_pagar(2026, 9, "fim_de_mes", [tipo])] == [ATIVO]


def test_DESLIGADO_marcado_a_mao_NAO_vai_no_arquivo(banco_auxilio, monkeypatch):
    """Os 77 de 06/10/2026: desligados marcados à mão (um "marcar todos" salvo)
    ficavam escondidos da lista e saíam no arquivo. Agora a marcação não vale:
    desligado não recebe auxílio — e salvar a seleção não grava a marcação."""
    from app.apps.analisesps import folha_auxilio as fx, folha_pagamento as fp
    from app.apps.analisesps.db import conexao
    fx.gravar_extras(fx.ALIMENTACAO, 2026, 9, SAIU, obra="CREPEOLINDA")
    # A marcação antiga, gravada antes da regra (direto, como estava no banco).
    fx.gravar_ajuste(fx.ALIMENTACAO, 2026, 9, SAIU, pagar=True, quem="ANTES")
    calculado = fx.calcular(fx.ALIMENTACAO, 2026, 9)
    saiu = next(p for p in calculado["pessoas"] if p["cpf"] == SAIU)
    assert not saiu["pagar"] and calculado["quantos_a_pagar"] == 1
    linhas = fp.resumo_direto(fx.ALIMENTACAO, {"ano": 2026, "mes": 9,
                                               "pagamento": "fim_de_mes"}, "beevale")["linhas"]
    assert SAIU not in {l["cpf"] for l in linhas}
    # E pela tela também não marca.
    r = fx.salvar_selecao(fx.ALIMENTACAO, 2026, 9, [{"cpf": SAIU, "pagar": True}], quem="X")
    assert SAIU in r["ignorados"]


def test_ARQUIVOS_GERADOS_mostram_a_SITUACAO_de_cada_SP(banco_auxilio, monkeypatch):
    """07/10/2026: *"fizesse uma leitura do número da SP pra saber o status de
    cada uma e colocar uma tag"*. Lido da base das SPs; a que ainda não chegou à
    base diz isso, sem inventar status."""
    from app.apps.analisesps import beevale, drive, folha_pagamento as fp
    from tests.test_analisesps_banco import semear, sp
    monkeypatch.setattr(beevale, "pasta_do_drive", lambda: ("PASTA", "teste"))
    subidos = []
    monkeypatch.setattr(drive, "subir_arquivo", lambda conteudo, nome, pasta, **k:
                        subidos.append(nome) or {"id": f"d{len(subidos)}", "link": "x"})
    fp.gerar_direto("alimentacao", {"ano": 2026, "mes": 9, "pagamento": "fim_de_mes"},
                    "beevale", quem="MARCELO")
    pagamento = next(a for a in fp.log() if a["destino"] == "beevale")
    fp.registrar_card(pagamento["id"], "900111,900222", "https://pipefy/900111")
    semear([sp("900111", status_pgt="Pago"),
            sp("900222", status_pgt="Pagar", agendado="Agendado")])
    rodadas = fp.rodadas()
    fp.situacao_das_sps(rodadas)
    sps = rodadas[0]["pagamentos"][0]["sps"]
    assert [(s["id"], s["status_pgt"], s["status_agend"], s["na_base"]) for s in sps] == [
        ("900111", "Pago", "", True), ("900222", "Pagar", "Agendado", True)]
    fp.registrar_card(pagamento["id"], "900333", "")
    rodadas = fp.rodadas()
    fp.situacao_das_sps(rodadas)
    assert rodadas[0]["pagamentos"][0]["sps"][0]["na_base"] is False


def test_TRANSPORTE_categorias_DIARIAS_entram_e_a_AUDITORIA_nao_perde_ninguem(banco_auxilio):
    """07/10/2026 — o dono comparou com a base do script: 27 de 38 colaboradores
    sumiam do transporte por serem "Diário", "Vale Transporte" ou "Diário e Vale
    Transporte". Os exemplos dele, com 22 dias úteis (setembro/2026)."""
    import io
    import openpyxl
    from app.apps.analisesps import folha_auxilio as fx
    from app.apps.analisesps.db import conexao
    gente = [("11144477735", "MARCIO", "Diário", "12.90"),
             ("22233344405", "EVANDRO", "Vale Transporte", "9.00"),
             ("33344455560", "ADEMIR", "Diário e Vale Transporte", "9.00"),
             ("44455566603", "MANASSES", "Mensal", "198.00"),
             ("55566677715", "ERRADO", "Diário", "198.00"),
             ("66677788820", "SEMREGRA", "Quinzenal Estranho", "50.00")]
    with conexao() as conn:
        for cpf, nome, cat, valor in gente:
            conn.execute(
                "INSERT INTO analisesps.colaborador (cpf, nome, fase, obra_codigo, "
                "  valor_transporte, modo_transporte) VALUES (?,?,?,?,?,?)",
                (cpf, nome, "Colaboradores Ativos", "CREPEOLINDA", D(valor), cat))
        conn.commit()
    for cpf, *_ in gente:
        fx.gravar_extras(fx.TRANSPORTE, 2026, 9, cpf, obra="CREPEOLINDA")

    calculado = fx.calcular(fx.TRANSPORTE, 2026, 9)
    valor = {p["nome"]: (p["pagar"], p["valor"]) for p in calculado["pessoas"]}
    assert valor["MARCIO"] == (True, D("283.80"))
    assert valor["EVANDRO"] == (True, D("198.00"))
    assert valor["ADEMIR"] == (True, D("198.00"))
    assert valor["MANASSES"] == (True, D("198.00"))
    assert valor["ERRADO"][0] is False, "198 'por dia' é o valor do mês: não paga"

    dados = fx.auditoria(fx.TRANSPORTE, 2026, 9, calculado)
    nomes = {l["nome"]: l for l in dados["linhas"]}
    assert set(nomes) >= {n for _, n, _, _ in gente}, "ninguém some"
    assert not nomes["ERRADO"]["pagar"] and "parece o valor do MÊS" in nomes["ERRADO"]["motivo"]
    assert not nomes["SEMREGRA"]["pagar"] and "não reconhecida" in nomes["SEMREGRA"]["motivo"]
    assert nomes["MARCIO"]["qtd"] == 22 and nomes["MARCIO"]["valor_base"] == D("283.80")
    r = dados["resumo"]
    assert r["base"] == r["processados"] == len(calculado["pessoas"])
    assert r["a_pagar"] + r["nao_pagos"] == r["base"]
    assert r["diferenca"] == r["total_base"] - r["total_final"]
    assert sum((l["diferenca"] for l in dados["diferencas"]), D("0")) == r["diferenca"]

    livro = openpyxl.load_workbook(io.BytesIO(fx.auditoria_xlsx(fx.TRANSPORTE, 2026, 9, dados)))
    assert livro.sheetnames == ["Auditoria", "Validação", "Não pagos", "Diferenças"]
    assert livro["Auditoria"].cell(row=1, column=20).value == "Motivo da Exclusão ou Ajuste"
