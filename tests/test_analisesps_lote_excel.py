# -*- coding: utf-8 -*-
"""Análise de SPs — o relatório do lote em Excel de verdade.

POR QUE ELE EXISTE ALÉM DO CSV. O CSV sai em blocos e aguenta a base larga; o
Excel monta o arquivo inteiro na memória, e por isso é do LOTE. O que ele
resolve e o CSV não: **valor é número** (soma e ordena sem converter) e **código
é texto** (o Excel não transforma um código de 47 dígitos em notação
científica, o que devolve o número arredondado e irrecuperável).
"""
from __future__ import annotations

import io

import pytest

from app.apps.analisesps import lote_excel


def abrir(dados: bytes):
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(dados))


def sp(sp_id="1443253401", valor=1000.0, **campos):
    base = {"id": sp_id, "vencimento_d": "2026-09-20",
            "credor": "SERTAO CASA E CONSTRUCAO",
            "documento": "29.066.773/0001-52",
            "descricao": "AQUISIÇÃO DE MATERIAL ELÉTRICO",
            "centro_custo": "CREPETRIUNFO", "tipo_despesa": "Ferramentas",
            "forma_pagamento": "Pix", "conta": "Bradesco - 50024-0",
            "nf": "1430", "status_pgt": "Pagar", "status_agend": "Agendar",
            "valor_num": valor}
    base.update(campos)
    return base


def montado(*grupos):
    if not grupos:
        grupos = (("Pagar amanhã", [sp()]),)
    lista = [{"titulo_exibido": t, "linhas": l, "nao_encontrados": [],
              "total": sum(x["valor_num"] for x in l)} for t, l in grupos]
    return {"quantidade": sum(len(l) for _, l in grupos),
            "total_geral": sum(g["total"] for g in lista), "grupos": lista}


# ---------------------------------------------------------------------------
# O QUE O EXCEL RESOLVE E O CSV NÃO
# ---------------------------------------------------------------------------
def test_o_valor_e_NUMERO_e_nao_texto():
    """É o motivo de este arquivo existir: no CSV o valor vai como texto
    "1.000,00" para o Excel brasileiro entender, e quem recebe tem de converter
    antes de somar."""
    aba = abrir(lote_excel.de_um_lote(montado())).active
    celula = aba.cell(row=2, column=len(lote_excel.COLUNAS))
    assert isinstance(celula.value, (int, float))
    assert celula.value == 1000.0
    assert "R$" in celula.number_format


def test_o_numero_da_SP_vai_como_TEXTO():
    """O Excel transforma um código comprido em 1,23457E+46 — e o número volta
    arredondado, irrecuperável. Com um ID de 10 dígitos ainda aguenta; com o
    código de barras de 47, não."""
    aba = abrir(lote_excel.de_um_lote(montado())).active
    assert aba.cell(row=2, column=2).number_format == "@"
    assert aba.cell(row=2, column=2).value == "1443253401"


def test_a_descricao_e_a_obra_entram_na_planilha():
    """As duas colunas que o dono pediu no relatório impresso valem aqui
    também — quem confere no Excel precisa das mesmas informações."""
    aba = abrir(lote_excel.de_um_lote(montado())).active
    cabecalho = [c.value for c in aba[1]]
    assert "Descrição" in cabecalho and "Obra" in cabecalho
    linha = {c: v for c, v in zip(cabecalho, [c.value for c in aba[2]])}
    assert linha["Descrição"] == "AQUISIÇÃO DE MATERIAL ELÉTRICO"
    assert linha["Obra"] == "CREPETRIUNFO"


def test_o_grupo_do_lote_vira_coluna():
    """O lote é organizado em grupos, e essa organização é trabalho de quem o
    montou. Achatar tudo numa lista só perderia a decisão dele."""
    dados = lote_excel.de_um_lote(montado(
        ("Pagar amanhã", [sp("1")]), ("Semana que vem", [sp("2")])))
    aba = abrir(dados).active
    assert [aba.cell(row=n, column=1).value for n in (2, 3)] == [
        "Pagar amanhã", "Semana que vem"]


def test_o_total_e_FORMULA_e_acompanha_o_filtro():
    """Quem filtra a planilha depois espera o total acompanhar. Um número fixo
    mentiria em silêncio — e é justamente para filtrar que se pede Excel."""
    aba = abrir(lote_excel.de_um_lote(montado())).active
    valores = [c.value for c in aba[aba.max_row]]
    assert "TOTAL" in valores
    formula = [v for v in valores if isinstance(v, str) and v.startswith("=")]
    assert formula and "SUBTOTAL(109" in formula[0], (
        "o total não acompanha o filtro")


def test_a_planilha_abre_com_o_cabecalho_travado_e_filtro():
    aba = abrir(lote_excel.de_um_lote(montado())).active
    assert aba.freeze_panes == "A2"
    assert aba.auto_filter.ref


# ---------------------------------------------------------------------------
# O NOME DA ABA — o Excel recusa coisas que um título de lote tem
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("titulo,esperado", [
    ("Pagar 15/09", "Pagar 15-09"),          # barra: o Excel recusa
    ("Depois: urgente", "Depois- urgente"),  # dois pontos: idem
    ("Lote [1]", "Lote -1-"),
    ("", "Lote"),
])
def test_titulo_de_lote_vira_nome_de_aba_que_o_Excel_aceita(titulo, esperado):
    """Um título de lote é texto livre, escrito por gente — "Pagar 15/09" tem
    barra. Deixar passar faria o arquivo INTEIRO não abrir por causa disso."""
    assert lote_excel._nome_de_aba(titulo, set()) == esperado


def test_nome_de_aba_comprido_e_cortado_em_31():
    assert len(lote_excel._nome_de_aba("a" * 60, set())) == 31


def test_duas_pessoas_com_o_mesmo_nome_nao_colidem():
    """O Excel recusa duas abas com o mesmo nome, e o nome vem do que a pessoa
    digitou — dois "Marcelo" acontecem."""
    usados = set()
    assert lote_excel._nome_de_aba("Marcelo", usados) == "Marcelo"
    assert lote_excel._nome_de_aba("Marcelo", usados) == "Marcelo (2)"


# ---------------------------------------------------------------------------
# TODOS OS LOTES
# ---------------------------------------------------------------------------
def test_todos_os_lotes_saem_UMA_ABA_POR_PESSOA():
    """Trinta lotes empilhados numa aba só seria o mesmo problema que o CSV já
    tem."""
    planilha = abrir(lote_excel.de_todos([
        {"nome": "Marcelo", "montado": montado()},
        {"nome": "Yara", "montado": montado()}]))
    assert planilha.sheetnames == ["Resumo", "Marcelo", "Yara"]


def test_o_resumo_vem_PRIMEIRO_e_soma_todo_mundo():
    """Quem abre um arquivo de oito abas quer ver o tamanho do todo antes de
    escolher em qual entrar."""
    planilha = abrir(lote_excel.de_todos([
        {"nome": "Marcelo", "montado": montado(("G", [sp("1", 100.0)]))},
        {"nome": "Yara", "montado": montado(("G", [sp("2", 250.0)]))}]))
    assert planilha.sheetnames[0] == "Resumo"
    resumo = planilha["Resumo"]
    assert [c.value for c in resumo[2]] == ["Marcelo", 1, 100.0]
    assert [c.value for c in resumo[3]] == ["Yara", 1, 250.0]


def test_a_soma_do_resumo_NAO_inclui_a_propria_celula():
    """REFERÊNCIA CIRCULAR: lendo a última linha DEPOIS de criar a linha do
    total, a soma incluía a si mesma — e o Excel abre com erro em vez de com o
    número. Achado conferindo a fórmula gerada."""
    planilha = abrir(lote_excel.de_todos([
        {"nome": "A", "montado": montado()},
        {"nome": "B", "montado": montado()}]))
    resumo = planilha["Resumo"]
    total = [c.value for c in resumo[resumo.max_row] if c.value]
    formula = [v for v in total if isinstance(v, str) and v.startswith("=")][0]
    assert formula == "=SUM(C2:C3)", (
        f"a soma pega a própria linha do total: {formula}")


def test_lote_vazio_nao_estoura_nem_inventa_total():
    dados = lote_excel.de_um_lote({"grupos": [], "total_geral": 0})
    aba = abrir(dados).active
    assert aba.max_row == 1, "só o cabeçalho"


def test_o_teto_de_linhas_existe_e_e_respeitado(monkeypatch):
    """O `openpyxl` monta o arquivo inteiro na memória — não há como sair em
    blocos. O teto é o que impede uma exportação enorme de derrubar o serviço
    dos outros 17 módulos."""
    monkeypatch.setattr(lote_excel, "MAXIMO_LINHAS", 3)
    dados = lote_excel.de_um_lote(montado(
        ("G", [sp(str(n)) for n in range(10)])))
    aba = abrir(dados).active
    # 1 cabeçalho + 3 linhas + 1 em branco + 1 total
    assert sum(1 for linha in aba.iter_rows(min_row=2)
               if linha[1].value and str(linha[1].value).isdigit()) == 3
