# -*- coding: utf-8 -*-
"""
DESPESAS COM COLABORADORES (DC) — 03/10/2026.

A aba "Data" da planilha da DC vira uma tela no padrão das outras folhas, e o
arquivo do BeeVale/SomaPay sai daqui. Estes testes provam o caminho com banco:
ler a aba, enriquecer pelo cadastro, selecionar, gerar, esconder o gerado e
acusar a duplicidade.
"""
from decimal import Decimal as D

import pytest

from tests.test_analisesps_telas import _como_mestre, app  # noqa: F401 — a fixture

pytestmark = pytest.mark.banco

ATIVO, OUTRO, FORA = "99713349334", "03513441363", "11122233396"

CABECALHO = ["ID", "Data Solicitação", "Vencimento", "Centro de Custo", "Tipo",
             "Período", "Valor diária?", "Responsável", "Requerente", "CPF",
             "Qtd", "Valor", "Valor diária", "Descrição"]
DATA = [
    CABECALHO,
    ["900100", "01/10/2026", "05/10/2026", "CREPEOLINDA", "Despesas com Alimentação",
     "", "", "JOAO", "MARIA", "997.133.493-34", "", "150,00", "", "Lanche da equipe"],
    ["900100", "01/10/2026", "05/10/2026", "CREPEOLINDA", "Despesas com Alimentação",
     "", "", "JOAO", "MARIA", "035.134.413-63", "3", "", "", "Lanche da equipe"],
    ["900200", "02/10/2026", "06/10/2026", "CREPEOLINDA", "Diárias de Viagem",
     "", "", "JOAO", "MARIA", "111.222.333-96", "", "80,00", "", "Fora do cadastro"],
    ["", "", "", "", "", "", "", "", "", "", "", "", "", ""],
]


@pytest.fixture
def banco_dc(banco_analisesps, monkeypatch):
    from app.apps.analisesps import dc, folha_cards
    from app.apps.analisesps.db import conexao
    with conexao() as conn:
        conn.execute("INSERT INTO analisesps.referencias_rateio (tipo, nome, codigo) "
                     " VALUES ('obra', 'CREPEOLINDA', '111')")
        conn.execute("INSERT INTO analisesps.contas_diarios (codigo, conta_pagamento) "
                     " VALUES ('CREPEOLINDA', '50024')")
        for cpf, nome, diaria in [(ATIVO, "ATIVO UM", None), (OUTRO, "OUTRO DOIS", D("40.00"))]:
            conn.execute(
                "INSERT INTO analisesps.colaborador (cpf, nome, fase, cargo, valor_diaria) "
                " VALUES (?,?,?,?,?)", (cpf, nome, "Colaboradores Ativos", "PEDREIRO", diaria))
        conn.commit()

    abas = {dc.ABA_DATA: DATA,
            "Data base BeeVale": [["Plano Financeiro", "Carteira BeeVale"],
                                  ["Despesas com Alimentação", "Auxílio Alimentação"]]}

    def ler_aba(nome):
        if nome not in abas:
            raise RuntimeError("aba não existe")
        return abas[nome]
    monkeypatch.setattr(dc, "_ler_aba", ler_aba)
    monkeypatch.setattr(folha_cards, "_plano_financeiro", lambda: ({
        folha_cards._chave("Despesas com Alimentação"): [
            {"nome": "Despesas com Alimentação", "record_id": "R-ALI", "codigo_omie": "2.01.01"}],
        folha_cards._chave("Diárias de Viagem"): [
            {"nome": "Diárias de Viagem", "record_id": "R-VIA", "codigo_omie": "2.02.01"}],
    }, ""))
    dc._cache.clear()
    yield banco_analisesps
    dc._cache.clear()


@pytest.fixture
def cliente_mestre(banco_dc, app):  # noqa: F811
    return _como_mestre(app)


def test_a_aba_Data_vira_LINHAS_enriquecidas_pelo_cadastro(banco_dc):
    from app.apps.analisesps import dc
    calculado = dc.calcular()
    por_cpf = {p["cpf"]: p for p in calculado["pessoas"]}
    assert set(por_cpf) == {ATIVO, OUTRO, FORA}       # cabeçalho e linha vazia fora

    ativo = por_cpf[ATIVO]
    assert ativo["nome"] == "ATIVO UM" and ativo["cargo"] == "PEDREIRO"
    assert ativo["valor"] == D("150.00")                # o valor informado
    assert ativo["conta"] == "50024" and ativo["omie"] == "111"
    assert ativo["categoria"] == "2.01.01" and ativo["record_id"] == "R-ALI"
    assert ativo["carteira"] == "Auxílio Alimentação"   # da aba Data base BeeVale
    assert ativo["pagar"] and ativo["link_card"].endswith("/900100")

    # Sem valor informado: quantidade × diária do cadastro.
    assert por_cpf[OUTRO]["valor"] == D("120.00")

    # Fora do cadastro: não entra (o banco recusa sem o nome).
    fora = por_cpf[FORA]
    assert fora["impossivel"] and not fora["pagar"]
    assert fora["carteira"] == dc.CARTEIRA_PADRAO       # tipo sem carteira = Produção

    assert calculado["total"] == D("270.00") and calculado["quantos_a_pagar"] == 2
    assert calculado["cards"] == 2


@pytest.mark.parametrize("linha,cadastrada,esperado,origem", [
    # Pede a diária cadastrada: quantidade × cadastro, mesmo com valor informado.
    ({"valor_diaria_flag": "Valor da diária cadastrada", "quantidade": "2",
      "valor": "500,00"}, D("90.00"), D("180.00"), "cadastro"),
    ({"valor_diaria_flag": "Sim", "quantidade": "1,5"}, D("90.00"), D("135.00"), "cadastro"),
    # Pede a cadastrada e o cadastro não tem: sem valor (não entra).
    ({"valor_diaria_flag": "Sim", "quantidade": "2"}, None, D("0.00"), ""),
    # Sem pedido: o informado manda; sem ele, qtd × diária informada; sem ela,
    # qtd × diária do cadastro.
    ({"valor_diaria_flag": "Não", "valor": "70,00", "quantidade": "2"}, D("90.00"),
     D("70.00"), "informado"),
    ({"quantidade": "2", "valor_diaria": "80,00"}, D("90.00"), D("160.00"), "informada"),
    ({"quantidade": "2"}, D("90.00"), D("180.00"), "cadastro"),
    ({"quantidade": "2"}, None, D("0.00"), ""),
])
def test_o_VALOR_da_linha_pela_diaria_cadastrada_ou_informada(linha, cadastrada,
                                                               esperado, origem):
    """Dono, 03/10/2026: *"A DC pode vir com valor ou não quando se trata de
    diária. Ela pode pedir que seja paga pelo valor de diária cadastrada, nesse
    caso o sistema calcula."*"""
    from app.apps.analisesps import dc
    valor, motivos, de_onde = dc.valor_da_linha(linha, cadastrada)
    assert (valor, de_onde) == (esperado, origem)
    if not valor:
        assert motivos


def test_a_lista_AGRUPA_por_obra_com_o_total_a_pagar(banco_dc):
    from app.apps.analisesps import dc
    calculado = dc.calcular()
    grupos = dc.agrupar(calculado["pessoas"], "tipo_despesa")
    assert [(g["rotulo"], g["linhas"], g["a_pagar"], g["total"]) for g in grupos] == [
        ("Despesas com Alimentação", 2, 2, D("270.00")),
        ("Diárias de Viagem", 1, 0, D("0.00"))]
    assert grupos[1]["pendencias"] == 1
    assert dc.agrupar(calculado["pessoas"], "")[0]["rotulo"] == ""
    assert calculado["resumos"]["conta"] == [
        {"rotulo": "50024", "linhas": 2, "pessoas": 2, "total": D("270.00")}]


def test_a_SELECAO_guarda_so_a_excecao(banco_dc):
    from app.apps.analisesps import dc
    chave = next(p["chave"] for p in dc.calcular()["pessoas"] if p["cpf"] == OUTRO)
    dc.salvar_selecao([{"chave": chave, "pagar": False}], quem="MARCELO")
    calculado = dc.calcular()
    assert not next(p for p in calculado["pessoas"] if p["cpf"] == OUTRO)["pagar"]
    assert calculado["total"] == D("150.00")
    # Remarcar volta ao calculado e apaga a exceção.
    dc.salvar_selecao([{"chave": chave, "pagar": True}], quem="MARCELO")
    assert dc.ajustes() == {}


def test_trocar_a_OBRA_muda_a_conta(banco_dc):
    from app.apps.analisesps import dc
    from app.apps.analisesps.db import conexao
    with conexao() as conn:
        conn.execute("INSERT INTO analisesps.contas_diarios (codigo, conta_pagamento) "
                     " VALUES ('CREPEAREIAS', '60010')")
        conn.commit()
    chave = next(p["chave"] for p in dc.calcular()["pessoas"] if p["cpf"] == ATIVO)
    dc.escolher_obra(chave, "crepeareias", quem="MARCELO")
    p = next(p for p in dc.calcular()["pessoas"] if p["chave"] == chave)
    assert p["obra"] == "CREPEAREIAS" and p["obra_trocada"] and p["conta"] == "60010"
    dc.escolher_obra(chave, "", quem="MARCELO")
    assert dc.ajustes() == {}


def test_GERAR_sobe_o_arquivo_guarda_o_lote_e_TIRA_as_linhas_da_tela(banco_dc, monkeypatch):
    import io
    import openpyxl
    from app.apps.analisesps import beevale, dc, drive, folha_pagamento as fp
    subidos = []
    monkeypatch.setattr(beevale, "pasta_do_drive", lambda: ("PASTA", "teste"))
    monkeypatch.setattr(drive, "subir_arquivo", lambda conteudo, nome, pasta, **k:
                        subidos.append((nome, k.get("mime"), conteudo))
                        or {"id": f"d{len(subidos)}", "link": f"https://drive/{len(subidos)}"})
    saida = fp.gerar_direto("dc", {}, "beevale", quem="MARCELO")
    pagamento = [a for a in saida["arquivos"] if a["destino"] == "beevale"]
    assert len(pagamento) == 1 and pagamento[0]["total"] == D("270.00")
    assert "Despesas com colaboradores" in pagamento[0]["nome"]

    # A carteira de cada linha vai no arquivo do BeeVale.
    conteudo = next(c for n, m, c in subidos if n == pagamento[0]["nome"])
    aba = openpyxl.load_workbook(io.BytesIO(conteudo)).active
    carteiras = {aba.cell(row=r, column=1).value: aba.cell(row=r, column=3).value
                 for r in range(2, aba.max_row + 1)}
    assert carteiras == {"ATIVO UM": "Auxílio Alimentação", "OUTRO DOIS": "Auxílio Alimentação"}
    assert any(m == fp.MIME_PDF for _, m, _ in subidos)   # o relatório da conta

    # O lote guarda o que foi; a tela não mostra mais.
    analise = next(a["id"] for a in saida["arquivos"] if a["destino"] == fp.ANALISE)
    lote = dc.lote_da_analise(analise)
    assert sorted(l["cpf"] for l in lote["linhas"]) == sorted([ATIVO, OUTRO])
    assert all(l["destino"] == "beevale" for l in lote["linhas"])
    calculado = dc.calcular()
    assert {p["cpf"] for p in calculado["pessoas"]} == {FORA}
    assert calculado["geradas"] == 2
    assert {p["cpf"] for p in dc.calcular(mostrar_geradas=True)["pessoas"]} == {ATIVO, OUTRO, FORA}

    # Excluir a geração devolve as linhas para a tela.
    fp.excluir_arquivos([a["id"] for a in saida["arquivos"]], quem="MARCELO")
    assert dc.lote_da_analise(analise) is None
    assert {p["cpf"] for p in dc.calcular()["pessoas"]} == {ATIVO, OUTRO, FORA}


def test_a_DUPLICIDADE_acusa_o_mesmo_CPF_e_tipo_dos_ultimos_10_dias(banco_dc):
    from app.apps.analisesps import dc
    linhas = [l for l in dc.linhas_a_pagar() if l["cpf"] == ATIVO]
    dc.registrar_lote(None, [dict(linhas[0], chave="outra-chave", card_id="800000")],
                      {"50024": "beevale"}, quem="MARCELO")
    p = next(p for p in dc.calcular()["pessoas"] if p["cpf"] == ATIVO)
    assert p["duplicidade"] and any("possível duplicidade" in m for m in p["motivos"])
    assert "800000" in " ".join(p["motivos"])


def test_sem_a_aba_das_CARTEIRAS_vale_a_TABELA_GRAVADA(banco_dc, monkeypatch):
    """06/10/2026: *"não pode ser assim (…) eu já disse quais são os tipos, por
    que não grava logo"*. Sem a aba, vale a tabela que ele passou — sem aviso."""
    from app.apps.analisesps import dc
    monkeypatch.setattr(dc, "_ler_aba", lambda nome: DATA if nome == dc.ABA_DATA
                        else (_ for _ in ()).throw(RuntimeError("sem aba")))
    dc._cache.clear()
    calculado = dc.calcular()
    ativo = next(p for p in calculado["pessoas"] if p["cpf"] == ATIVO)
    assert ativo["carteira"] == "Auxílio Alimentação"
    assert not any("Data base BeeVale" in a for a in calculado["avisos"])
    mapa, aviso = dc.carteiras(recarregar=True)
    assert aviso == ""
    assert {k: mapa[dc._sem_acento(k)] for k in dc.CARTEIRAS_DA_DC} == {
        "Despesas com Alimentação": "Auxílio Alimentação",
        "Despesas com Transporte": "Despesas com Transporte",
        "Diárias": "Diárias",
        "Gratificações e Extras": "Gratiticações e Extras",
        "Produção": "Produção",
        "Salários e Ordenados": "Diárias"}


def test_a_aba_com_cabecalho_TIPO_DC_e_lida(banco_dc, monkeypatch):
    """O cabeçalho da aba dele é "Tipo DC | Tipo BeeVale" — antes não era
    reconhecido."""
    from app.apps.analisesps import dc
    assert dc._carteiras_de([["Tipo DC", "Tipo BeeVale"],
                             ["Bonificação", "Gratiticações e Extras"]]) == {
        "bonificacao": "Gratiticações e Extras"}


def test_a_TELA_mostra_as_linhas_e_os_blocos_da_lateral(banco_dc, cliente_mestre):
    resposta = cliente_mestre.get("/analisesps/folha/dc")
    assert resposta.status_code == 200
    html = resposta.get_data(as_text=True)
    assert "Despesas com colaboradores" in html
    assert "ATIVO UM" in html and "card 900100" in html
    assert "Gerar arquivos" in html and "Planilha de cadastro" in html
    assert "Divisão por obra e conta" in html
    assert "Auxílio Alimentação" in html


def test_a_tela_abre_AGRUPADA_por_obra_e_troca_o_agrupamento(banco_dc, cliente_mestre):
    html = cliente_mestre.get("/analisesps/folha/dc").get_data(as_text=True)
    assert html.count('class="grupo-cab"') == 1          # uma obra só
    assert "Resumo do que vai ser pago" in html and "Tipo de despesa</th>" in html
    html = cliente_mestre.get("/analisesps/folha/dc?agrupar=card_id").get_data(as_text=True)
    assert html.count('class="grupo-cab"') == 2          # dois cards
    html = cliente_mestre.get("/analisesps/folha/dc?agrupar=").get_data(as_text=True)
    assert 'class="grupo-cab"' not in html and "card 900100" in html


def test_a_tela_diz_QUANDO_a_aba_foi_lida(banco_dc, cliente_mestre):
    """Dono, 05/10/2026: *"seria interessante ter a informação do momento em que
    ela foi atualizada, data e hora"*."""
    from app.apps.analisesps import dc
    html = cliente_mestre.get("/analisesps/folha/dc").get_data(as_text=True)
    assert "lida em <b>" in html
    assert dc.calcular()["lida_em"] is not None


def test_a_tela_filtra_por_TIPO_DE_DESPESA(banco_dc, cliente_mestre):
    html = cliente_mestre.get(
        "/analisesps/folha/dc?tipo_despesa=Di%C3%A1rias+de+Viagem").get_data(as_text=True)
    assert "card 900200" in html and "card 900100" not in html


def test_o_RELATORIO_da_DC_sai_em_pdf(banco_dc, cliente_mestre):
    resposta = cliente_mestre.get("/analisesps/folha/dc/relatorio.pdf")
    assert resposta.status_code == 200
    assert resposta.data[:5] == b"%PDF-"


def test_o_RELATORIO_da_tela_com_as_GERADAS_traz_a_SP(banco_dc, monkeypatch):
    """Com "mostrar as já geradas", elas entram no relatório — cada uma com a SP."""
    from app.apps.analisesps import dc
    calculado = dc.calcular()
    for p in calculado["pessoas"]:
        p["gerada"], p["sp"] = True, {"id": "1234567", "link": ""}
    montado = dc.montado_do_relatorio(calculado, geradas_entram=True)
    entram = [p for p in montado["pessoas"] if p["entra"]]
    assert entram and all("SP 1234567" in p["situacao_rotulo"] for p in entram)
    assert not any(p["entra"] for p in dc.montado_do_relatorio(calculado)["pessoas"]), \
        "sem pedir, gerada não entra (é o relatório do que falta pagar)"


# ---------------------------------------------------------------------------
# NO PIPEFY: uma SP por conta, rateio por obra e por categoria, e os cards de
# origem marcados e movidos (o que o `geraspbeevale.gs` fazia).
# ---------------------------------------------------------------------------
def _gerar_dc(monkeypatch, destino="beevale", trocados=None):
    from app.apps.analisesps import beevale, drive, folha_pagamento as fp
    subidos = []
    trocados = [] if trocados is None else trocados
    monkeypatch.setattr(drive, "substituir_conteudo", lambda arquivo, conteudo, mime:
                        trocados.append((arquivo, mime)))
    monkeypatch.setattr(beevale, "pasta_do_drive", lambda: ("PASTA", "teste"))
    monkeypatch.setattr(drive, "subir_arquivo", lambda conteudo, nome, pasta, **k:
                        subidos.append(nome)
                        or {"id": f"d{len(subidos)}", "link": f"https://drive/{len(subidos)}"})
    saida = fp.gerar_direto("dc", {}, destino, quem="MARCELO")
    return next(a["id"] for a in saida["arquivos"] if a["destino"] == fp.ANALISE)


@pytest.fixture
def com_duas_categorias(banco_dc, monkeypatch):
    """Uma terceira linha, de outro tipo de despesa, na mesma conta."""
    from app.apps.analisesps import dc
    data = DATA[:-1] + [
        ["900300", "02/10/2026", "06/10/2026", "CREPEOLINDA", "Diárias de Viagem",
         "", "", "JOAO", "MARIA", "997.133.493-34", "", "50,00", "", "Viagem"]]
    original = dc._ler_aba
    monkeypatch.setattr(dc, "_ler_aba", lambda nome: data if nome == dc.ABA_DATA
                        else original(nome))
    dc._cache.clear()
    return banco_dc


def test_a_SP_da_DC_tem_RATEIO_POR_CATEGORIA_e_move_os_cards(com_duas_categorias,
                                                             monkeypatch, cliente_mestre):
    from app.apps.analisesps import dc, folha_cards as fcd, pipefy
    from tests.test_analisesps_folha_cards import PipefyFalso
    pipe = PipefyFalso(monkeypatch)
    movidos, trocados, desenhados = [], [], []
    monkeypatch.setattr(pipefy, "mover_card", lambda card, fase, *a, **k:
                        movidos.append((str(card), int(fase))))
    analise = _gerar_dc(monkeypatch, trocados=trocados)
    from app.apps.analisesps import folha_relatorio as fr
    original_pdf = fr.pdf
    monkeypatch.setattr(fr, "pdf", lambda dados: desenhados.append(dados)
                        or original_pdf(dados))

    vista = fcd.previa(analise)
    assert vista["bloqueios"] == []
    assert vista["avisos"] == [], "gerada com o relatório em PDF: nada a avisar"
    grupo = vista["grupos"][0]
    assert grupo["verba"] == "dc" and grupo["cards_de_origem"] == ["900100", "900300"]
    sp = grupo["sps"][0]
    assert sp["conta"] == "50024" and sp["valor"] == D("320.00")
    assert sp["tipo_sp"] == "R-ALI"                    # o da maior parte (270 × 50)
    assert "Mais de um tipo de despesa" in sp["observacao"]

    fcd.lancar(analise, quem="MARCELO")
    criada = pipe.do_pipe(fcd.PIPE_SP)
    assert len(criada) == 1
    valores = criada[0]["valores"]
    assert valores["valor"] == "320.00" and valores["tipo_de_despesa"] == "R-ALI"
    rateio = valores["rateio_m_ltiplo"]
    assert '"codigo_categoria":"2.01.01"' in rateio and '"codigo_categoria":"2.02.01"' in rateio
    assert '"cCodDep":"111"' in rateio
    descricao = valores["descri_o"]
    assert "IDs de origem: 900100, 900300" in descricao
    assert "Valor BeeVale (+1,50%): R$ 324,80" in descricao
    assert "Planilha de pagamento: https://drive/" in descricao
    assert "Relatório (PDF): https://drive/" in descricao
    # "Folha de Pgt" + "BeeVale" (05/10/2026).
    assert pipe.atualizacoes_de(criada[0]["id"])["etiquetas"] == ["318116255", "317521565"]

    # Os cards de origem: mover_card = Sim e a fase de processados.
    assert sorted(movidos) == [("900100", dc.FASE_PROCESSADO), ("900300", dc.FASE_PROCESSADO)]
    assert pipe.atualizacoes_de("900100") == {dc.CAMPO_MOVER: "Sim"}
    assert dc.lote_da_analise(analise)["cards_movidos"]

    # O RELATÓRIO EM PDF ganha o número da SP (dono, 06/10/2026), gravado POR
    # CIMA do arquivo da geração — o link que está no card continua valendo.
    numero = criada[0]["id"]
    from app.apps.analisesps import folha_pagamento as fp
    pdf_da_conta = next(a for a in fp.log() if a["destino"] == fp.RELATORIO)
    assert trocados == [(pdf_da_conta["drive_id"], "application/pdf")]
    assert f"SP nº {numero}" in desenhados[-1]["subtitulo"]
    assert desenhados[-1]["quantas"] == 3, "as três linhas pagas (dois cards), não zero"
    assert all(f"SP {numero}" in p["situacao_rotulo"] for p in desenhados[-1]["pessoas"])
    # E na tela, com as geradas à mostra, cada linha diz a SP que a pagou.
    assert {p["sp"]["id"] for p in dc.calcular(mostrar_geradas=True)["pessoas"]
            if p["gerada"]} == {numero}
    tela = cliente_mestre.get("/analisesps/folha/dc?geradas=1").get_data(as_text=True)
    assert f">SP {numero}</a>" in tela

    with pytest.raises(fcd.ErroDosCards):
        fcd.lancar(analise)
    assert len(pipe.criados) == 1 and len(movidos) == 2 and len(trocados) == 1


def test_se_o_Pipefy_recusa_MOVER_lancar_de_novo_so_termina_a_mudanca(banco_dc, monkeypatch):
    from app.apps.analisesps import dc, folha_cards as fcd, pipefy
    from tests.test_analisesps_folha_cards import PipefyFalso
    pipe = PipefyFalso(monkeypatch)
    movidos, recusar = [], {"vez": True}

    def mover(card, fase, *a, **k):
        if recusar["vez"]:
            raise pipefy.ErroDoPipefy("fase não permite")
        movidos.append(str(card))
    monkeypatch.setattr(pipefy, "mover_card", mover)
    analise = _gerar_dc(monkeypatch, destino="somapay")

    saida = fcd.lancar(analise, quem="MARCELO")
    assert saida["avisos"] and saida["cards_movidos"] == []
    assert len(pipe.do_pipe(fcd.PIPE_SP)) == 1
    assert "automa_o_2" not in pipe.do_pipe(fcd.PIPE_SP)[0]["valores"]   # só no BeeVale
    assert not dc.lote_da_analise(analise)["cards_movidos"]
    # A tela libera o botão só para terminar (antes ficava escondido).
    assert fcd.previa(analise)["so_terminar"]

    recusar["vez"] = False
    saida = fcd.lancar(analise, quem="MARCELO")
    assert saida["cards_movidos"] == ["900100"] and movidos == ["900100"]
    assert len(pipe.do_pipe(fcd.PIPE_SP)) == 1, "nenhuma SP a mais"
    assert dc.lote_da_analise(analise)["cards_movidos"]
    assert not fcd.previa(analise)["so_terminar"], "tudo terminado: botão fechado"


def test_DIARIAS_usa_a_linha_de_SALARIOS_E_ORDENADOS_do_plano(banco_dc, monkeypatch):
    """07/10/2026: lançar a DC barrou *tipo de despesa "Diárias" sem Código Omie
    (…) sem Record ID* — "Diárias" não existe no Plano Financeiro. Vale a linha
    que os diaristas já usam, "Salários e Ordenados" (só quando o nome falta)."""
    from app.apps.analisesps import dc, folha_cards
    plano = {folha_cards._chave("Salários e Ordenados"): [
        {"nome": "Salários e Ordenados", "record_id": "R-SAL", "codigo_omie": "2.03.01"}]}
    assert dc.classificacao_no_plano("Diárias", plano) == ("2.03.01", "R-SAL")
    assert dc.classificacao_no_plano("Outra Coisa", plano) == ("", "")
    # Se um dia o plano tiver "Diárias", vale a linha dela.
    plano[folha_cards._chave("Diárias")] = [
        {"nome": "Diárias", "record_id": "R-DIA", "codigo_omie": "2.09.09"}]
    assert dc.classificacao_no_plano("Diárias", plano) == ("2.09.09", "R-DIA")


def test_DC_ja_GERADA_sem_classificacao_lanca_resolvendo_AGORA(banco_dc, monkeypatch):
    """A geração grava a categoria de cada linha; a que foi gerada antes da regra
    (sem categoria) é resolvida na hora de lançar, sem gerar de novo."""
    from app.apps.analisesps import dc, folha_cards as fcd
    from tests.test_analisesps_folha_cards import PipefyFalso
    data = DATA[:1] + [
        ["900500", "02/10/2026", "06/10/2026", "CREPEOLINDA", "Diárias",
         "", "", "JOAO", "MARIA", "997.133.493-34", "", "100,00", "", "Diária"]]
    original = dc._ler_aba
    monkeypatch.setattr(dc, "_ler_aba", lambda nome: data if nome == dc.ABA_DATA
                        else original(nome))
    dc._cache.clear()
    PipefyFalso(monkeypatch)
    monkeypatch.setattr(fcd, "_plano_financeiro", lambda: ({}, ""))   # sem a linha
    analise = _gerar_dc(monkeypatch)
    assert dc.lote_da_analise(analise)["linhas"][0]["categoria"] == ""
    assert any("Diárias" in b for b in fcd.previa(analise)["bloqueios"])

    monkeypatch.setattr(fcd, "_plano_financeiro", lambda: ({
        fcd._chave("Salários e Ordenados"): [
            {"nome": "Salários e Ordenados", "record_id": "R-SAL", "codigo_omie": "2.03.01"}]},
        ""))
    vista = fcd.previa(analise)
    assert not any("Diárias" in b for b in vista["bloqueios"]), vista["bloqueios"]
    sp = vista["grupos"][0]["sps"][0]
    assert sp["tipo_sp"] == "R-SAL" and '"codigo_categoria":"2.03.01"' in sp["rateio"]
