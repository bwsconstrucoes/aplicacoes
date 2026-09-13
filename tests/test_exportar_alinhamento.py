"""A planilha tem de pôr cada valor debaixo do cabeçalho CERTO.

A ARMADILHA, achada em 12/09/2026 ao fazer a resposta do assistente virar
relatório: a linha é um dicionário, e o exportador pegava `linha.values()` —
os valores na ordem em que o dicionário foi montado, não na ordem das colunas.

Dois jeitos de isso dar errado, e nenhum deles avisa:

  · o dicionário montado em outra ordem → valor debaixo do cabeçalho errado;
  · a linha sem um dos campos → todos os seguintes andam uma casa à esquerda.

Nos dois casos o arquivo sai bonito, sem erro nenhum, com o número errado no
lugar certo. É o pior defeito que um relatório pode ter, porque quem lê não
tem como desconfiar.

Roda SEM banco: é aritmética de planilha, não consulta.
"""
from __future__ import annotations

from app.apps.erp.core.comum import exportar as svc

# O formato que as respostas do assistente devolvem (`_resposta`, em
# core/perguntas/respostas.py): coluna com chave, rótulo e tipo.
COLUNAS = [
    {"chave": "numero_sp", "rotulo": "SP", "tipo": "texto"},
    {"chave": "credor", "rotulo": "Credor", "tipo": "texto"},
    {"chave": "valor", "rotulo": "Valor", "tipo": "dinheiro"},
    {"chave": "vencimento", "rotulo": "Vencimento", "tipo": "data"},
]


def _linhas(colunas, linhas):
    return svc._linhas_como_listas(linhas, svc.normalizar_colunas(colunas))


def test_dicionario_fora_de_ordem_sai_na_ordem_das_colunas():
    """O caso que quebra em silêncio: quem escreve a resposta monta o
    dicionário na ordem que lhe convém, e não tem obrigação nenhuma de seguir
    a ordem das colunas."""
    linha = {"valor": "1500,00", "numero_sp": "SP0001",
             "vencimento": "2026-10-10", "credor": "ALFA LTDA"}
    assert _linhas(COLUNAS, [linha]) == [
        ["SP0001", "ALFA LTDA", "1500,00", "2026-10-10"]]


def test_linha_com_campo_faltando_deixa_o_buraco_no_lugar_certo():
    """Sem isto, a ausência do credor empurraria valor e vencimento uma casa
    para a esquerda — e o vencimento apareceria na coluna do valor."""
    linha = {"numero_sp": "SP0002", "valor": "800,00",
             "vencimento": "2026-11-01"}
    assert _linhas(COLUNAS, [linha]) == [
        ["SP0002", None, "800,00", "2026-11-01"]]


def test_campo_a_mais_na_linha_nao_entra_no_arquivo():
    """A resposta pode carregar o id do registro para a tela usar. Ele não é
    informação para quem lê o relatório, e não pode virar uma coluna torta."""
    linha = {"numero_sp": "SP0003", "credor": "BETA", "valor": "10,00",
             "vencimento": "2026-12-01", "titulo_id": 99}
    assert _linhas(COLUNAS, [linha]) == [
        ["SP0003", "BETA", "10,00", "2026-12-01"]]


def test_tela_que_manda_so_os_rotulos_continua_funcionando():
    """As telas de lista mandam a tabela já desenhada, sem chave nenhuma. Elas
    não podem quebrar por causa desta correção — são a maioria dos usos."""
    assert _linhas(["SP", "Credor"], [{"a": "SP0004", "b": "GAMA"}]) == [
        ["SP0004", "GAMA"]]
    assert _linhas(["SP", "Credor"], [["SP0005", "DELTA"]]) == [
        ["SP0005", "DELTA"]]


def test_a_planilha_de_verdade_sai_com_o_valor_debaixo_do_cabecalho_certo():
    """Prova de ponta a ponta: gera o .xlsx e lê de volta."""
    import io

    from openpyxl import load_workbook

    conteudo = svc.para_excel(
        "Teste de alinhamento", COLUNAS,
        [{"valor": 1500, "numero_sp": "SP0001",
          "vencimento": "10/10/2026", "credor": "ALFA LTDA"}],
        quem="teste")
    aba = load_workbook(io.BytesIO(conteudo)).active

    # acha a linha do cabeçalho (o arquivo tem título e filtros antes)
    linha_cab = None
    for linha in aba.iter_rows(values_only=True):
        if linha and linha[0] == "SP":
            linha_cab = linha
            break
    assert linha_cab is not None, "não achei o cabeçalho na planilha"
    coluna_do_credor = list(linha_cab).index("Credor")

    dados = None
    achou_cab = False
    for linha in aba.iter_rows(values_only=True):
        if achou_cab and linha and linha[0]:
            dados = linha
            break
        if linha and linha[0] == "SP":
            achou_cab = True
    assert dados is not None, "não achei a linha de dados"
    assert dados[coluna_do_credor] == "ALFA LTDA", (
        "o credor saiu debaixo do cabeçalho errado — é o defeito que este "
        "arquivo inteiro existe para impedir")
